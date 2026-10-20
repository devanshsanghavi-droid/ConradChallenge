"""
importers.py: read any model file safely (iteration 5, task 15), and make an imported model mean what the model
assumes (task 15b).

    python -m residualmap.importers FILE [FILE ...]               # what each file is, by its content, and what to do
    python -m residualmap.importers --prepare FILE [--out DIR]    # the model readiness report for an EPANET model

A utility sends what it has: an EPANET .inp saved by a vendor package (a byte-order mark, ';' comment lines first, a
Windows code page, UTF-16), the same file renamed .txt or kept as a .bak, an emailed zip, a CAD drawing, a GIS layer, a
vendor model database or a spreadsheet.  This module identifies a file by its content, never by its name, and turns
an EPANET model into a file WNTR reads exactly as the model means it, with a plain readiness report.

- sniff(data, name) reads the first 64 KB (the whole file for zips, SQLite, OLE and JSON, all in memory; the whole text
  when the first 64 KB looks like the start of an EPANET file but holds none of its node or link sections) and returns
  a Sniff: kind, plain label, detail, confidence, whether the file name decided anything (by_extension), whether the
  name marks a backup copy, and advice in plain words.  The content decides the kind.  The name either adds a note
  (a .wtg.mdb or .wtg.sqlite name says which program uses the database, at high confidence) or breaks a tie the
  content cannot settle (at low confidence), and either way the result says so.  sniff never raises: a file it cannot
  read comes back as 'unknown', with advice.  Every test runs in time linear in the bytes it reads (no backtracking
  pattern over the file, csv.Sniffer on a bounded sample with no quote characters, SQLite queries under a step
  budget).
- decode_text(data) decodes a UTF-8 byte-order mark, then a UTF-16 one, then strict UTF-8, then cp1252 (any of the five
  bytes cp1252 leaves undefined is read as latin-1, the rest as cp1252), and says which.  The BOMs come first: a file
  with a UTF-8 BOM also decodes as plain UTF-8, with a leading U+FEFF that WNTR rejects (ENSyntaxError 201 at line 1).
- scan_inp_ids(text) checks the IDs in the text, before WNTR parses: WNTR keeps the later of two lines with one ID in a
  section without a word, loads a junction that shares a tank's ID and then the simulator fails, and stops on an ID
  of 32 characters with an AssertionError; EPANET itself counts bytes, so an accented ID that fits in a Windows code
  page can be too long once the copy is UTF-8.  Each of these blocks, with 'export with IDs, not labels'.  Pattern and
  curve IDs get the same length rule.
- prepare_inp(src, out_dir, disinfectant) returns (inp_path, Readiness).  A clean UTF-8 file given as a path comes back
  as the same path, untouched (so the app's upload name and grid cache do not change); clean bytes are written as
  <sha1 of the bytes>[:8]_<name>.  A file that needed a fix is written as <sha1 of the new text>[:8]_<stem>.inp in
  out_dir, which callers keep git-ignored (outputs/app_uploads, outputs/private/inputs).  The original is never
  changed.  A file that is not a model, or has a blocking problem, gives (None, Readiness) and leaves nothing in
  out_dir: a copy is written under a temporary name and put in place only after WNTR has read it, an entry point
  was found and EPANET ran the model's first hour (a test run, set up as the app's first run, in a temporary folder;
  WNTR reads a junction no pipe connects, and EPANET then stops).  An upload name too long for a file name once the
  digest is added is shortened (prepared_name).  Nothing in a file makes it raise: any failure after the file is read
  is a blocking problem.
- Task 15b, the model's physics, each a row of the readiness report (fix_reactions_text, dw_to_hw, _physics_report):
  reaction orders other than 1 are set to 1; per-pipe BULK and per-tank TANK lines are cleared and a LIMITING
  POTENTIAL set to 0 (per-pipe WALL lines are ignored, since the grid sets every wall rate); a Darcy-Weisbach model is
  converted to Hazen-Williams (each pipe's C matched to its headloss at its reference velocity) and kept only when the
  water age of the copy stays within DW_GATE_MEDIAN_H and DW_GATE_P95_H of the file's, and Chezy-Manning blocks; the
  start clock time is recorded and a weekly pattern warned about; with pressure-control valves every junction at
  elevation 0 is flagged and one elevation everywhere blocks; coordinates are typed as schematic, to scale in feet or
  metres, or missing; sources at nodes the tool does not dose (boosters), a source the dose overrides, and more than
  one inlet are warned about; pumps with no control, rule or speed pattern, no time patterns, zero demand and
  disconnected parts are warned about, and a WaterGEMS export's connector pipes reported; one C value on every pipe
  means the old-pipe factor cannot be learned (grid_for gives 'full_fixed_gamma'); and the run cost gives the grid's
  runs, its file and its temporary files, and minutes from one timed 7-day run, with a hard warning above
  LARGE_JUNCTIONS and size_gate's refusal when the build would leave under 1.5 GB free.  A model blocked now leaves
  no copy behind, not even a content-named copy an earlier version accepted.
- Task 15b also fixed three of task 15's leftovers in the sniff: a SWMM model whose SWMM-only sections come after the
  first 64 KB (the whole text is searched for them), a zip holding a model under another name (its members' contents
  are typed), and a UTF-8 file with a few stray bytes of a Windows code page (decode_text reads it as UTF-8 with those
  bytes as cp1252, and the report says the file mixes two encodings; since the second review the byte pairs a cp1252
  file makes by chance count for cp1252 and lines are counted, not characters: _utf8_kind, _mixed_utf8).
- Task 16, grab logs and tap maps (read_table, guess_log_columns, read_grab_log, read_tap_map): a grab log as the
  utility's lab or field system exports it, CSV in any encoding decode_text reads and with any of the separators
  above (decimal commas read), or an Excel workbook through pandas and openpyxl (typed date and time cells kept;
  zip limits first; since the review of task 16, openpyxl in read-only mode with row and column caps, a cover sheet
  passed over; a text table's title or 'Lab:' lines and Excel's 'sep=;' line skipped).  The columns are guessed from
  the header ('Sample Point', 'Sample Date/Time', 'Total Chlorine (mg/L)', or the template's names) or given; one
  datetime column or a date and a time; US or European dates, month names, and a column whose dates all read two
  ways is refused with the question until a date format is given ('MDY', 'DMY', 'YMD' or a strptime format).
  '<0.02', 'ND' and 'BDL' are censored below detection, '>2.2' above the measuring range: kept and listed with their
  reason, never fitted.  Free and total chlorine are never mixed: a log with both columns is refused until its
  reading column is chosen, and then only that one is read; a column naming the other species is refused; a
  long-format export (one row per analyte) keeps only the chosen species' rows (until the final verification this
  failed when the species sat in brackets or the analytes did not name chlorine in the reader's words: see below).  A reading's unit is read from its
  header or a units column (µg/L divided by 1000), and one above 10 mg/L is not used.  Tap IDs that Excel turned
  into numbers are matched by number, with a warning.  A tap with 12 or more readings a day, minutes apart, on 2 or
  more days is an online analyzer and keeps one reading a day, the one nearest 10:00 (or the hour given), unless it
  is kept whole.  Most grab samples outside 06:00 to 18:00 warns of a UTC export.  Each time of day becomes a model
  hour: the time minus START CLOCKTIME, mod 24 h, minutes truncated (model_hour; clock_label and
  daytime_model_hours go the other way, for display and the route's 07:00 to 17:00 window).  Every row that is not
  used is listed with its reason (dropped_lines, at most 20 runs of rows per reason).  xlsx_bytes writes the Excel
  template.  After the final verification of task 16: the species is read with the bracket contents kept ('Chlorine
  (Free)', 'Cl2 (total)'), 'Chloramine' is total chlorine, and 'Free Cl', 'Total Cl' and the codes FCL2, CL2-F, TCL2,
  CL2-T name chlorine; two reading columns that are not both of one named species ask which is the reading; a
  long-format export whose analytes cannot pick the species (codes the reader does not know, two chlorine analytes
  that name no species, an analyte-like column beside a plain 'Result') is refused with the question until the
  analyte value is chosen, instead of read whole; a detection limit, pressure or temperature column is no reading;
  units without a slash ('ugL', 'mg Cl2/L', 'mg/litre') and a mangled 'Âµ' are read, and a log with readings above
  and below 10 mg/L warns of an unnamed µg/L; a text's lines are counted before it is split, and records stop at the
  row limit; a comma title before a semicolon export, and an Excel title row of two cells, are skipped; an Excel 00:00
  among timed datetimes is a time only at an online analyzer's tap; a time column holding a date and a time gives its
  time; a tap ID with a superscript digit is a tap like any other; a two-column junction export is no tap map; and a
  not-used reason that names a value is one line for all its values when more than 5 differ, the list capped at 30
  lines.
- Section headers and their names are read the way WNTR's reader reads them: after any whitespace (a form feed
  included), and a singular name such as [PIPE] is the [PIPES] section; the ID scan also stops at [END], as WNTR does.
- Zips are read in memory only, never extracted to disk.  A member with an absolute path or a '..' part is refused,
  and so is a zip whose members add up to more than 200 MB uncompressed.  The member that is read (the model) is
  refused if it is over 10 MB uncompressed and compressed more than 100 to 1; members that are only listed are not
  held to that ratio, because a shapefile's space-padded DBF packs past it at city size.  Files that macOS or Windows
  add to a zip (__MACOSX/, ._ copies, .DS_Store, Thumbs.db, desktop.ini) are left out of the listing.

Nothing here is applied inside simulate.load, so the bundled networks run bit for bit as before.  Uses the standard
library, WNTR and what WNTR itself needs (numpy, networkx).
"""
from __future__ import annotations

import argparse
import codecs
import contextlib
import csv
import hashlib
import io
import json
import math
import os
import re
import sqlite3
import struct
import sys
import tempfile
import time
import warnings
import zipfile
from dataclasses import asdict, dataclass, field

HEAD_BYTES = 64 * 1024                 # how much of a file the text and signature tests read
MAX_ID_BYTES = 31                      # EPANET's longest ID, in bytes of the UTF-8 copy (WNTR counts characters)
ZIP_MAX_TOTAL = 200_000_000           # uncompressed bytes in all members (200 MB)
ZIP_RATIO_MAX = 100                    # compression ratio refused, for the member that is read ...
ZIP_RATIO_MIN_SIZE = 10_000_000        # ... when it is over this many uncompressed bytes (10 MB)
ZIP_MAX_MEMBERS = 10000
JSON_MAX_BYTES = 50 * 1024 * 1024      # JSON files parsed whole up to this size, by their first 64 KB above it
TEXT_SCAN_MAX = 50 * 1024 * 1024       # an EPANET-looking text whose sections start after 64 KB is searched whole up to this
JSON_MAX_DEPTH = 64                    # nesting followed when coordinates are collected from a GeoJSON geometry
CSV_SNIFF_CHARS = 4096                 # csv.Sniffer sees this much, its quote characters taken out (_sniffer_sample)
XML_MAX_SKIP = 1000                    # comments, processing instructions and DOCTYPEs skipped before the root element
SQLITE_MAX_STEPS = 2000                # progress-handler calls (1,000 SQLite steps each) before a query is stopped
MAX_NAME_BYTES = 160                   # a prepared file's name, in UTF-8 bytes (a file name has 255, and the grid
                                       # cache adds up to about 60 for its own words)
TEST_RUN_ROW = "Test run"              # the readiness row of the EPANET run of the first hour
RUN_COST_ROW = "Run cost"              # the readiness row of the grid's size and time (task 15b)
LARGE_WARNING_START = "Large model:"   # the start of the warning for more than LARGE_JUNCTIONS junctions
FLOAT32_BYTES = 4                      # the grid cache stores ln C as float32: runs x 24 hours x junctions
CP1252_UNDEFINED = b"\x81\x8d\x8f\x90\x9d"   # the five bytes cp1252 leaves undefined


def _cp1252_gaps(err):
    """Error handler for cp1252: a byte cp1252 leaves undefined is read as latin-1 (the same code point)."""
    if isinstance(err, UnicodeDecodeError):
        return err.object[err.start:err.end].decode("latin-1"), err.end
    raise err


codecs.register_error("residualmap_cp1252_gaps", _cp1252_gaps)

# provenance of a readiness row
FROM_FILE, DERIVED, ASSUMED, MISSING = "from your file", "derived", "assumed", "missing, please supply"
SOURCES = (FROM_FILE, DERIVED, ASSUMED, MISSING)
DISINFECTANT_WORDS = {"free_chlorine": "free chlorine", "chloramine": "chloramine (total chlorine)"}

INP_SECTIONS = ("TITLE", "JUNCTIONS", "RESERVOIRS", "TANKS", "PIPES", "PUMPS", "VALVES", "TAGS", "DEMANDS", "STATUS",
                "PATTERNS", "CURVES", "CONTROLS", "RULES", "ENERGY", "EMITTERS", "QUALITY", "SOURCES", "REACTIONS",
                "MIXING", "TIMES", "REPORT", "OPTIONS", "COORDINATES", "VERTICES", "LABELS", "BACKDROP", "END",
                "LEAKAGE")
CORE_SECTIONS = {"JUNCTIONS", "PIPES", "RESERVOIRS", "TANKS"}
# section headers that only a SWMM model (sewers, storm drains) or an EPANET-MSX reaction file has; both also use some
# EPANET headers ([JUNCTIONS], [PIPES], [TANKS]), so they are told apart before the EPANET test
# (only names no EPANET file and few other text files use: a generic one like [MAP] or [STORAGE] decides nothing)
SWMM_SECTIONS = ("OUTFALLS", "CONDUITS", "XSECTIONS", "SUBCATCHMENTS", "SUBAREAS", "RAINGAGES", "ORIFICES", "WEIRS",
                 "DIVIDERS", "TRANSECTS", "INFILTRATION", "DWF", "EVAPORATION")
MSX_SECTIONS = ("SPECIES", "COEFFICIENTS", "TERMS", "PARAMETERS", "DIFFUSIVITY")   # [SPECIES], which MSX requires, decides
# a header starts a line, after any whitespace (WNTR strips every kind, a form feed included), once the line ends are
# made '\n' (so a classic Mac CR-only file, or CR-CR-LF line ends from a double conversion, are read like any other);
# it ends at whitespace, a ';' comment or the line end.  The name is matched in Python (_canonical_section), the way
# WNTR's reader does it, so the pattern only finds the bracketed word: linear, with no alternation to backtrack.
_INP_HEADER = re.compile(r"(?m)^[^\S\n]*\[([A-Za-z]+)\](?=[^\S\n]|;|$)")
NODE_SECTIONS = ("JUNCTIONS", "RESERVOIRS", "TANKS")
LINK_SECTIONS = ("PIPES", "PUMPS", "VALVES")
NAMED_SECTIONS = ("PATTERNS", "CURVES")   # their first token is a pattern or curve ID, repeated on continuation lines


def _canonical_section(name: str) -> str:
    """A section name as WNTR's .inp reader reads it: upper case, and a name with a missing or an extra S
    ([JUNCTION], [PIPE], [TANK], [RESERVOIR], [OPTION], [TITLES]) is the EPANET section it stands for.  A SWMM or MSX
    name, or any other, comes back as it is (upper case)."""
    n = name.strip().upper()
    if n in INP_SECTIONS:
        return n
    if n + "S" in INP_SECTIONS:
        return n + "S"
    if n.endswith("S") and n[:-1] in INP_SECTIONS:
        return n[:-1]
    return n


_KNOWN_SECTIONS = frozenset(INP_SECTIONS + SWMM_SECTIONS + MSX_SECTIONS)


def _section_headers(text: str) -> set[str]:
    """The section names (upper case, as WNTR reads them) of every header line in the text, whatever its line ends."""
    flat = text.replace("\r\n", "\n").replace("\r", "\n")
    found = {_canonical_section(m.group(1)) for m in _INP_HEADER.finditer(flat)}
    return found & _KNOWN_SECTIONS

DWG_VERSIONS = {"AC1.2": "R1.2", "AC1.40": "R1.40", "AC1.50": "R2.05", "AC2.10": "R2.10", "AC1001": "R2.22",
                "AC1002": "R2.5", "AC1003": "R2.6", "AC1004": "R9", "AC1006": "R10", "AC1009": "R11 and R12",
                "AC1012": "R13", "AC1014": "R14", "AC1015": "AutoCAD 2000", "AC1018": "AutoCAD 2004",
                "AC1021": "AutoCAD 2007", "AC1024": "AutoCAD 2010", "AC1027": "AutoCAD 2013",
                "AC1032": "AutoCAD 2018"}
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
SQLITE_MAGIC = b"SQLite format 3\x00"
GPKG_APP_IDS = {0x47504B47: "GPKG", 0x47503130: "GP10", 0x47503131: "GP11"}
BINARY_DXF = b"AutoCAD Binary DXF\r\n\x1a\x00"
EPANET_BIN_MAGIC = 516114521
SHP_FILE_CODE = 9994

_EPANET_EXPORT = ("If a hydraulic model of this system exists, send its EPANET export instead: in WaterCAD or WaterGEMS, "
                  "make the scenario you want active and use File, Export, EPANET, with IDs rather than labels; in "
                  "InfoWater, use its EPANET export; most modelling packages can write an EPANET .inp.")
_SKELETON_LATER = ("Building a skeleton model from a drawing or a GIS layer is planned, not built yet, and even then it "
                   "carries no demands, tanks, pumps or controls.")

# kind -> (category, plain label, advice).  Categories: model, other_model (a model of something else, or part of
# one), results, drawing, bim, vendor_model, gis, table, document, text, archive, refused, unknown.
KINDS: dict[str, tuple[str, str, str]] = {
    "epanet_inp": ("model", "an EPANET model (.inp)",
                   "It is checked below and then run."),
    "epanet_inp_zip": ("model", "a zip holding an EPANET model",
                       "The model is read from the zip in memory (the zip is not unpacked) and saved as a prepared "
                       "copy."),
    "swmm_inp": ("other_model", "a SWMM sewer or stormwater model (.inp)",
                 "SWMM models sewers and storm drains, not a pressurised drinking-water network, although its files "
                 "also end in .inp. Send the water distribution model's EPANET .inp instead."),
    "epanet_msx": ("other_model", "an EPANET-MSX reaction file (.msx)",
                   "It holds water quality reactions, not the pipe network. Send the network's EPANET .inp: "
                   "ResidualMap sets and calibrates the chlorine decay itself."),
    "epanet_bin": ("results", "an EPANET binary results file (.bin)",
                   "It holds the results of a run, not the model. Send the .inp file that produced it."),
    "dxf_ascii": ("drawing", "a CAD drawing (DXF, text)",
                  "A drawing holds lines, not a hydraulic model. " + _EPANET_EXPORT + " " + _SKELETON_LATER),
    "dxf_binary": ("drawing", "a CAD drawing (binary DXF)",
                   "A drawing holds lines, not a hydraulic model. Save it again as a text (ASCII) DXF. " + _EPANET_EXPORT
                   + " " + _SKELETON_LATER),
    "dwg": ("drawing", "an AutoCAD drawing (DWG)",
            "DWG is a closed format that cannot be read here: save it as DXF (AutoCAD: Save As, DXF) or export the water "
            "mains as a GIS layer. " + _EPANET_EXPORT),
    "dgn_v7": ("drawing", "a MicroStation drawing (DGN V7)",
               "Export it as DXF or as a GIS layer. " + _EPANET_EXPORT),
    "dgn_v8": ("drawing", "a MicroStation drawing (DGN V8)",
               "DGN V8 needs MicroStation itself to convert: export it as DXF (File, Export, DXF) or as a GIS layer. In a "
               "WaterGEMS for MicroStation project the drawing is not the model. " + _EPANET_EXPORT),
    "ifc": ("bim", "a BIM model (IFC)",
            "IFC describes buildings and plant equipment, not the hydraulics of a pipe network. " + _EPANET_EXPORT
            + " Otherwise send the water mains as a GIS layer."),
    "pdf": ("document", "a PDF (a map book, drawing set or report)",
            "A PDF map cannot be read as a network. " + _EPANET_EXPORT + " Otherwise send the water mains as a GIS "
            "layer."),
    "access_db": ("vendor_model", "a Microsoft Access (Jet or ACE) database",
                  "Older WaterCAD and WaterGEMS versions keep the model in an Access database (.wtg.mdb), whose layout is "
                  "not published, so it is not read here. In WaterCAD or WaterGEMS use File, Export, EPANET and send "
                  "the .inp."),
    "sqlite_model": ("vendor_model", "a SQLite database that is not a GeoPackage",
                     "If it came from WaterGEMS or WaterCAD (a .wtg.sqlite file) it is the model database, whose layout "
                     "is not published, so it is not read here: in WaterGEMS or WaterCAD use File, Export, EPANET and "
                     "send the .inp."),
    "wtg_settings": ("vendor_model", "a WaterCAD or WaterGEMS project file (.wtg)",
                     "The .wtg file holds project settings; the model is in the .wtg.sqlite database next to it, and "
                     "neither is read here. In WaterCAD or WaterGEMS use File, Export, EPANET and send the .inp."),
    "vendor_model_zip": ("vendor_model", "a zip of WaterCAD or WaterGEMS model files",
                         "These files are not read here. In WaterCAD or WaterGEMS use File, Export, EPANET and send the "
                         ".inp."),
    "geopackage": ("gis", "a GIS layer (GeoPackage)",
                   "A GIS layer holds pipes and their sizes, not a hydraulic model. " + _EPANET_EXPORT + " "
                   + _SKELETON_LATER),
    "geojson": ("gis", "a GIS layer (GeoJSON)",
                "A GIS layer holds pipes and their sizes, not a hydraulic model. " + _EPANET_EXPORT + " "
                + _SKELETON_LATER),
    "esri_json": ("gis", "a GIS layer (Esri JSON)",
                  "A GIS layer holds pipes and their sizes, not a hydraulic model. " + _EPANET_EXPORT + " "
                  + _SKELETON_LATER),
    "gml": ("gis", "a GIS layer (GML)",
            "A GIS layer holds pipes and their sizes, not a hydraulic model. " + _EPANET_EXPORT + " " + _SKELETON_LATER),
    "kml": ("gis", "a Google Earth file (KML)",
            "Lines drawn in Google Earth are not a hydraulic model. " + _EPANET_EXPORT + " " + _SKELETON_LATER),
    "kmz": ("gis", "a Google Earth file (KMZ)",
            "Lines drawn in Google Earth are not a hydraulic model. " + _EPANET_EXPORT + " " + _SKELETON_LATER),
    "shapefile_zip": ("gis", "a zipped shapefile",
                      "A GIS layer holds pipes and their sizes, not a hydraulic model. " + _EPANET_EXPORT + " "
                      + _SKELETON_LATER),
    "shp": ("gis", "a shapefile's geometry file (.shp)",
            "A .shp alone has no attributes: send it zipped with its .shx, .dbf and .prj. A GIS layer is not a hydraulic "
            "model. " + _EPANET_EXPORT),
    "file_gdb_zip": ("gis", "a zipped Esri file geodatabase (.gdb)",
                     "A file geodatabase cannot be read here: export the water mains as a shapefile or GeoPackage. "
                     + _EPANET_EXPORT),
    "wntr_json": ("model", "a WNTR model file (JSON)",
                  "Write it out as an EPANET .inp (wntr.network.write_inpfile) and send that."),
    "json_other": ("text", "a JSON file that is not a GIS layer or a model",
                   "It is not a model, a drawing or a GIS layer that ResidualMap recognises. Tell us which program made "
                   "it."),
    "xml_other": ("text", "an XML file",
                  "It is not a model, a drawing or a GIS layer that ResidualMap recognises. Tell us which program made "
                  "it."),
    "gpx": ("gis", "a GPS track (GPX)",
            "GPS points are useful for sample taps later; they are not a hydraulic model. " + _EPANET_EXPORT),
    "xlsx": ("table", "an Excel workbook (.xlsx)",
             "If it is a grab-sample log, upload it under '2. Your samples' with the demo off, with its tap map; the "
             "Get started page checks it first. It is not a model."),
    "xls_legacy": ("table", "a legacy Excel workbook (.xls)",
                   "Save it as .xlsx or CSV. If it is a grab-sample log, upload that under '2. Your samples'."),
    "office_other": ("document", "a Word or PowerPoint file",
                     "It is not a model. " + _EPANET_EXPORT),
    "delimited_text": ("table", "a table in text form (CSV or similar)",
                       "If it is a grab-sample log or a tap map, upload it under '2. Your samples' with the demo off; "
                       "the Get started page checks it first. It is not a model."),
    "text_other": ("text", "plain text with no EPANET sections",
                   "It has none of the EPANET section headers ([JUNCTIONS], [PIPES] and so on). If it should be a model, "
                   "open it in EPANET or the modelling package and save it again as .inp."),
    "zip_other": ("archive", "a zip with no model, drawing or GIS layer recognised",
                  "Send the model's EPANET export (.inp) on its own, or tell us what the zip holds."),
    "ole_other": ("unknown", "an old Microsoft Office or MicroStation file (OLE) of unknown type",
                  "Not sure what it is. Tell us which program made it. " + _EPANET_EXPORT),
    "refused_zip": ("refused", "a zip that is refused",
                    "Nothing in it was opened or unpacked. Send the model file on its own."),
    "empty": ("unknown", "an empty file", "The file has no content. Send it again."),
    "unknown": ("unknown", "a file of unknown type",
                "Not sure what this file is: it is not an EPANET model, a drawing, a GIS layer or a table that "
                "ResidualMap recognises. Tell us which program made it. " + _EPANET_EXPORT),
}


class ZipRefused(ValueError):
    """A zip that is not opened: an unsafe member path, or too large when unpacked."""


@dataclass
class Sniff:
    kind: str
    detail: str = ""
    confidence: str = "high"           # high, medium or low; low means 'not sure', and the advice asks the user
    by_extension: bool = False         # the file name decided the kind or added to the detail (the detail says so)
    backup: bool = False               # the name ends in .bak: a backup copy of whatever the content is
    encoding: str | None = None        # text files: how the text was decoded
    members: list[str] = field(default_factory=list)   # zips: the members that decided the kind

    @property
    def category(self) -> str:
        return KINDS[self.kind][0]

    @property
    def label(self) -> str:
        return KINDS[self.kind][1]

    @property
    def advice(self) -> str:
        return KINDS[self.kind][2]

    @property
    def is_model(self) -> bool:
        return self.kind in ("epanet_inp", "epanet_inp_zip")

    def sentence(self) -> str:
        """One plain sentence: what the file is, by its content."""
        what = f"a backup copy of {self.label}" if self.backup else self.label
        s = f"This file is {what}"
        if self.detail:
            s += f": {self.detail}"
        s += "."
        if self.confidence == "low":
            s = "Not sure. " + s
        return s

    def as_dict(self) -> dict:
        d = asdict(self)
        d.update(category=self.category, label=self.label)
        return d


# ----------------------------------------------------------------------------- text
def decode_text(data: bytes) -> tuple[str, str]:
    """(text, how): a UTF-8 byte-order mark (stripped), a UTF-16 one, UTF-16 without one (NUL bytes on every other
    position), strict UTF-8, UTF-8 with a few stray bytes of a Windows code page (task 15b: _mixed_utf8, when the
    file's valid UTF-8 characters are not the byte pairs a cp1252 file makes by chance), then cp1252, where any of the
    five bytes cp1252 leaves undefined is read as latin-1 and every other byte as cp1252 (so a euro sign or a curly
    quote elsewhere in the file still reads right).
    how is 'utf-8' only for plain UTF-8 with no BOM, the one case that needs no rewrite."""
    for bom, codec, how in ((b"\xef\xbb\xbf", "utf-8", "utf-8 with a byte-order mark"),
                            (b"\xff\xfe", "utf-16-le", "utf-16 (little-endian, with a byte-order mark)"),
                            (b"\xfe\xff", "utf-16-be", "utf-16 (big-endian, with a byte-order mark)")):
        if data.startswith(bom):
            try:
                return data[len(bom):].decode(codec), how
            except UnicodeDecodeError:
                return data[len(bom):].decode(codec, errors="replace"), how + ", with undecodable bytes replaced"
    guess = _utf16_without_bom(data[:4096])
    if guess:
        return data.decode(guess, errors="replace"), f"utf-16 ({guess[-2:]}, no byte-order mark)"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    mixed = _mixed_utf8(data)
    if mixed is not None:
        text, n = mixed[:2]
        return text, f"{MIXED_PREFIX} {n} stray byte{'s' if n != 1 else ''} read as cp1252 (Windows Western European)"
    try:
        return data.decode("cp1252"), "cp1252 (Windows Western European)"
    except UnicodeDecodeError:
        n = sum(data.count(bytes([b])) for b in CP1252_UNDEFINED)
        return (data.decode("cp1252", errors="residualmap_cp1252_gaps"),
                f"cp1252 (Windows Western European), with {n} byte{'s' if n != 1 else ''} it leaves undefined read as "
                f"latin-1")


MIXED_PREFIX = "utf-8 mixed with"       # decode_text's word for UTF-8 with a few bytes of a Windows code page
_STRAY = re.compile("[\udc80-\udcff]")   # a byte that is not UTF-8, as surrogateescape decodes it
_NON_ASCII = re.compile("[^\x00-\x7f]")
_LINE_BREAK = re.compile("\r\n|\r|\n")


def _cp1252_byte(m: re.Match) -> str:
    return bytes([ord(m.group(0)) - 0xDC00]).decode("cp1252", errors="residualmap_cp1252_gaps")


def _utf8_kind(t: str, i: int) -> str:
    """How the valid multi-byte UTF-8 character t[i] reads in a file that also holds bytes that are not UTF-8 (task
    15b's second review).  A file in cp1252 throughout holds such characters only by chance, as an accented capital
    (bytes C2 to DF) followed by a cp1252 symbol or punctuation byte (80 to BF), and the character those two bytes make
    in UTF-8 says which it is:
    'strong': a lowercase accented letter (U+00E0 to U+00FF, but not U+00F0), a curly quote, a dash or a bullet
      (U+2010 to U+203A) or the euro sign.  In cp1252 these bytes read as 'Ã' or 'â€' and a symbol ('Ã©', 'â€™'),
      which no text in cp1252 writes, so one of them shows the file holds UTF-8.
    'plausible': another character a model's text uses (U+00A0 to U+017F, the Latin-1 and Latin Extended-A letters,
      the degree sign and the no-break space; the Romanian letters with a comma below; the Greek letters; any other
      character of three or four bytes), or a letter of Cyrillic to NKo (U+0400 to U+07FF) next to another letter of
      those scripts or Greek, as in a word.
    'chance': anything else, the byte pairs a cp1252 file makes by chance: the C1 controls (U+0080 to U+009F, 'Â' and
      a quote or dash), Latin Extended-B, the IPA and the combining marks (U+0180 to U+037F, so 'É', 'È', 'Ç', 'Ì' or
      'Í' and a quote, dash, degree sign or no-break space: 'CITÉ’' is the bytes C9 92, U+0252), and a Cyrillic to NKo
      letter alone between Latin letters ('Ó' or 'Ñ' and a dash)."""
    o = ord(t[i])
    if (0xE0 <= o <= 0xFF and o != 0xF0) or 0x2010 <= o <= 0x203A or o == 0x20AC:
        return "strong"
    if 0xA0 <= o <= 0x17F or 0x218 <= o <= 0x21B or 0x386 <= o <= 0x3CE or o >= 0x800:
        return "plausible"
    if 0x400 <= o <= 0x7FF and any(0 <= k < len(t) and 0x370 <= ord(t[k]) <= 0x7FF for k in (i - 1, i + 1)):
        return "plausible"
    return "chance"


def _mixed_utf8(data: bytes) -> tuple[str, int, list[int], int, int] | None:
    """A UTF-8 file with a few stray bytes of a Windows code page (task 15b; one pasted degree sign or accented letter
    is enough to fail strict UTF-8): (text, the number of bytes read as cp1252, the first three lines that hold them,
    the number of lines that hold them, the number of chance pairs among them), or None to decode the file as cp1252
    throughout, as task 15 did.
    Each valid UTF-8 character is weighed with _utf8_kind.  The file is read as UTF-8 when it holds a 'strong' one, or
    when the lines that hold 'plausible' ones are at least as many as the lines that hold stray bytes or chance pairs.
    Then every stray byte and every chance pair is read as cp1252 (any of the five bytes cp1252 leaves undefined as
    latin-1) and every other UTF-8 character is kept.  Decoding such a file as cp1252 throughout, as task 15 did,
    turned every UTF-8 accent into two characters: an accented ID changed, a letter whose second byte is 0xA0 became a
    no-break space that WNTR splits on, and the model was blocked with WNTR's or EPANET's error rather than a word
    about the encoding.
    Lines, not occurrences, are counted (the second review): an ID is written on its own line, at each pipe end and
    in the coordinates, so counting occurrences let one ID with a chance pair outweigh the stray bytes of a whole file
    ('CITÉ’15' with a degree sign in the title: 3 against 1).  In the two encodings a stray byte is read the same way
    and a chance pair too; only the 'plausible' and 'strong' characters read differently, so the rule weighs only the
    evidence about them.  What it still misreads: a file in cp1252 throughout whose only chance pairs make plausible
    letters (an 'Ã', 'Ä' or 'Å' before a quote or dash: 'KÅ“' reads as 'Kœ') on at least as many lines as its other
    bytes that are not UTF-8; and a UTF-8 file whose only accents are capitals or other 'plausible' characters, with
    stray bytes on more lines than they are, reads as cp1252 throughout."""
    t = data.decode("utf-8", errors="surrogateescape")
    if not _STRAY.search(t):
        return None
    import bisect
    breaks = [m.start() for m in _LINE_BREAK.finditer(t)]
    utf8_lines, other_lines, strong, chance, n_bytes = set(), set(), False, set(), 0
    for m in _NON_ASCII.finditer(t):
        i = m.start()
        line = bisect.bisect_left(breaks, i) + 1
        if "\udc80" <= t[i] <= "\udcff":
            other_lines.add(line)
            n_bytes += 1
            continue
        kind = _utf8_kind(t, i)
        if kind == "chance":
            chance.add(i)
            other_lines.add(line)
            n_bytes += len(t[i].encode("utf-8"))
        else:
            utf8_lines.add(line)
            strong = strong or kind == "strong"
    if not utf8_lines or not (strong or len(utf8_lines) >= len(other_lines)):
        return None

    def cp(m: re.Match) -> str:
        if m.start() in chance:
            return m.group(0).encode("utf-8").decode("cp1252", errors="residualmap_cp1252_gaps")
        return _cp1252_byte(m) if "\udc80" <= m.group(0) <= "\udcff" else m.group(0)
    return _NON_ASCII.sub(cp, t), n_bytes, sorted(other_lines)[:3], len(other_lines), len(chance)


def _utf16_without_bom(head: bytes) -> str | None:
    """'utf-16-le' or 'utf-16-be' when a text file without a byte-order mark has NUL bytes on every other position."""
    h = head[:4096]
    if len(h) < 16 or h.count(0) < len(h) // 4:
        return None
    odd, even = h[1::2].count(0), h[0::2].count(0)
    if odd > 0.9 * (len(h) // 2) and even < 0.05 * len(h):
        return "utf-16-le"
    if even > 0.9 * (len(h) // 2) and odd < 0.05 * len(h):
        return "utf-16-be"
    return None


def _head_text(data: bytes) -> tuple[str, str]:
    """The first HEAD_BYTES as text, with how they were decoded (a multi-byte character cut by the 64 KB limit does not
    turn UTF-8 into cp1252)."""
    head = data[:HEAD_BYTES + 4]
    guess = None if head.startswith((b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff")) else _utf16_without_bom(head)
    if guess:
        return head[:HEAD_BYTES].decode(guess, errors="replace"), f"utf-16 ({guess[-2:]}, no byte-order mark)"
    if len(data) <= HEAD_BYTES:
        return decode_text(data)
    text, how = decode_text(head[:HEAD_BYTES])
    if how.startswith(("cp1252", MIXED_PREFIX)):
        try:
            head[:HEAD_BYTES].decode("utf-8")
        except UnicodeDecodeError as e:
            if e.start >= HEAD_BYTES - 3:           # only the cut at the end fails
                return head[:e.start].decode("utf-8"), "utf-8"
    return text, how


def _is_binary(head: bytes) -> bool:
    h = head[:8192]
    if not h:
        return False
    ctrl = sum(1 for b in h if b < 9 or (13 < b < 32 and b != 27))
    return ctrl > 0.01 * len(h)


# ----------------------------------------------------------------------------- zips, in memory
def _unsafe_member(name: str) -> bool:
    n = name.replace("\\", "/")
    return n.startswith("/") or bool(re.match(r"^[A-Za-z]:", n)) or any(p == ".." for p in n.split("/"))


def zip_junk(name: str) -> bool:
    """A file that macOS or Windows adds to a zip by itself: anything under __MACOSX/ (Finder's Compress), an
    AppleDouble '._' copy, .DS_Store, Thumbs.db or desktop.ini.  It is left out of what a zip is said to hold."""
    parts = [p for p in name.replace("\\", "/").split("/") if p]
    base = parts[-1] if parts else ""
    return "__MACOSX" in parts or base.startswith("._") or base.lower() in (".ds_store", "thumbs.db", "desktop.ini")


def zip_listing(data: bytes, max_total: int = ZIP_MAX_TOTAL) -> list[zipfile.ZipInfo]:
    """The members of a zip held in memory, after the safety rules (no unsafe path, at most ZIP_MAX_MEMBERS members,
    at most max_total bytes unpacked in all); raises ZipRefused.  Nothing is unpacked, and the compression ratio is
    tested only for the member that is read (read_zip_member)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except Exception as e:  # noqa: BLE001  (BadZipFile, ValueError, struct.error, ...: any of them means not readable)
        raise ZipRefused(f"it is not a readable zip ({type(e).__name__}: {e})") from None
    if len(infos) > ZIP_MAX_MEMBERS:
        raise ZipRefused(f"it has {len(infos)} members, more than {ZIP_MAX_MEMBERS}")
    bad = [i.filename for i in infos if _unsafe_member(i.filename)]
    if bad:
        raise ZipRefused(f"a member's path leaves the zip's own folder ({bad[0]})")
    total = sum(i.file_size for i in infos)
    if total > max_total:
        raise ZipRefused(f"its members add up to {total / 1e6:.0f} MB unpacked, more than {max_total / 1e6:.0f} MB")
    return infos


def read_zip_member(data: bytes, member: str, limit: int = ZIP_MAX_TOTAL, ratio_max: float = ZIP_RATIO_MAX,
                    ratio_min_size: int = ZIP_RATIO_MIN_SIZE) -> bytes:
    """One member of a zip held in memory, read with a cap: a member over ratio_min_size bytes unpacked that is packed
    more than ratio_max to 1 is refused before it is opened, and a member whose real size exceeds its declared size
    (or the cap) is refused, so a zip that lies in its directory cannot fill memory."""
    infos = {i.filename: i for i in zip_listing(data)}
    if member not in infos:
        raise ZipRefused(f"no member {member}")
    i = infos[member]
    if i.file_size > ratio_min_size and i.file_size > ratio_max * max(i.compress_size, 1):
        raise ZipRefused(f"member {member} would unpack to {i.file_size / 1e6:.0f} MB, "
                         f"{i.file_size / max(i.compress_size, 1):.0f} times its packed size (more than "
                         f"{ratio_max:g} to 1)")
    cap = min(i.file_size, limit)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf, zf.open(member) as fh:
            out = fh.read(cap + 1)
    except Exception as e:  # noqa: BLE001  (a corrupt offset raises ValueError 'negative seek value'; any failure refuses)
        raise ZipRefused(f"member {member} cannot be read ({type(e).__name__}: {e})") from None
    if len(out) > cap:
        raise ZipRefused(f"member {member} is larger than its zip directory says")
    return out


def _sniff_zip(data: bytes) -> Sniff:
    try:
        infos = zip_listing(data)
    except ZipRefused as e:
        return Sniff("refused_zip", detail=str(e), confidence="high")
    files = [i.filename for i in infos if not i.is_dir()]
    names = [n for n in files if not zip_junk(n)]
    low = [n.lower() for n in names]
    s = _sniff_zip_names(names, low)
    if s.kind == "zip_other" and names:
        models = _zip_models_by_content(data, names)
        if models:
            s = Sniff("epanet_inp_zip", detail="it holds " + ", ".join(models[:5]) + (" and more" if len(models) > 5
                                                                                     else "")
                      + f", an EPANET model by {'its' if len(models) == 1 else 'their'} content (the name does not "
                        f"end in .inp)", members=models)
    junk = len(files) - len(names)
    if junk:
        s.detail += (f"{'; ' if s.detail else ''}{junk} file{'s' if junk > 1 else ''} that macOS or Windows adds "
                     f"(__MACOSX, ._ copies, .DS_Store, Thumbs.db) left out")
    return s


ZIP_CONTENT_SNIFF_MAX = 50             # members whose content is read when no member's name says what the zip holds


def _zip_models_by_content(data: bytes, names: list[str]) -> list[str]:
    """Members that are EPANET models by their content, read in memory (task 15b: a zip holding a model renamed .txt
    was typed as 'a zip with nothing recognised').  Only when no member's name decided the kind; at most
    ZIP_CONTENT_SNIFF_MAX members, each under the same rules as the member that is read (read_zip_member: the
    ratio rule over 10 MB, the size cap), none over TEXT_SCAN_MAX unpacked, and a member that is itself a zip is not
    opened."""
    sizes = {}
    try:
        sizes = {i.filename: i.file_size for i in zip_listing(data)}
    except ZipRefused:
        return []
    out = []
    for n in names[:ZIP_CONTENT_SNIFF_MAX]:
        if sizes.get(n, TEXT_SCAN_MAX + 1) > TEXT_SCAN_MAX:
            continue
        try:
            b = read_zip_member(data, n)
        except ZipRefused:
            continue
        if b.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
            continue
        try:
            if _sniff_content(b, n).kind == "epanet_inp":
                out.append(n)
        except Exception:  # noqa: BLE001  (a member that cannot be read decides nothing)
            continue
    return out


def _sniff_zip_names(names: list[str], low: list[str]) -> Sniff:
    if not names:
        return Sniff("zip_other", detail="no files", confidence="medium")
    if "[content_types].xml" in low:
        if any(n.startswith("xl/") for n in low):
            return Sniff("xlsx", detail=f"{len(names)} parts, typed by its [Content_Types].xml and xl/ folder")
        return Sniff("office_other", detail="typed by its [Content_Types].xml")
    inp = [n for n, l in zip(names, low) if l.endswith(".inp") or l.endswith(".inp.bak")]
    if inp:
        return Sniff("epanet_inp_zip", detail="it holds " + ", ".join(inp[:5]) + (" and more" if len(inp) > 5 else ""),
                     members=inp)
    kml = [n for n, l in zip(names, low) if l.endswith(".kml")]
    if kml:
        return Sniff("kmz", detail="it holds " + ", ".join(kml[:3]), members=kml)
    gdb = sorted({re.sub(r"(?i)(\.gdb)/.*$", r"\1", n) for n in names if re.search(r"(?i)\.gdb/", n)})
    if gdb and any(l.endswith((".gdbtable", ".gdbtablx")) for l in low):
        return Sniff("file_gdb_zip", detail="geodatabase " + ", ".join(gdb), members=gdb)
    shp = [n for n, l in zip(names, low) if l.endswith(".shp")]
    if shp:
        layers = []
        for n in shp:
            base = n[:-4]
            have = {ext for ext in (".shx", ".dbf", ".prj", ".cpg") if (base + ext).lower() in low}
            miss = [ext for ext in (".shx", ".dbf", ".prj") if ext not in have]
            layers.append(base + (f" (missing {', '.join(miss)})" if miss else ""))
        return Sniff("shapefile_zip", detail=f"{len(shp)} layer{'s' if len(shp) > 1 else ''}: " + "; ".join(layers),
                     members=shp)
    wtg = [n for n, l in zip(names, low) if re.search(r"\.wtg(\.sqlite|\.mdb)?(\.bak)?$", l)]
    if wtg:
        return Sniff("vendor_model_zip", detail="it holds " + ", ".join(wtg[:4]), members=wtg)
    exts = sorted({os.path.splitext(l)[1] or "(no extension)" for l in low})
    return Sniff("zip_other", detail=f"{len(names)} files ({', '.join(exts[:8])})", confidence="medium")


# ----------------------------------------------------------------------------- OLE, SQLite
def ole_names(data: bytes, max_dir_sectors: int = 64) -> list[str]:
    """Names of the storages and streams in an OLE compound file's directory (the first max_dir_sectors directory
    sectors, followed through the first FAT sectors held in data).  [] when the directory is not in data."""
    if len(data) < 512 or not data.startswith(OLE_MAGIC):
        return []
    ssz = 1 << struct.unpack_from("<H", data, 30)[0]
    if ssz not in (512, 4096):
        return []
    first_dir = struct.unpack_from("<I", data, 48)[0]
    fat_secs = [s for s in struct.unpack_from("<109I", data, 76) if s < 0xFFFFFFFA]

    def sector(i):
        off = (i + 1) * ssz
        return data[off:off + ssz] if off + ssz <= len(data) else None

    fat: list[int] = []
    for s in fat_secs:
        b = sector(s)
        if b is None:
            break
        fat.extend(struct.unpack(f"<{ssz // 4}I", b))
    names, seen, s = [], set(), first_dir
    while s < 0xFFFFFFFA and s not in seen and len(seen) < max_dir_sectors:
        seen.add(s)
        b = sector(s)
        if b is None:
            break
        for k in range(0, ssz, 128):
            e = b[k:k + 128]
            nlen, typ = struct.unpack_from("<H", e, 64)[0], e[66]
            if typ in (1, 2, 5) and 2 <= nlen <= 64:
                names.append(e[:nlen - 2].decode("utf-16-le", errors="replace"))
        s = fat[s] if s < len(fat) else 0xFFFFFFFE
    return names


def _sqlite_open(data: bytes) -> sqlite3.Connection:
    """An in-memory copy of an uploaded SQLite database (sqlite3 deserialize; nothing written), with a step budget:
    any query that runs past SQLITE_MAX_STEPS progress calls is stopped (sqlite3.OperationalError 'interrupted'), so
    a crafted file cannot keep a query busy."""
    con = sqlite3.connect(":memory:")
    calls = [0]

    def budget() -> int:
        calls[0] += 1
        return 1 if calls[0] > SQLITE_MAX_STEPS else 0

    try:
        con.deserialize(data)
        con.set_progress_handler(budget, 1000)
    except BaseException:
        con.close()
        raise
    return con


def _sqlite_tables(data: bytes) -> dict[str, tuple[str, str]] | None:
    """{name: (type, sql)} for the tables and views of a SQLite database held in memory, read from sqlite_master only
    (no view is run), or None when it cannot be read."""
    try:
        con = _sqlite_open(data)
    except (sqlite3.Error, AttributeError, OverflowError, ValueError, MemoryError):
        return None
    with contextlib.closing(con):
        try:
            rows = con.execute("select name, type, sql from sqlite_master where type in ('table', 'view') "
                               "limit 10000").fetchall()
        except (sqlite3.Error, OverflowError, ValueError):
            return None
        return {str(n): (str(t), str(q or "")) for n, t, q in rows}


def _gpkg_layers(data: bytes, tables: dict) -> str:
    """'layers: ...' from gpkg_contents, read only when it is a plain table (a view or a virtual table would run code
    from the file), at most 50 rows, under the step budget."""
    typ, sql = tables.get("gpkg_contents", ("", ""))
    if not typ:
        return "layers not listed (no gpkg_contents table)"
    if typ != "table" or sql.lstrip().upper().startswith("CREATE VIRTUAL"):
        what = "a view" if typ == "view" else "a virtual table"
        return f"layers not listed (its gpkg_contents is {what}, not a plain table, so it is not read)"
    try:
        con = _sqlite_open(data)
        with contextlib.closing(con):
            rows = con.execute("select table_name, data_type from gpkg_contents limit 50").fetchall()
    except (sqlite3.Error, AttributeError, OverflowError, ValueError, MemoryError):
        return "layers not readable"
    layers = [f"{t} ({d})" for t, d in rows]
    return f"layers: {', '.join(layers) if layers else 'none listed'}"


# ----------------------------------------------------------------------------- JSON and XML
_GEOJSON_TYPES = {"FeatureCollection", "Feature", "Point", "MultiPoint", "LineString", "MultiLineString", "Polygon",
                  "MultiPolygon", "GeometryCollection"}


def _coords(geom, out: list, cap: int = 2000, max_depth: int = JSON_MAX_DEPTH) -> None:
    """Up to cap (x, y) pairs from a GeoJSON geometry, depth first and left to right, with an explicit stack and a
    nesting cap (so a crafted file cannot exhaust the recursion limit); members of the wrong type are skipped."""
    stack = [(geom, 0)]
    while stack and len(out) < cap:
        g, depth = stack.pop()
        if g is None or depth > max_depth:
            continue
        if isinstance(g, dict):
            subs = g.get("geometries")
            if isinstance(subs, list):
                stack.extend((x, depth + 1) for x in reversed(subs[:cap]))
            if "coordinates" in g:
                stack.append((g["coordinates"], depth + 1))
        elif isinstance(g, (list, tuple)):
            if len(g) >= 2 and all(isinstance(v, (int, float)) for v in g[:2]):
                try:
                    out.append((float(g[0]), float(g[1])))
                except OverflowError:
                    continue
            else:
                stack.extend((x, depth + 1) for x in reversed(g[:cap]))


def _crs_name(d: dict) -> str:
    """The GeoJSON crs member's properties.name, when every part of it has the type the specification gives."""
    crs = d.get("crs")
    props = crs.get("properties") if isinstance(crs, dict) else None
    name = props.get("name") if isinstance(props, dict) else None
    return name if isinstance(name, str) else ""


def _geojson_detail(d: dict) -> str:
    feats = d.get("features") if d.get("type") == "FeatureCollection" else None
    bad_feats = feats is not None and not isinstance(feats, list)
    if bad_feats:
        feats = []
    pts: list = []
    for f in (feats or [d])[:500]:
        _coords(f.get("geometry", f) if isinstance(f, dict) else None, pts)
    crs = _crs_name(d)
    m = re.search(r"EPSG:+(\d+)", crs)
    epsg = m.group(1) if m else ("4326" if "CRS84" in crs else None)
    lonlat = bool(pts) and all(abs(x) <= 180 and abs(y) <= 90 for x, y in pts)
    n = f"{len(feats)} feature{'s' if len(feats) != 1 else ''}, " if feats is not None else ""
    if bad_feats:
        n = "a features member that is not a list, "
    if epsg and epsg != "4326":
        where = f"projected coordinates (EPSG:{epsg}, from its crs member)"
        if lonlat:
            where += ", although the numbers fall in the lon/lat range"
    elif lonlat:
        where = "lon/lat coordinates (WGS84)"
    elif pts:
        where = "coordinates outside the lon/lat range, and no crs member says which system they are in"
    else:
        where = "no coordinates"
    return n + where


def _sniff_json(text: str, complete: bool) -> Sniff | None:
    try:
        d = json.loads(text) if complete else None
    except RecursionError:
        return Sniff("json_other", detail="nested too deeply to read", confidence="low")
    except ValueError:
        d = None
    if d is None:
        if complete:
            return None
        if re.search(r'"type"\s*:\s*"FeatureCollection"', text):
            return Sniff("geojson", detail="too large to read whole here; typed by its first 64 KB", confidence="medium")
        if re.search(r'"geometryType"\s*:\s*"esriGeometry', text):
            return Sniff("esri_json", detail="typed by its first 64 KB", confidence="medium")
        return None
    if isinstance(d, dict):
        keys = set(d)
        if str(d.get("version", "")).lower().startswith("wntr") or {"nodes", "links", "options"} <= keys:
            return Sniff("wntr_json", detail=f"WNTR version {d.get('version', 'not given')}")
        if str(d.get("geometryType", "")).startswith("esriGeometry") or (
                "spatialReference" in keys and "features" in keys):
            sr = d.get("spatialReference")
            if isinstance(sr, dict):
                wkid = sr.get("latestWkid") or sr.get("wkid")
            else:                      # some exports write the code alone
                wkid = sr if isinstance(sr, (int, str)) and not isinstance(sr, bool) else None
            feats = d.get("features")
            n = len(feats) if isinstance(feats, list) else 0
            return Sniff("esri_json", detail=f"{d.get('geometryType', 'features')}, {n} feature{'s' if n != 1 else ''}"
                                             + (f", spatial reference {wkid}" if wkid else ""))
        if d.get("type") in _GEOJSON_TYPES:
            return Sniff("geojson", detail=_geojson_detail(d))
        return Sniff("json_other", detail="keys " + ", ".join(sorted(map(str, keys))[:6]), confidence="medium")
    return Sniff("json_other", detail="a JSON list", confidence="medium")


_XML_NAME = re.compile(r"<([A-Za-z_][\w.\-]*)(?::([A-Za-z_][\w.\-]*))?")
_XML_SPACE = re.compile(r"\s*")


def _xml_root(text: str) -> tuple[str | None, str] | None:
    """(prefix, local name) of an XML text's root element, after any leading processing instructions, comments and
    DOCTYPE (its internal subset in square brackets included), or None.  A linear scan with str.find, at most
    XML_MAX_SKIP of them: a pattern with a repeated group of lazy matches (the first version) backtracks
    exponentially on a run of comments with no element after it."""
    i, n = 0, len(text)
    for _ in range(XML_MAX_SKIP + 1):
        i = _XML_SPACE.match(text, i).end()
        if text.startswith("<?", i):
            j = text.find("?>", i + 2)
            if j < 0:
                return None
            i = j + 2
        elif text.startswith("<!--", i):
            j = text.find("-->", i + 4)
            if j < 0:
                return None
            i = j + 3
        elif text[i:i + 9].upper() == "<!DOCTYPE":
            j, depth = i + 9, 0
            while j < n:
                c = text[j]
                if text.startswith("<!--", j):      # a comment in the internal subset
                    k = text.find("-->", j + 4)
                    if k < 0:
                        return None
                    j = k + 2
                elif c in "\"'":                    # a quoted literal may hold '>' or brackets
                    k = text.find(c, j + 1)
                    if k < 0:
                        return None
                    j = k
                elif c == "[":
                    depth += 1
                elif c == "]":
                    depth -= 1
                elif c == ">" and depth <= 0:
                    break
                j += 1
            if j >= n:
                return None
            i = j + 1
        else:
            m = _XML_NAME.match(text, i)
            if not m:
                return None
            return (m.group(1), m.group(2)) if m.group(2) else (None, m.group(1))
    return None


def _sniff_xml(text: str) -> Sniff | None:
    root = _xml_root(text)
    if root is None:
        return None
    prefix, local = root
    tag = f"{prefix}:{local}" if prefix else local
    low = local.lower()
    if low == "kml":
        return Sniff("kml", detail=f"{len(re.findall(r'<Placemark', text))} placemarks, "
                                   f"{len(re.findall(r'<LineString', text))} lines in the first 64 KB")
    if low == "haestaddatasource":
        v = re.search(r'Version="([^"]+)"', text)
        return Sniff("wtg_settings", detail="root element HaestadDataSource" + (f", version {v.group(1)}" if v else ""))
    if low == "gpx":
        return Sniff("gpx", detail="root element gpx")
    if (prefix or "").lower() in ("gml", "wfs", "ogr") or "www.opengis.net/gml" in text[:4096]:
        return Sniff("gml", detail=f"root element {tag}")
    return Sniff("xml_other", detail=f"root element {tag}", confidence="medium")


# ----------------------------------------------------------------------------- delimited text and DXF
_DELIMS = {",": "commas", ";": "semicolons", "\t": "tabs", "|": "vertical bars"}


def _sniffer_sample(sample: str) -> str:
    """What csv.Sniffer is given: the first CSV_SNIFF_CHARS characters with every quote character taken out.
    Sniffer's quote regexes (a lazy match between two quotes, across lines) are super-linear in the number of quote
    characters: 90 KB of a,"a,"... took 27.7 s, and even inside a 4 KB sample a line of a few hundred quote marks
    took seconds (the final verification), the time growing about as the fourth power of the count.  With no quote
    characters those regexes fail at once and Sniffer falls back to counting each delimiter per line, which is
    linear.  Quoting still matters, but only to the consistency test, which reads the lines with csv.reader and
    decides."""
    return sample[:CSV_SNIFF_CHARS].replace('"', "").replace("'", "")


def _sniff_table(text: str) -> Sniff | None:
    lines = [ln for ln in text.splitlines() if ln.strip()][:25]
    if len(lines) < 2:
        return None
    sample = "\n".join(lines)
    delim = None
    try:            # csv.Sniffer only orders the delimiters tried below; see _sniffer_sample for what it is given
        delim = csv.Sniffer().sniff(_sniffer_sample(sample), delimiters="".join(_DELIMS)).delimiter
    except csv.Error:   # the consistency test below reads all 25 lines with csv.reader's quoting, in linear time
        pass
    best = None
    for d in ([delim] if delim else []) + [x for x in _DELIMS if x != delim]:
        try:
            rows = list(csv.reader(lines, delimiter=d))
        except csv.Error:
            continue
        n0 = len(rows[0])
        if n0 < 2:
            continue
        same = sum(len(r) == n0 for r in rows) / len(rows)
        if same >= 0.8:
            best = (d, rows)
            break
    if best is None:
        return None
    d, rows = best
    cols = [c.strip() for c in rows[0]][:8]
    comma_decimals = d != "," and any(re.fullmatch(r"-?\d+,\d+", c.strip()) for r in rows[1:] for c in r)
    detail = f"separated by {_DELIMS[d]}; first row: {', '.join(cols)}" + ("; decimal commas" if comma_decimals else "")
    return Sniff("delimited_text", detail=detail, confidence="medium")


_DXF_SECTIONS = ("HEADER", "CLASSES", "TABLES", "BLOCKS", "ENTITIES", "OBJECTS", "THUMBNAILIMAGE")


def _is_ascii_dxf(text: str) -> tuple[bool, int]:
    """(is a DXF, number of 999 comment groups skipped): group code and value lines, comments first allowed, the first
    real group a 0 SECTION followed by a 2 with a section name."""
    lines = text.splitlines()[:400]
    i, comments = 0, 0
    while i + 1 < len(lines) and lines[i].strip() == "999":
        i, comments = i + 2, comments + 1
    if i + 3 >= len(lines):
        return False, comments
    ok = (lines[i].strip() == "0" and lines[i + 1].strip().upper() == "SECTION" and lines[i + 2].strip() == "2"
          and lines[i + 3].strip().upper() in _DXF_SECTIONS)
    return ok, comments


# ----------------------------------------------------------------------------- the sniff
def sniff(data: bytes, name: str = "") -> Sniff:
    """What a file is, by its content.  The name adds a note or breaks a tie (and marks a .bak as a backup copy); when
    it was used the result says so (by_extension, and the detail).  Never raises: a file whose reading fails in any
    test comes back as 'unknown', with advice, so the uploader always has something plain to say."""
    try:
        s = _sniff_content(data, name)
    except Exception as e:  # noqa: BLE001  (a crafted or broken file must not reach the user as a traceback)
        s = Sniff("unknown", detail=f"it could not be read ({type(e).__name__})", confidence="low")
    low = os.path.basename(name or "").lower()
    if low.endswith(".bak"):
        s.backup = True
    return s


def _only_comments(text: str) -> bool:
    """Every non-blank line starts with ';' (a vendor's comment header, which can run past 64 KB)."""
    lines = [ln.strip(" \t﻿") for ln in text.replace("\r", "\n").split("\n")]
    lines = [ln for ln in lines if ln]
    return bool(lines) and all(ln.startswith(";") for ln in lines)


def _by_name(s: Sniff, note: str) -> Sniff:
    s.by_extension = True
    s.detail = (s.detail + "; " if s.detail else "") + note + " (from the file name)"
    return s


def _sniff_content(data: bytes, name: str) -> Sniff:
    low = os.path.basename(name or "").lower()
    while low.endswith(".bak"):
        low = low[:-4]
    ext = os.path.splitext(low)[1]
    head = data[:HEAD_BYTES]
    if not data:
        return Sniff("empty")
    # binary signatures, strongest first
    if head.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        return _sniff_zip(data)
    if head.startswith(OLE_MAGIC):
        names = ole_names(data)
        if any(n.startswith("Dgn~") for n in names):
            return Sniff("dgn_v8", detail="an OLE compound file with streams named Dgn~")
        if any(n in ("Workbook", "Book") for n in names):
            which = "Workbook" if "Workbook" in names else "Book"
            return Sniff("xls_legacy", detail=f"an OLE compound file with a {which} stream")
        if any(n in ("WordDocument", "PowerPoint Document") for n in names):
            return Sniff("office_other", detail="an OLE compound file with a " +
                         ("WordDocument" if "WordDocument" in names else "PowerPoint Document") + " stream")
        s = Sniff("ole_other", detail="an OLE compound file" + (f" with streams {', '.join(names[:5])}" if names else
                                                                 " whose directory could not be read"),
                  confidence="low")
        if ext == ".xls":
            return _by_name(Sniff("xls_legacy", detail=s.detail, confidence="low"), "the name ends in .xls")
        if ext == ".dgn":
            return _by_name(Sniff("dgn_v8", detail=s.detail, confidence="low"), "the name ends in .dgn")
        return s
    if head.startswith(SQLITE_MAGIC):
        app = struct.unpack_from(">I", head, 68)[0] if len(head) >= 72 else 0
        tables = _sqlite_tables(data)
        if app in GPKG_APP_IDS or (tables and "gpkg_contents" in tables):
            layers = _gpkg_layers(data, tables or {})
            why = (f"application id {GPKG_APP_IDS[app]}" if app in GPKG_APP_IDS else
                   f"a gpkg_contents {(tables or {}).get('gpkg_contents', ('table',))[0]}")
            return Sniff("geopackage", detail=f"{why}; {layers}")
        s = Sniff("sqlite_model", detail=f"{len(tables)} table{'s' if len(tables) != 1 else ''}" if tables is not None
                   else "tables not readable")
        if re.search(r"\.wtg\.sqlite$", low):
            return _by_name(s, "the name ends in .wtg.sqlite, which WaterGEMS and WaterCAD use for the model database")
        return s
    if head[4:19] in (b"Standard Jet DB", b"Standard ACE DB"):
        s = Sniff("access_db", detail=head[4:19].decode())
        if re.search(r"\.wtg\.mdb$", low):
            return _by_name(s, "the name ends in .wtg.mdb, which older WaterCAD and WaterGEMS versions use for the model")
        if low.endswith(".mdb") and "map" in low:
            return _by_name(s, "the name looks like a GIS personal geodatabase")
        return s
    pdf_at = head[:1024].find(b"%PDF-")
    if pdf_at >= 0:
        v = head[pdf_at + 5:pdf_at + 8].decode("latin-1", errors="replace")
        return Sniff("pdf", detail=f"PDF version {v}")
    if re.match(rb"AC(?:1\d{3}|[12]\.\d{1,2})", head[:6]):
        code = head[:6].decode("latin-1").rstrip("\x00")
        ver = DWG_VERSIONS.get(code) or DWG_VERSIONS.get(code[:5]) or "version not listed"
        return Sniff("dwg", detail=f"{ver} format ({code})")
    if head.startswith(BINARY_DXF):
        return Sniff("dxf_binary", detail="it starts with the binary DXF sentinel")
    if len(head) >= 4 and head[0] in (0x08, 0xC8) and head[1:4] == b"\x09\xfe\x02":
        return Sniff("dgn_v7", detail="3D" if head[0] == 0xC8 else "2D")
    if len(head) >= 8 and struct.unpack_from("<i", head, 0)[0] == EPANET_BIN_MAGIC:
        ver, nn, nt, nl = struct.unpack_from("<4i", head, 4) if len(head) >= 20 else (0, 0, 0, 0)
        return Sniff("epanet_bin", detail=f"EPANET version {ver}, {nn} nodes ({nt} tanks and reservoirs), {nl} links")
    if len(head) >= 36 and struct.unpack_from(">i", head, 0)[0] == SHP_FILE_CODE and \
            struct.unpack_from("<i", head, 28)[0] == 1000:
        st = struct.unpack_from("<i", head, 32)[0]
        kind = {1: "points", 3: "lines", 5: "polygons", 8: "multipoints", 11: "points with z", 13: "lines with z",
                15: "polygons with z", 21: "points with m", 23: "lines with m"}.get(st, f"shape type {st}")
        return Sniff("shp", detail=kind)
    # text
    bom = head.startswith((b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff"))
    utf16 = bom and head[:2] in (b"\xff\xfe", b"\xfe\xff")
    guess16 = None if bom else _utf16_without_bom(head)
    # trailing NUL bytes (a copy padded to a block size) do not make a text binary: WNTR stops reading at [END]
    if not (utf16 or guess16) and _is_binary(head.rstrip(b"\x00")):
        return Sniff("unknown", detail="binary content with no signature recognised", confidence="low")
    text, how = _head_text(data)
    found_all = _section_headers(text)
    deciding = CORE_SECTIONS | set(SWMM_SECTIONS) | {"SPECIES"}
    where = "in the first 64 KB"
    if (not found_all & deciding and HEAD_BYTES < len(data) <= TEXT_SCAN_MAX
            and (found_all or _only_comments(text))):
        # it starts like an EPANET file ([TITLE], [PATTERNS] or ';' comment lines) but its node and link sections
        # come later: a long title, a large pattern block first, a long comment header.  The file is in memory.
        full, how_full = decode_text(data)
        found_full = _section_headers(full)
        if found_full & deciding:
            found_all, how = found_full, how_full
            where = "in the whole file (its first 64 KB holds none of [JUNCTIONS], [PIPES], [RESERVOIRS] or [TANKS])"
    other_model = set(SWMM_SECTIONS) | {"SPECIES"}
    if (found_all & CORE_SECTIONS and not found_all & other_model and HEAD_BYTES < len(data) <= TEXT_SCAN_MAX
            and where == "in the first 64 KB"):
        # task 15b: a large SWMM model can hold only shared headers ([JUNCTIONS], [CONDUITS] later) in its first 64 KB,
        # and was typed as EPANET, then blocked with WNTR's syntax error; the whole text is searched for the
        # sections only SWMM or EPANET-MSX use (the file is in memory; the search is linear)
        full, how_full = decode_text(data)
        found_full = _section_headers(full)
        if found_full & other_model:
            found_all, how = found_full, how_full
            where = "in the whole file (after the first 64 KB, which holds only headers EPANET also uses)"
    swmm = sorted(found_all & set(SWMM_SECTIONS))
    if swmm:
        return Sniff("swmm_inp", detail=f"SWMM section{'s' if len(swmm) > 1 else ''} "
                                        f"{', '.join(f'[{x}]' for x in swmm[:4])} found {where}", encoding=how)
    msx = sorted(found_all & set(MSX_SECTIONS))
    if "SPECIES" in found_all and not found_all & {"JUNCTIONS", "RESERVOIRS"}:
        return Sniff("epanet_msx", detail=f"EPANET-MSX section{'s' if len(msx) > 1 else ''} "
                                          f"{', '.join(f'[{x}]' for x in msx[:4])} found {where}, and no "
                                          f"[JUNCTIONS] or [RESERVOIRS]", encoding=how)
    found = found_all & set(INP_SECTIONS)
    if found & CORE_SECTIONS:
        lead = text.lstrip("﻿ \t\r\n")
        extra = "; ';' comment lines come first" if lead.startswith(";") else ""
        n = f"{len(found)} EPANET section{'s' if len(found) > 1 else ''}"
        return Sniff("epanet_inp", detail=f"{n} found {where}, text in {how}{extra}",
                     confidence="high" if len(found) >= 2 else "medium", encoding=how)
    t = text.lstrip("﻿ \t\r\n")
    if t.startswith("ISO-10303-21"):
        m = re.search(r"FILE_SCHEMA\s*\(\s*\(\s*'([^']+)'", text)
        if m and m.group(1).upper().startswith("IFC"):
            return Sniff("ifc", detail=f"schema {m.group(1)}", encoding=how)
        return Sniff("ifc", detail="a STEP file (ISO-10303-21) with no IFC schema in its header", confidence="medium",
                     encoding=how)
    is_dxf, n999 = _is_ascii_dxf(text)
    if is_dxf:
        return Sniff("dxf_ascii", detail="group codes and values" + (f", {n999} comment group{'s' if n999 > 1 else ''} (999) first" if n999 else ""),
                     encoding=how)
    if t.startswith("<"):
        s = _sniff_xml(t)
        if s:
            s.encoding = how
            return s
    if t.startswith(("{", "[")):
        complete = len(data) <= JSON_MAX_BYTES
        s = _sniff_json(decode_text(data)[0].lstrip("﻿ \t\r\n") if complete else t, complete)
        if s:
            s.encoding = how
            return s
    s = _sniff_table(text)
    if s:
        s.encoding = how
        return s
    printable = sum(ch.isprintable() or ch in "\t\r\n" for ch in text[:8192]) / max(1, len(text[:8192]))
    if printable > 0.95:
        hint = " (the name ends in .inp, but the content has no EPANET sections)" if ext == ".inp" else ""
        return Sniff("text_other", detail=f"text in {how}{hint}", confidence="medium", encoding=how)
    return Sniff("unknown", detail="no signature or structure recognised", confidence="low")


# ----------------------------------------------------------------------------- IDs, before WNTR
@dataclass
class IdScan:
    counts: dict = field(default_factory=dict)        # section -> number of data lines
    blocking: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    n_coordinates: int = 0                            # data lines in [COORDINATES]
    coordinate_ids: set = field(default_factory=set)  # the node IDs those lines name


def _where(sections: list[str]) -> str:
    if len(set(sections)) == 1:
        return f"{len(sections)} times in [{sections[0]}]"
    return " and ".join(f"[{s}]" for s in sorted(set(sections)))


def scan_inp_ids(text: str) -> IdScan:
    """IDs in the node sections ([JUNCTIONS], [RESERVOIRS], [TANKS]) and the link sections ([PIPES], [PUMPS], [VALVES]),
    read from the text before WNTR parses it.  Blocking problems, each with 'export with IDs, not labels':
    - a duplicate within either group: WNTR keeps the later of two lines in one section without a word, and loads a
      junction that shares a tank's ID, after which the simulator fails;
    - an ID longer than EPANET allows, here or in [PATTERNS] and [CURVES]: EPANET counts bytes (31), and the prepared
      copy is UTF-8, so an accented ID of 31 characters in a Windows code page can be too long once copied (WNTR,
      which counts characters, stops on 32 with an AssertionError).
    IDs that differ only by upper and lower case are a warning (some EPANET versions treat them as one ID).  Also
    collects the node IDs of the [COORDINATES] lines, for the readiness report.
    The text is read the way WNTR's reader reads it, so no header form WNTR accepts gets past the scan: lines end at
    LF, CR or CR-LF only (str.splitlines would also split at a form feed or a vertical tab, which WNTR reads as
    spaces), a header's name is matched like WNTR's ([PIPE] is [PIPES]: _canonical_section), and [END] ends the
    file, as it does for WNTR."""
    out = IdScan()
    seen = {"node": {}, "link": {}}
    long_ids: list[tuple[str, str]] = []
    section = None
    for ln in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        s = ln.split(";", 1)[0].strip()
        if not s:
            continue
        if s.startswith("["):
            section = _canonical_section(s.split("]", 1)[0].lstrip("["))
            if section == "END":
                break
            continue
        if section is None:
            continue
        tok = s.split()
        if section in NODE_SECTIONS or section in LINK_SECTIONS:
            group = "node" if section in NODE_SECTIONS else "link"
            out.counts[section] = out.counts.get(section, 0) + 1
            seen[group].setdefault(tok[0], []).append(section)
            if len(tok[0].encode("utf-8")) > MAX_ID_BYTES:
                long_ids.append((section, tok[0]))
        elif section in NAMED_SECTIONS:
            if len(tok[0].encode("utf-8")) > MAX_ID_BYTES:
                long_ids.append((section, tok[0]))
        elif section == "COORDINATES":
            out.n_coordinates += 1
            out.coordinate_ids.add(tok[0])
    words = {"node": "nodes (junctions, reservoirs, tanks)", "link": "links (pipes, pumps, valves)"}
    for group in ("node", "link"):
        dup = {i: secs for i, secs in seen[group].items() if len(secs) > 1}
        if dup:
            ex = ", ".join(f"{i} ({_where(secs)})" for i, secs in list(dup.items())[:5])
            harm = []
            if any(len(set(secs)) < len(secs) for secs in dup.values()):
                harm.append("WNTR would keep only the later of two lines in one section, without a word.")
            if any(len(set(secs)) > 1 for secs in dup.values()):
                harm.append("WNTR would load an ID listed in two sections (a junction that shares a tank's ID), and "
                            "the simulation would then fail.")
            out.blocking.append(f"{len(dup)} ID{'s are' if len(dup) > 1 else ' is'} used more than once among the "
                                f"{words[group]}: {ex}. {' '.join(harm)} Export the model with IDs, not labels.")
        fold: dict[str, set] = {}
        for i in seen[group]:
            fold.setdefault(i.upper(), set()).add(i)
        case = [sorted(v) for v in fold.values() if len(v) > 1]
        if case:
            out.warnings.append(f"{len(case)} pair{'s' if len(case) > 1 else ''} of {group} IDs differ only by upper "
                                f"and lower case ({', '.join(' and '.join(c) for c in case[:3])}); some EPANET versions "
                                f"treat them as one ID.")
    if long_ids:
        first = {}
        for sec, i in long_ids:
            first.setdefault(i, sec)
        accented = any(not i.isascii() for i in first)
        size = lambda i: f"{len(i)} characters" if i.isascii() else f"{len(i.encode('utf-8'))} bytes"
        ex = ", ".join(f"{i} ({size(i)}, in [{sec}])" for i, sec in list(first.items())[:3])
        out.blocking.append(f"{len(first)} ID{'s are' if len(first) > 1 else ' is'} longer than EPANET allows "
                            f"({MAX_ID_BYTES} characters"
                            + ("; EPANET counts bytes, and an accented letter takes 2 or more in the UTF-8 copy it reads"
                               if accented else "")
                            + f"): {ex}. Export the model with IDs, not labels.")
    return out


# ----------------------------------------------------------------------------- physics settings (task 15b)
# The decay model assumes first-order reactions, one bulk rate for every pipe and tank, and a Hazen-Williams C on every
# pipe (the old-pipe factor and the roughness axis read it).  A file can say otherwise in ways WNTR reads without a
# word; these settings and checks turn each into a fix, a block or a warning in the readiness report.
G_SI = 9.81                                # m/s2
NU_20C_M2S = 1.1e-5 * 0.3048 ** 2          # EPANET's kinematic viscosity of water at 20 C, 1.1e-5 ft2/s, in m2/s; the file's
                                           # VISCOSITY option is a multiple of it
HW_K, HW_Q_EXP, HW_D_EXP = 10.67, 1.852, 4.8704   # Hazen-Williams in SI, as the plan states it:
                                                  # h/L = 10.67 Q^1.852 / (C^1.852 D^4.8704)
V_REF_FLOOR_MS = 0.3                       # a pipe's reference velocity: the median |V| over the scored day, at least this
DW_GATE_MEDIAN_H, DW_GATE_P95_H = 1.0, 4.0  # the Darcy-Weisbach gate (plan open decision 3, lead decision 3), pre-registered
                                            # before any run: water age may move by at most these over the scored day's
                                            # junction-hours between the file and its converted copy
PRESSURE_VALVES = ("PRV", "PSV", "PBV")    # valves whose setting is a pressure: they need real elevations
SCHEMATIC_SPREAD = 1.5                     # 95th over 5th percentile of drawn over declared pipe length above this: schematic
FEET_PER_METRE = 1 / 0.3048
UNIT_RATIO_TOL = 0.1                       # a median ratio within 10% of 3.281 is feet, of 1.0 is metres
CONNECTOR_LENGTH_M, CONNECTOR_DIAM_M = 0.3048, 99 * 0.0254   # the 1 ft, 99 inch pipes a WaterGEMS export adds
CONNECTOR_TAG = "addedforinpconversion"
CONNECTOR_MIN_PIPES = 3                    # untagged pipes of that size reported only from this many (Net6 has one)
LARGE_JUNCTIONS = 3000                     # above this a grid build needs an explicit confirm (app checkbox, --allow-large)
EFFECTIVE_SPEEDUP = 4.0                    # the plan's assumed effective parallel speedup of a grid build, not the worker
                                           # count: its feasibility review measured 2.3 on Net3 and 4.9 on ky10 (11
                                           # workers).  Task 15b's pre-registered bar on Net3 missed (the build took
                                           # 2.10 times the estimate, against a bar of 2); the measured runs are in
                                           # outputs/intake/run_cost_timing.json.  Not restated: that is a decision.
TEMP_MB_PER_JUNCTION = 9.5 / 959           # one 7-day run's EPANET files (.inp, .rpt, .bin): 9.5 MB on ky4 (959
                                           # junctions); ky4, ky10 and Net6 are measured in outputs/intake/
                                           # task15b_physics.json (part run_cost, temp_files_one_run)
VOLATILE_COST = ("seconds_per_run", "estimate_s", "minutes", "workers", "temp_files_mb")   # vary by run or machine:
                                                                                           # not in the record
REACTION_ORDERS = ("BULK", "WALL", "TANK")
ALL_INP_SECTIONS = frozenset(INP_SECTIONS)


def _lines_keep(text: str) -> list[list[str]]:
    """The text as [line, line end] pairs, split at LF, CR or CR-LF only (as WNTR's reader splits it), so that joining
    them gives the text back exactly."""
    parts = re.split(r"(\r\n|\r|\n)", text)
    return [[parts[i], parts[i + 1] if i + 1 < len(parts) else ""] for i in range(0, len(parts), 2)]


def _data_lines(pairs: list[list[str]]):
    """(index, section, tokens) for every data line of the text (comments taken off), with section names read as
    WNTR reads them; stops at [END]."""
    section = None
    for i, (body, _) in enumerate(pairs):
        s = body.split(";", 1)[0].strip()
        if not s:
            continue
        if s.startswith("["):
            section = _canonical_section(s.split("]", 1)[0].lstrip("["))
            if section == "END":
                return
            continue
        if section is not None:
            yield i, section, s.split()


def _num(tok: str) -> float | None:
    try:
        v = float(tok)
    except ValueError:
        return None
    return v if math.isfinite(v) else None


def fix_reactions_text(text: str) -> tuple[str, dict]:
    """The [REACTIONS] settings the decay model cannot use as they are, fixed in the text (the rest of the file is
    left exactly as it is):
    - a reaction order other than 1 (bulk, wall or tank) is set to 1 (ky10 and Net6 declare 0, and then the bulk rate
      has no effect: the grid's members are all the same);
    - per-pipe BULK and per-tank TANK lines are turned into comments: they override the global bulk rate the grid
      sets, so those pipes and tanks would ignore the calibrated rate;
    - a LIMITING POTENTIAL other than 0 is set to 0 (it changes the first-order law).
    Per-pipe WALL lines and a ROUGHNESS CORRELATION are left in the copy: the grid sets every pipe's wall rate itself,
    so they are ignored, and the report says so.  Each changed line keeps a ';ResidualMap:' comment with what the file
    had.  Returns (text, info)."""
    pairs = _lines_keep(text)
    info = {"orders": {}, "orders_fixed": {}, "bulk_lines": [], "tank_lines": [], "wall_lines": 0,
            "limiting_potential": 0.0, "limiting_potential_fixed": None, "roughness_correlation": 0.0}
    for i, sec, tok in list(_data_lines(pairs)):
        if sec != "REACTIONS" or len(tok) < 3:
            continue
        k1, k2 = tok[0].upper(), tok[1].upper()
        if k1 == "ORDER" and k2 in REACTION_ORDERS:
            v = _num(tok[2])
            info["orders"][k2.lower()] = v
            if v != 1.0:
                info["orders_fixed"][k2.lower()] = tok[2]
                pairs[i][0] = f" ORDER {k2} 1 ;ResidualMap: the file had {tok[2]}; the decay model is first order"
        elif k1 in ("BULK", "TANK"):
            info["bulk_lines" if k1 == "BULK" else "tank_lines"].append(tok[1])
            pairs[i][0] = (";" + pairs[i][0].strip() + " ;ResidualMap: cleared, the calibrated bulk rate applies to "
                                                       "every pipe and tank")
        elif k1 == "WALL":
            info["wall_lines"] += 1
        elif k1 == "LIMITING":
            v = _num(tok[2])
            info["limiting_potential"] = v
            if v != 0.0:
                info["limiting_potential_fixed"] = tok[2]
                pairs[i][0] = f" LIMITING POTENTIAL 0 ;ResidualMap: the file had {tok[2]}"
        elif k1 == "ROUGHNESS":
            info["roughness_correlation"] = _num(tok[2])
    return "".join(a + b for a, b in pairs), info


def set_headloss_text(text: str, roughness: dict[str, str]) -> str:
    """The text with [OPTIONS] HEADLOSS set to H-W and each named pipe's roughness field replaced (the sixth field of
    its [PIPES] line); every other byte is left as it is."""
    pairs = _lines_keep(text)
    for i, sec, tok in list(_data_lines(pairs)):
        if sec == "OPTIONS" and tok[0].upper() == "HEADLOSS" and len(tok) >= 2:
            pairs[i][0] = f" Headloss            H-W ;ResidualMap: the file had {tok[1]}"
        elif sec == "PIPES" and len(tok) >= 6 and tok[0] in roughness:
            body, sep, comment = pairs[i][0].partition(";")
            a, b = [m.span() for m in re.finditer(r"\S+", body)][5]
            pairs[i][0] = body[:a] + roughness[tok[0]] + body[b:] + sep + comment
    return "".join(a + b for a, b in pairs)


def hw_equivalent_c(diameter_m: float, roughness_m: float, velocity_ms: float, nu_m2s: float) -> float:
    """The Hazen-Williams C that gives the Darcy-Weisbach headloss of a pipe at velocity V:
    C = [10.67 Q^1.852 (2 g D) / (D^4.8704 f V^2)]^(1/1.852), Q = V pi D^2 / 4, with f from Swamee-Jain,
    f = 0.25 / [log10(e / (3.7 D) + 5.74 / Re^0.9)]^2, Re = V D / nu (what EPANET uses above Re 4000)."""
    D, e, V = diameter_m, roughness_m, velocity_ms
    re_ = V * D / nu_m2s
    f = 0.25 / math.log10(e / (3.7 * D) + 5.74 / re_ ** 0.9) ** 2
    q = V * math.pi * D * D / 4
    return (HW_K * q ** HW_Q_EXP * 2 * G_SI * D / (D ** HW_D_EXP * f * V * V)) ** (1 / HW_Q_EXP)


def _run_7day(path: str, parameter: str = "AGE") -> tuple:
    """A 7-day run of the file, set up as simulate.load sets up every run, in a temporary folder that is removed:
    (water age in hours at every junction, pipe velocity in m/s), each hour of the scored day (the last)."""
    import wntr
    from .simulate import _last_day, load
    wn = load(path)
    wn.options.quality.parameter = parameter
    with tempfile.TemporaryDirectory(prefix="rm_epanet_") as d, warnings.catch_warnings():
        warnings.simplefilter("ignore")
        res = wntr.sim.EpanetSimulator(wn).run_sim(file_prefix=os.path.join(d, "run"))
    age = _last_day(res.node["quality"], wn.junction_name_list) / 3600.0
    vel = _last_day(res.link["velocity"], wn.pipe_name_list).abs()
    return age, vel


def dw_to_hw(read_path: str, text: str, work_dir: str) -> tuple[str, dict]:
    """A Darcy-Weisbach model converted to Hazen-Williams, and the gate's numbers.  Each pipe's C is the one that gives
    its Darcy-Weisbach headloss at its reference velocity (the median |V| over the file's own scored day, at least
    V_REF_FLOOR_MS), with nu from the file's VISCOSITY option.  The converted copy is run too, and water age is compared
    at every junction-hour of the scored day: the gate passes when the median change is at most DW_GATE_MEDIAN_H and
    the 95th percentile at most DW_GATE_P95_H.  Exact only at the reference velocity."""
    import numpy as np
    from .simulate import load
    wn = load(read_path)
    age0, vel0 = _run_7day(read_path)
    nu = float(wn.options.hydraulic.viscosity or 1.0) * NU_20C_M2S
    c, v_ref = {}, {}
    for pn in wn.pipe_name_list:
        p = wn.get_link(pn)
        v_ref[pn] = max(float(np.median(vel0[pn].values)), V_REF_FLOOR_MS)
        c[pn] = hw_equivalent_c(float(p.diameter), float(p.roughness), v_ref[pn], nu)
    new_text = set_headloss_text(text, {pn: f"{v:.3f}" for pn, v in c.items()})
    conv = os.path.join(work_dir, "converted.inp")
    with open(conv, "w", encoding="utf-8", newline="") as fh:
        fh.write(new_text)
    age1, _ = _run_7day(conv)
    d = np.abs(age1[age0.columns].values - age0.values).ravel()
    med, p95 = float(np.median(d)), float(np.percentile(d, 95))
    vals = np.array(list(c.values()))
    return new_text, {"pipes": len(c), "c_min": float(vals.min()), "c_max": float(vals.max()),
                      "c": {k: round(v, 3) for k, v in c.items()}, "v_ref": {k: round(v, 4) for k, v in v_ref.items()},
                      "age_change_median_h": med, "age_change_p95_h": p95,
                      "passes": med <= DW_GATE_MEDIAN_H and p95 <= DW_GATE_P95_H, "converted_path": conv}


def control_scan(text: str) -> dict:
    """What the [CONTROLS] and [RULES] act on and read, and the [TAGS] a WaterGEMS export adds, from the text:
    the link IDs a control or rule sets, the node IDs a control or rule reads (any NODE in a simple control; a NODE or
    JUNCTION whose PRESSURE or HEAD a rule's premise reads), and the number of each.  In a rule only the action
    clauses set a link (THEN and ELSE, and the AND lines after them); a link named in a premise (IF, and the AND and OR
    lines after it) is only read, so a pump that a rule reads and never sets is not controlled (found by the review:
    'IF PUMP 10 STATUS IS OPEN THEN PIPE 20 STATUS IS OPEN' left pump 10 closed for 7 days without a warning)."""
    pairs = _lines_keep(text)
    links, nodes, n_controls, n_rules, tags = set(), set(), 0, 0, []
    clause = None                         # in [RULES]: 'premise' or 'action', the clause the current line belongs to
    for _, sec, tok in _data_lines(pairs):
        up = [t.upper() for t in tok]
        if sec == "CONTROLS":
            n_controls += 1
            if up[0] == "LINK" and len(tok) > 1:
                links.add(tok[1])
            nodes.update(tok[k + 1] for k in range(len(up) - 1) if up[k] == "NODE")
        elif sec == "RULES":
            if up[0] == "RULE":
                n_rules += 1
                clause = None
                continue
            if up[0] == "IF":
                clause = "premise"
            elif up[0] in ("THEN", "ELSE"):
                clause = "action"
            elif up[0] == "PRIORITY":
                clause = None
            for k in range(len(up) - 1):
                if clause == "action" and up[k] in ("LINK", "PIPE", "PUMP", "VALVE"):
                    links.add(tok[k + 1])
                elif (clause == "premise" and up[k] in ("NODE", "JUNCTION") and k + 2 < len(up)
                      and up[k + 2] in ("PRESSURE", "HEAD")):
                    nodes.add(tok[k + 1])
        elif sec == "TAGS" and len(tok) >= 3 and any(t.lower() == CONNECTOR_TAG for t in tok[2:]):
            tags.append(tok[1])
    return {"links": links, "nodes": nodes, "n_controls": n_controls, "n_rules": n_rules, "connector_tags": tags}


def coordinate_scale(wn, have_xy: set) -> dict:
    """Drawn pipe length (with vertices) over declared length, for every pipe with both ends drawn: a spread (95th
    percentile over 5th) above SCHEMATIC_SPREAD is a schematic drawing; otherwise the median names the units (3.281
    per metre is feet, 1.0 is metres)."""
    import numpy as np
    if not have_xy:
        return {"status": "no coordinates", "pipes_measured": 0}
    r = []
    for pn in wn.pipe_name_list:
        p = wn.get_link(pn)
        if p.start_node_name not in have_xy or p.end_node_name not in have_xy or not p.length > 0:
            continue
        pts = [p.start_node.coordinates, *p.vertices, p.end_node.coordinates]
        d = sum(math.dist(pts[k], pts[k + 1]) for k in range(len(pts) - 1))
        if d > 0:
            r.append(d / p.length)
    if len(r) < 2:
        return {"status": "too few pipes drawn", "pipes_measured": len(r)}
    p5, p50, p95 = (float(x) for x in np.percentile(r, [5, 50, 95]))
    spread = p95 / p5
    out = {"pipes_measured": len(r), "ratio_p5": round(p5, 4), "ratio_median": round(p50, 4),
           "ratio_p95": round(p95, 4), "spread": round(spread, 2)}
    if spread > SCHEMATIC_SPREAD:
        return {"status": "schematic", "units": "unknown", **out}
    units = ("feet" if abs(p50 / FEET_PER_METRE - 1) <= UNIT_RATIO_TOL else
             "metres" if abs(p50 - 1) <= UNIT_RATIO_TOL else "unknown")
    return {"status": "to scale", "units": units, **out}


def grid_for(disinfectant: str, uniform_c: bool) -> str:
    """The simulator grid a file gets: chloramine's for chloramine; for free chlorine, 'full', or 'full_fixed_gamma'
    (the old-pipe factor held at 0) when every pipe has one C value and the factor cannot be learned."""
    if disinfectant == "chloramine":
        return "chloramine"
    return "full_fixed_gamma" if uniform_c else "full"


def disk_free_gb(path: str) -> float:
    """Free disk in GB (simgp.disk_free_gb); here so a saved check can stand in a value."""
    from .simgp import disk_free_gb as free
    return free(path)


def size_gate(cost: dict, path: str = ".", free_gb: float | None = None) -> tuple[bool, str]:
    """(ok, plain message): a grid build is refused when its file plus the temporary EPANET files of one run per worker
    would leave less than simgp.MIN_FREE_GB free on the volume of `path`."""
    from .simgp import MIN_FREE_GB
    free = disk_free_gb(path) if free_gb is None else float(free_gb)
    need = (cost["grid_file_mb"] + cost["temp_files_mb"]) / 1000
    left = free - need
    msg = (f"The grid needs about {cost['grid_file_mb']:g} MB for its file and up to {cost['temp_files_mb']:g} MB of "
           f"temporary EPANET files while it runs ({cost['workers']} at a time); this disk has {free:.2f} GB free, so "
           f"{left:.2f} GB would be left")
    if left < MIN_FREE_GB:
        return False, (msg + f", below the {MIN_FREE_GB:g} GB this tool keeps free. Free some disk space, or use a "
                             f"smaller zone of the model.")
    return True, msg + "."


# ----------------------------------------------------------------------------- the readiness report
@dataclass
class Readiness:
    name: str
    disinfectant: str = "free_chlorine"
    sha1: str = ""                                   # of the file as received
    sniff: dict = field(default_factory=dict)
    not_a_model: bool = False
    rows: list = field(default_factory=list)         # (item, value, source), source in SOURCES
    blocking: list = field(default_factory=list)
    fixes: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    run_cost: dict = field(default_factory=dict)
    prepared: str | None = None                      # the prepared file's name (never its folder)
    same_file: bool = False                          # the prepared file is the file given, untouched
    flags: dict = field(default_factory=dict)        # task 15b: what the physics checks found, by check (bundled_dict)

    def add(self, item: str, value, source: str) -> None:
        assert source in SOURCES, source
        self.rows.append((item, str(value), source))

    @property
    def ok(self) -> bool:
        return not self.not_a_model and not self.blocking

    def advice(self) -> str:
        """For a file that is not a model: what it is and what to send instead, in plain words."""
        s = Sniff(**{k: v for k, v in self.sniff.items() if k in Sniff.__dataclass_fields__}) if self.sniff else None
        if s is None:
            return "This file could not be read."
        return f"{s.sentence()} ResidualMap reads EPANET models. {s.advice}"

    def summary(self) -> str:
        if self.not_a_model:
            return self.advice()
        if self.blocking:
            n = len(self.blocking)
            return f"{n} problem{'s' if n > 1 else ''} stop{'' if n > 1 else 's'} this model from running."
        if any(item == TEST_RUN_ROW for item, _, _ in self.rows):
            bits = ["WNTR reads it and EPANET runs its first hour without an error, and one full 7-day run (timed for "
                    "the run cost); the grid's runs come next. Its physics settings are checked in the rows below "
                    "(reaction orders, bulk lines, headloss formula, start time, valves and elevations, coordinates, "
                    "sources, pumps and controls, connectivity, pipe roughness)."]
        else:                          # prepare_inp(parse=False)
            bits = ["Prepared without a check: WNTR has not read it and EPANET has not run it; only the reaction "
                    "settings in its text were checked."]
        if self.fixes:
            bits.append(f"{len(self.fixes)} fix{'es' if len(self.fixes) > 1 else ''} applied to a copy; your file is "
                        f"not changed.")
        if self.warnings:
            bits.append(f"{len(self.warnings)} warning{'s' if len(self.warnings) > 1 else ''}.")
        return " ".join(bits)

    def to_markdown(self) -> str:
        out = [f"**Model readiness: {self.name}**", "", self.summary(), ""]
        if self.not_a_model:
            return "\n".join(out)
        if self.blocking:
            out += ["**Problems that stop the run:**"] + [f"- {b}" for b in self.blocking] + [""]
        if self.rows:
            out += ["| item | value | source |", "|---|---|---|"]
            out += [f"| {i} | {v.replace('|', '/')} | {s} |" for i, v, s in self.rows] + [""]
        if self.fixes and self.blocking:
            out += ["**Fixes that would be applied** (no copy is kept while a problem stops the run; your file is not "
                    "changed):"] + [f"- {f}" for f in self.fixes] + [""]
        elif self.fixes:
            out += ["**Fixes applied to a copy** (your file is not changed):"] + [f"- {f}" for f in self.fixes] + [""]
        if self.warnings:
            out += ["**Warnings:**"] + [f"- {w}" for w in self.warnings] + [""]
        return "\n".join(out).rstrip() + "\n"

    def as_dict(self) -> dict:
        return {"name": self.name, "disinfectant": self.disinfectant, "sha1": self.sha1, "sniff": self.sniff,
                "not_a_model": self.not_a_model, "ok": self.ok, "summary": self.summary(),
                "rows": [{"item": i, "value": v, "source": s} for i, v, s in self.rows], "blocking": self.blocking,
                "fixes": self.fixes, "warnings": self.warnings, "run_cost": self.run_cost, "prepared": self.prepared,
                "same_file": self.same_file, "flags": self.flags}

    def bundled_dict(self) -> dict:
        """What outputs/intake/readiness_bundled.json keeps of a report: everything that does not vary from run to run
        or machine to machine (the measured run time, the minutes estimated from it and the worker count are left out;
        they go to the checks' run log), and no file name."""
        cost = {k: v for k, v in self.run_cost.items() if k not in VOLATILE_COST}
        flags = {k: v for k, v in self.flags.items() if k != "run_cost"}
        return {"ok": self.ok, "blocking": self.blocking, "fixes": self.fixes, "warnings": self.warnings, "run_cost": cost,
                "flags": flags, "rows": [{"item": i, "value": v, "source": s} for i, v, s in self.rows
                                         if i not in (RUN_COST_ROW, "File", "Text encoding")]}

    def to_json(self) -> str:
        return json.dumps(self.as_dict(), indent=1)


def stem(name: str) -> str:
    """The name without its folder, any .bak and one .inp or .txt (or other) extension, made safe for a file name."""
    b = os.path.basename(name or "model")
    while b.lower().endswith(".bak"):
        b = b[:-4]
    root, ext = os.path.splitext(b)
    b = root if ext else b
    b = re.sub(r"[^A-Za-z0-9._-]+", "_", b).strip("._") or "model"
    return b[:80]


def _sha1(b: bytes) -> str:
    return hashlib.sha1(b).hexdigest()


def run_cost(n_junctions: int, disinfectant: str, uniform_c: bool = False, seconds_per_run: float | None = None,
             workers: int | None = None) -> dict:
    """The simulator grid for one dose (task 15b): its EPANET runs, its cache file (ln C as float32 for every run, hour
    of the last day and junction: 64.8 KB per junction for the 675-run free-chlorine grid, 138.2 KB for chloramine's
    1440 runs), the temporary EPANET files of one run per worker while it builds, and, given one timed 7-day run, the
    minutes it takes (runs times that run, over the measured effective speedup EFFECTIVE_SPEEDUP, not the worker
    count).  'large': more junctions than LARGE_JUNCTIONS, which needs an explicit confirm."""
    from .simgp import GRIDS
    grid = grid_for(disinfectant, uniform_c)
    runs = 1
    for axis in GRIDS[grid]:
        runs *= len(axis)
    mb = runs * 24 * n_junctions * FLOAT32_BYTES / 1e6
    workers = workers or max(1, (os.cpu_count() or 2) - 1)
    out = {"grid": grid, "runs_per_dose": runs, "grid_file_mb": round(mb, 1), "junctions": n_junctions,
           "large": n_junctions > LARGE_JUNCTIONS, "workers": workers,
           "temp_files_mb": round(TEMP_MB_PER_JUNCTION * n_junctions * workers, 1)}
    if seconds_per_run is not None:
        est = estimate_seconds(float(seconds_per_run), runs)
        out["seconds_per_run"] = round(float(seconds_per_run), 3)
        out["estimate_s"] = est                   # unrounded: what run_cost_gate scores
        out["minutes"] = round(est / 60, 1)
    return out


def estimate_seconds(seconds_per_run: float, runs: int) -> float:
    """The grid build's time estimate (task 15b, item 10, as the plan states it): one timed 7-day run times the number
    of runs, over the assumed effective parallel speedup EFFECTIVE_SPEEDUP.  On a small model, where starting the
    workers takes most of the time, the build takes longer than this, and on a large one less
    (outputs/intake/run_cost_timing.json: Net3's builds took 1.4 to 2.1 times the estimate, ky10's 0.64)."""
    return seconds_per_run * runs / EFFECTIVE_SPEEDUP


def prepared_name(digest: str, name: str, keep_name: bool) -> str:
    """The prepared copy's file name: <digest>_<name> when keep_name (a clean upload, so the app's upload name and
    grid cache name stay as they were before task 15) and the name ends in .inp and fits in MAX_NAME_BYTES of UTF-8;
    otherwise <digest>_<stem>.inp, whose stem is at most 80 safe characters.  A file name has at most 255 bytes, and the
    grid cache's name adds its own words to this one, so a long upload name is shortened here, never refused."""
    base = os.path.basename(name or "")
    if keep_name and base.lower().endswith(".inp") and len(f"{digest}_{base}".encode("utf-8")) <= MAX_NAME_BYTES:
        return f"{digest}_{base}"
    return f"{digest}_{stem(name)}.inp"


def prepare_inp(src, out_dir: str, disinfectant: str = "free_chlorine", name: str | None = None,
                parse: bool = True) -> tuple[str | None, Readiness]:
    """(inp_path, Readiness) for a model file given as a path or as bytes.

    A clean UTF-8 file given as a path comes back as that path, untouched.  Clean bytes are written to out_dir as
    <sha1 of the bytes>[:8]_<name> (the app's upload name; a name not ending in .inp, or too long, gets <stem>.inp:
    prepared_name).  A file that needed a fix (a byte-order mark, UTF-16, a Windows code page, a model inside a zip)
    is written to out_dir as <sha1 of the new text>[:8]_<stem>.inp; the original is never changed.  A file that is not
    a model, or has a blocking problem, gives (None, Readiness) and leaves nothing in out_dir: a new copy is written
    under a temporary name, read by WNTR from there (after the ID scan, never before), and moved to its name only when
    WNTR read it, the model has an entry point and EPANET ran its first hour; on a blocking problem the temporary file
    is removed.  parse=True does that load and test run, for the report's counts and checks; a failure in either is
    a blocking problem.  With parse=False the copy is moved into place at once, unread.

    Never raises for what is in the file: a failure anywhere after the file is read (a zip that breaks on reading, a
    copy that cannot be saved, an error no test foresaw) comes back as a blocking problem in plain words, with
    nothing left in out_dir.  A wrong disinfectant, or a path that does not exist, still raises."""
    if disinfectant not in DISINFECTANT_WORDS:
        raise ValueError(f"disinfectant must be one of {sorted(DISINFECTANT_WORDS)}")
    if isinstance(src, (bytes, bytearray, memoryview)):
        data, path = bytes(src), None
        name = name or "upload.inp"
    else:
        path = os.fspath(src)
        name = name or os.path.basename(path)
        with open(path, "rb") as fh:
            data = fh.read()
    rep = Readiness(name=os.path.basename(name), disinfectant=disinfectant, sha1=_sha1(data))
    try:
        return _prepare(data, path, name, out_dir, disinfectant, parse, rep)
    except Exception as e:  # noqa: BLE001  (a crafted or broken file must reach the user as a plain problem)
        rep.prepared, rep.same_file = None, False
        rep.blocking.append(f"This file could not be prepared ({type(e).__name__}); nothing was kept. Send the "
                            f"model's EPANET export (.inp) again, on its own.")
        return None, rep


def _stale_names(name: str, contents: list[bytes]) -> set[str]:
    """Every content name a prepared copy of this file can have had, under task 15's scheme and this one: the digest
    of the bytes as received, as decoded, or after the reaction fixes, with the upload name kept or the safe stem.
    The digest makes each name this file's own."""
    out = set()
    for b in contents:
        d = _sha1(b)[:8]
        out.update(prepared_name(d, name, keep) for keep in (True, False))
    return out


def _remove_stale(out_dir: str, names: set[str], keep: str | None) -> None:
    """A blocked model leaves no prepared copy behind (task 15b): a content-named copy that an earlier version (or an
    earlier, less strict check) accepted is removed from out_dir.  The file given (keep) is never removed."""
    for n in names:
        p = os.path.join(out_dir, n)
        if keep is not None and os.path.abspath(p) == os.path.abspath(keep):
            continue
        if os.path.isfile(p) and not os.path.islink(p):
            with contextlib.suppress(OSError):
                os.remove(p)


def _write_temp(out_dir: str, b: bytes) -> str:
    """b in a new temporary file in out_dir (removed again if the write fails)."""
    fd, tmp = tempfile.mkstemp(prefix=".prepare_", suffix=".inp", dir=out_dir)
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(b)
    except BaseException:
        with contextlib.suppress(OSError):
            os.remove(tmp)
        raise
    return tmp


def _decode_fix(rep: Readiness, data: bytes, how: str) -> None:
    """The decoding fix (and, for a file that mixes two encodings, a warning naming the lines)."""
    if how == "utf-8":
        return
    if how.startswith(MIXED_PREFIX):
        _, n, lines, n_lines, pairs = _mixed_utf8(data)
        chance = (f", {2 * pairs if pairs > 1 else 'two'} of them in {pairs if pairs > 1 else 'a'} pair"
                  f"{'s' if pairs > 1 else ''} that {'are' if pairs > 1 else 'is'} valid UTF-8 only by chance (an "
                  f"accented capital before a quote, a dash or a degree sign, as in 'CITÉ’')" if pairs else "")
        rep.fixes.append(f"Read the text as UTF-8, with {n} stray byte{'s' if n != 1 else ''} read as cp1252 (Windows "
                         f"Western European){chance}, and saved it as UTF-8 (WNTR reads UTF-8 only).")
        rep.warnings.append(f"The file mixes two text encodings: it is UTF-8 except for {n} byte{'s' if n != 1 else ''} "
                            f"(on line{'s' if n_lines != 1 else ''} {', '.join(map(str, lines))}"
                            f"{' and more' if n_lines > len(lines) else ''}), which {'were' if n != 1 else 'was'} "
                            f"read as cp1252 (Windows Western European), as a pasted degree sign or accented letter "
                            f"usually is. If a name or an ID there looks wrong, correct it in the modelling package and "
                            f"export again.")
        return
    what = {"utf-8 with a byte-order mark": "Removed the byte-order mark at the start (WNTR stops at line 1 on it)"}
    rep.fixes.append(what.get(how, f"Decoded the text from {how} and saved it as UTF-8 (WNTR reads UTF-8 only)") + ".")


def _ids_list(ids: list[str], k: int = 6) -> str:
    return ", ".join(ids[:k]) + (f" and {len(ids) - k} more" if len(ids) > k else "")


def _reaction_fixes(rep: Readiness, info: dict) -> None:
    """The reaction settings fix_reactions_text changed, as fixes (task 15b, items 1 and 2), kept for the rows."""
    rep.flags["reactions"] = {"orders": info["orders"], "orders_fixed": info["orders_fixed"],
                              "bulk_lines": len(info["bulk_lines"]), "tank_lines": len(info["tank_lines"]),
                              "wall_lines": info["wall_lines"], "limiting_potential": info["limiting_potential"],
                              "limiting_potential_fixed": info["limiting_potential_fixed"],
                              "roughness_correlation": info["roughness_correlation"]}
    of = info["orders_fixed"]
    if of:
        rep.fixes.append("Set the reaction order to 1, as the decay model assumes (first order), where the file had "
                         + " and ".join(f"{k} {v}" for k, v in of.items()) + ". An order of 0 is another law (a fixed "
                         "amount lost per day), so the decay rates would not mean what they say, and the bulk rate can "
                         "have no effect at all.")
    b, t = info["bulk_lines"], info["tank_lines"]
    if b or t:
        parts = ([f"the BULK lines of {len(b)} pipe{'s' if len(b) != 1 else ''} ({_ids_list(b)})"] if b else []) + \
                ([f"the TANK lines of {len(t)} tank{'s' if len(t) != 1 else ''} ({_ids_list(t)})"] if t else [])
        rep.fixes.append("Cleared " + " and ".join(parts) + ": a per-pipe or per-tank bulk rate overrides the one the "
                         "model calibrates, so those pipes and tanks would ignore it. The calibrated bulk rate now "
                         "applies to every pipe and tank.")
    if info["limiting_potential_fixed"] is not None:
        rep.fixes.append(f"Set LIMITING POTENTIAL to 0, where the file had {info['limiting_potential_fixed']}: with a "
                         f"limiting potential the bulk decay is not first order in the chlorine.")


def _reaction_rows(rep: Readiness) -> None:
    r = rep.flags.get("reactions")
    if r is None:
        return
    if r["orders_fixed"]:
        had = ", ".join(f"{k} {v}" for k, v in r["orders_fixed"].items())
        rep.add("Reaction orders", f"the file has {had}; the copy uses 1 (first order) for bulk, wall and tank, as the "
                                   f"decay model assumes", DERIVED)
    else:
        rep.add("Reaction orders", "1 (first order) for bulk, wall and tank, as the decay model assumes"
                + ("" if len(r["orders"]) == 3 else " (EPANET's default where the file sets none)"), FROM_FILE)
    bits = []
    if r["bulk_lines"] or r["tank_lines"]:
        bits.append(f"{r['bulk_lines']} per-pipe BULK and {r['tank_lines']} per-tank TANK lines cleared in the copy, "
                    f"so the calibrated bulk rate applies everywhere")
    if r["limiting_potential_fixed"] is not None:
        bits.append(f"LIMITING POTENTIAL {r['limiting_potential_fixed']} set to 0")
    if r["wall_lines"] or r["roughness_correlation"]:
        bits.append((f"{r['wall_lines']} per-pipe WALL lines" if r["wall_lines"] else "")
                    + (" and " if r["wall_lines"] and r["roughness_correlation"] else "")
                    + (f"a ROUGHNESS CORRELATION of {r['roughness_correlation']:g}" if r["roughness_correlation"] else "")
                    + " ignored: the model sets every pipe's wall rate from its C")
    rep.add("Per-pipe and per-tank rates", "; ".join(bits) if bits else "none: one bulk rate for every pipe and tank, "
                                                                        "as the decay model assumes",
            DERIVED if bits else FROM_FILE)


def _prepare(data: bytes, path: str | None, name: str, out_dir: str, disinfectant: str, parse: bool,
             rep: Readiness) -> tuple[str | None, Readiness]:
    s = sniff(data, name)
    rep.sniff = s.as_dict()
    if s.kind == "epanet_inp_zip":
        if len(s.members) != 1:
            rep.blocking.append(f"The zip holds {len(s.members)} EPANET models ({', '.join(s.members[:5])}): send one "
                                f"model at a time.")
            return None, rep
        try:
            data = read_zip_member(data, s.members[0])
        except ZipRefused as e:
            rep.blocking.append(f"The zip is refused: {e}.")
            return None, rep
        rep.fixes.append(f"Read {s.members[0]} from the zip in memory (the zip was not unpacked); the model is saved as "
                         f"a prepared copy.")
        name, path = s.members[0], None
        inner = sniff(data, name)
        if inner.kind != "epanet_inp":
            rep.not_a_model = True
            rep.sniff = inner.as_dict()
            return None, rep
        s = inner
    elif not s.is_model:
        rep.not_a_model = True
        return None, rep
    rep.add("File", "an EPANET model (.inp), identified by its content"
            + ("; the name marks it as a backup copy" if s.backup else ""), FROM_FILE)
    text, how = decode_text(data)
    rep.add("Text encoding", how, FROM_FILE)
    _decode_fix(rep, data, how)
    ids = scan_inp_ids(text)
    rep.blocking += ids.blocking
    rep.warnings += ids.warnings
    fixed, rinfo = (text, None) if rep.blocking else fix_reactions_text(text)
    stale = _stale_names(name, [data, text.encode("utf-8"), fixed.encode("utf-8")])
    if rep.blocking:
        _remove_stale(out_dir, stale, path)
        return None, rep
    _reaction_fixes(rep, rinfo)
    changed = fixed != text
    # write only when needed, under a temporary name until WNTR has read it, an entry point is found, the physics
    # checks and the test runs passed; a temporary file left by any failure (the save itself included) is removed in
    # the finally, and so is a content-named copy of this file that an earlier version accepted (task 15b)
    tmp, tmp2, accepted, out_path = None, None, False, None
    try:
        if path is not None and how == "utf-8" and not changed:
            read_path = out_path = path
            rep.same_file = True
        else:
            new = data if how == "utf-8" and not changed else fixed.encode("utf-8")
            fname = prepared_name(_sha1(new)[:8], name, keep_name=how == "utf-8" and not rep.fixes)
            os.makedirs(out_dir, exist_ok=True)
            out_path = os.path.join(out_dir, fname)
            read_path = tmp = _write_temp(out_dir, new)
        rep.prepared = os.path.basename(out_path)
        replacement = None
        if parse:
            accepted, replacement = _read_and_report(read_path, out_path, rep, ids, disinfectant, fixed)
        else:
            accepted = True
        if accepted and replacement is not None:      # a Darcy-Weisbach model, converted: the copy holds the new text
            new2 = replacement.encode("utf-8")
            out_path = os.path.join(out_dir, prepared_name(_sha1(new2)[:8], name, keep_name=False))
            os.makedirs(out_dir, exist_ok=True)
            tmp2 = _write_temp(out_dir, new2)
            os.replace(tmp2, out_path)
            tmp2 = None
            rep.prepared, rep.same_file = os.path.basename(out_path), False
        elif accepted and tmp is not None:
            os.replace(tmp, out_path)
            tmp = None
    except OSError as e:
        accepted = False
        rep.blocking.append(f"The prepared copy could not be saved ({e.strerror or type(e).__name__}); nothing was "
                            f"kept.")
    finally:
        for t in (tmp, tmp2):
            if t is not None:
                with contextlib.suppress(OSError):
                    os.remove(t)
        if not accepted:
            rep.prepared, rep.same_file = None, False
            _remove_stale(out_dir, stale, path)
    return (out_path if accepted else None), rep


def _test_run(wn) -> str | None:
    """EPANET on the model's first hour, set up the way the app's first run is (simulate.load's time steps, water
    age), in a temporary folder that is removed afterwards.  None when it runs; otherwise the problem in plain words.
    WNTR reading a file does not mean EPANET runs it: a node that no pipe, pump or valve connects reads fine and then
    stops EPANET with its Error 200 (found by the final verification of task 15)."""
    import wntr
    from .simulate import QUALITY_STEP_S
    t = wn.options.time
    t.duration, t.hydraulic_timestep, t.report_timestep, t.quality_timestep = 3600, 3600, 3600, QUALITY_STEP_S
    wn.options.quality.parameter = "AGE"
    try:
        with tempfile.TemporaryDirectory(prefix="rm_test_run_") as d, warnings.catch_warnings():
            warnings.simplefilter("ignore")
            prefix = os.path.join(d, "check")
            try:
                wntr.sim.EpanetSimulator(wn).run_sim(file_prefix=prefix)
            except Exception as e:  # noqa: BLE001
                msg = str(e).replace(prefix, "the test copy").replace(d, "").replace(" %s", "").replace("%s", "")
                linked = set()
                for _, link in wn.links():
                    linked.update((link.start_node_name, link.end_node_name))
                lone = [n for n in wn.node_name_list if n not in linked]
                hint = ""
                if lone:
                    hint = (f" {len(lone)} node{'s' if len(lone) > 1 else ''} in the file "
                            f"{'are' if len(lone) > 1 else 'is'} not connected to any pipe, pump or valve "
                            f"({', '.join(lone[:5])}{' and more' if len(lone) > 5 else ''}), and EPANET stops on such "
                            f"a node.")
                return (f"EPANET could not run the first hour of this model ({type(e).__name__}: {msg[:200]})." + hint
                        + " Open the model in EPANET or the modelling package, fix what it reports, and export it "
                          "again.")
    except OSError as e:          # the temporary folder itself
        return f"The EPANET test run could not start ({e.strerror or type(e).__name__})."
    return None


ELEVATION_EVIDENCE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs", "intake",
                                  "elevation_prv_ky10.json")


def _elevation_evidence() -> str:
    """One sentence from the committed outputs/intake/elevation_prv_ky10.json, or '' when it is not there."""
    try:
        with open(ELEVATION_EVIDENCE) as fh:
            d = json.load(fh)
        return (f" In a test on the public model ky10 (its {d['pressure_valves']} PRVs), setting every junction to one "
                f"elevation moved the water age by more than {d['threshold_h']:g} h at {d['junctions_moved']} of "
                f"{d['junctions']} junctions, by up to {d['max_change_h']:.1f} h (outputs/intake/elevation_prv_ky10.json).")
    except (OSError, KeyError, ValueError, TypeError):
        return ""


def _clock_minute(hours: float) -> int:
    """A time in hours as the minute of the day, rounded once (5:59:59 is minute 360; 23:59:50 and 0:00:20 are 0)."""
    return int(round(hours * 60)) % 1440


def _clock(hours: float) -> str:
    """HH:MM of a time in hours, rounded to the minute once (so 5:59:59 is 06:00, never 05:60)."""
    m = _clock_minute(hours)
    return f"{m // 60:02d}:{m % 60:02d}"


def _clock_report(rep: Readiness, clock_h: float) -> None:
    """The start clock row and, for a start other than midnight, the model-hours warning (task 15b, item 4).  Both
    quote the start rounded to the minute once, and the warning goes by that minute (the second review: a start within
    30 s of midnight is 00:00 and gets no warning, and the offset quoted is the rounded one)."""
    m = _clock_minute(clock_h)
    off = f"{m / 60:g} h"
    rep.add("Start clock time", f"{off} ({_clock(clock_h)}): the model's hour 0 is this time of day; recorded, the file "
                                f"is not changed", FROM_FILE)
    if m:
        rep.warnings.append(f"The model starts at {_clock(clock_h)} (START CLOCKTIME), so its hour 0 is {_clock(clock_h)}, "
                            f"not midnight. When this file is uploaded to the app (or run with the pilot's --inp), the "
                            f"app and the pilot convert (task 16): a grab log's times of day become model hours (the "
                            f"time minus {off}, mod 24 h), and the app shows the route's hours, the worst hour, the "
                            f"hour picker and each junction's hour of the daily minimum as times of day. Hours typed "
                            f"into the app's sample table stay model hours: take {off} off a sample's time of day "
                            f"before typing it. Add {off} to a model hour for the time of day.")


def _physics_report(rep: Readiness, wn, ctl: dict, have_xy: set, src_nodes: list[str]) -> set[str]:
    """The task-15b checks that read the model WNTR loaded (no EPANET run): headloss formula, start clock and pattern
    lengths, pressure-control valves and elevations, coordinates, sources and entry points, pumps and controls,
    demand, connectivity and export connectors.  Rows, warnings and blocking problems go to rep; what each found goes
    to rep.flags.  Returns the export's connector pipes when the Export connectors row recognises them (the pipes
    tagged AddedForINPConversion, and the pipes of 1 ft and 99 in once the row is shown), which the uniform-C test
    leaves out; else an empty set."""
    import networkx as nx
    from wntr.network import LinkStatus
    f = rep.flags
    _reaction_rows(rep)
    # 3. headloss formula (the conversion itself runs later, after the test run)
    hl = wn.options.hydraulic.headloss
    f["headloss"] = hl
    if hl == "H-W":
        rep.add("Headloss formula", "H-W (Hazen-Williams), as the decay model assumes: the wall decay and the "
                                    "roughness axis read each pipe's roughness as a C", FROM_FILE)
    elif hl == "D-W":
        rep.add("Headloss formula", "D-W (Darcy-Weisbach) in your file; converted to Hazen-Williams in the copy, each "
                                    "pipe's C matched to its headloss at its reference velocity (next rows)", DERIVED)
    else:
        rep.add("Headloss formula", f"{hl} (Chezy-Manning) in your file", FROM_FILE)
        rep.blocking.append("The file uses the Chezy-Manning headloss formula. The wall decay reads each pipe's roughness "
                            "as a Hazen-Williams C, and a Manning n has no such reading, so the decay rates would not "
                            "mean what they say. Export the model with Hazen-Williams, or ask your modeller.")
    # 4. start clock time and pattern lengths
    t = wn.options.time
    clock_h = float(t.start_clocktime or 0) / 3600
    f["start_clock_h"] = clock_h
    _clock_report(rep, clock_h)
    step = float(t.pattern_timestep or 3600)
    weekly, multi = [], []
    for pn in wn.pattern_name_list:
        dur_h = len(wn.get_pattern(pn).multipliers) * step / 3600
        if abs(dur_h - 168) < 1e-9:
            weekly.append(pn)
        elif dur_h > 24 + 1e-9:
            multi.append([pn, dur_h])
    f["weekly_patterns"], f["multi_day_patterns"] = weekly, multi
    if weekly:
        rep.warnings.append(f"{len(weekly)} weekly pattern{'s' if len(weekly) != 1 else ''} ({_ids_list(weekly)}: a "
                            f"week long). Every run is 7 days and only the last day is scored, so the scored day is "
                            f"the pattern's day 7, six days after the weekday it starts on. The file does not record "
                            f"that weekday: tell us which day of the week the pattern starts on, so the scored day can "
                            f"be named.")
    if multi:
        rep.warnings.append(f"{len(multi)} pattern{'s' if len(multi) != 1 else ''} longer than a day and not a week ("
                            + ", ".join(f"{pn}: {d:g} h" for pn, d in multi[:6]) + "): EPANET repeats each one over its "
                            "own length, so the scored day (the last of the 7-day run) is not a repeat of one daily "
                            "cycle.")
    # 5. pressure-control valves and elevations
    junctions = wn.junction_name_list
    kinds: dict[str, list[str]] = {}
    for vn in wn.valve_name_list:
        kinds.setdefault(str(wn.get_link(vn).valve_type), []).append(vn)
    pvalves = [vn for k in PRESSURE_VALVES for vn in kinds.get(k, [])]
    elev = {j: float(wn.get_node(j).elevation) for j in junctions}
    zeros = [j for j in junctions if elev[j] == 0.0]
    constant = len(junctions) > 1 and len(set(elev.values())) == 1
    pnodes = [j for j in junctions if j in ctl["nodes"]]
    need = set(pnodes)
    for vn in pvalves:
        link = wn.get_link(vn)
        need.update((link.start_node_name, link.end_node_name))
    need_j = [j for j in junctions if j in need]
    emitters = [j for j in junctions if (wn.get_node(j).emitter_coefficient or 0) > 0]
    pressure_driven = str(wn.options.hydraulic.demand_model).upper() in ("PDA", "PDD")
    f["pressure_valves"] = {k: len(kinds.get(k, [])) for k in PRESSURE_VALVES if kinds.get(k)}
    f["zero_elevation_junctions"] = zeros
    f["zero_elevation_flagged"] = bool(pvalves and zeros and not constant)
    f["constant_elevation"] = constant
    f["elevations_needed_at"] = len(need_j)
    if pvalves:
        what = ", ".join(f"{len(kinds[k])} {k}" for k in PRESSURE_VALVES if kinds.get(k))
        at_zero = [j for j in need_j if j in set(zeros)]
        rep.add("Elevations", f"{len(pvalves)} pressure-control valve{'s' if len(pvalves) != 1 else ''} ({what}): the "
                              f"elevations change the flows here; real elevations are needed at the valves"
                              + (", downstream of them and in pressure-based controls" if pnodes else
                                 " and downstream of them")
                              + f"; {len(need_j)} junction{'s' if len(need_j) != 1 else ''} "
                              + ("sit on either side of a valve or in such a control" if pnodes else
                                 "sit on either side of a valve")
                              + (f", {len(at_zero)} of them at elevation 0" if at_zero else
                                 ", none of them at elevation 0"), FROM_FILE)
        if constant:
            rep.blocking.append(
                f"Elevations needed at these valves: {_ids_list(pvalves)}. Every junction has the same elevation "
                f"({next(iter(elev.values())):g} m), and a pressure-reducing, -sustaining or -breaker valve works on the "
                f"pressure, the head minus the elevation, so with one made-up elevation the valves set the wrong flows "
                f"and the water age is wrong." + _elevation_evidence() + " Export the model with its real elevations, "
                f"at least at the valves, the junctions downstream of them and the junctions named in pressure-based "
                f"controls.")
        elif zeros:
            rep.warnings.append(
                f"{len(zeros)} junction{'s' if len(zeros) != 1 else ''} at elevation 0 in a model with pressure-control "
                f"valves: {_ids_list(zeros, 10)}. If a 0 means 'not surveyed', the pressure there is wrong, and near a "
                f"valve or a pressure-based control so are the flows." + _elevation_evidence() + " Check these "
                f"elevations in the modelling package.")
    elif pnodes or emitters or pressure_driven:
        users = ([f"{len(pnodes)} junctions read by pressure-based controls"] if pnodes else []) + \
                ([f"{len(emitters)} emitters"] if emitters else []) + (["pressure-driven demands"] if pressure_driven
                                                                       else [])
        rep.add("Elevations", "no pressure-control valves; " + " and ".join(users) + " use the pressure, so the "
                              "elevations there change the flows; read from your file", FROM_FILE)
        if constant:
            rep.warnings.append("Every junction has the same elevation, and the model uses the pressure ("
                                + " and ".join(users) + "): the flows there may be wrong. Export the model with its "
                                "real elevations.")
    else:
        rep.add("Elevations", "no pressure-control valves, pressure-based controls or emitters: in this demand-driven "
                              "run the junction elevations change the pressures, not the flows", FROM_FILE)
    # 6. coordinates
    cs = coordinate_scale(wn, have_xy)
    f["coordinates"] = cs
    later = "tap mapping by coordinates, when it comes (planned), will not be offered for this file; use junction IDs"
    if cs["status"] == "no coordinates":
        rep.add("Coordinate scale", f"no coordinates: the maps will show every node at one point, and {later}", MISSING)
    elif cs["status"] == "too few pipes drawn":
        rep.add("Coordinate scale", f"too few pipes drawn to tell: {later}", DERIVED)
    elif cs["status"] == "schematic":
        rep.add("Coordinate scale", f"schematic: the drawn pipe lengths are not in proportion to the lengths in the "
                                    f"file (95th over 5th percentile of their ratio {cs['spread']:g}, above "
                                    f"{SCHEMATIC_SPREAD:g}), so {later}", DERIVED)
    else:
        unit = {"feet": "in feet", "metres": "in metres"}.get(cs["units"], "in units that are not feet or metres")
        rep.add("Coordinate scale", f"to scale, {unit} (median drawn length {cs['ratio_median']:g} per metre of "
                                    f"declared length; 95th over 5th percentile {cs['spread']:g})", DERIVED)
    # 7. sources in the file and entry points
    dosed = set(src_nodes)
    srcs = [(s.node_name, str(s.source_type)) for _, s in wn.sources()]
    over = [n for n, _ in srcs if n in dosed]
    boost = [n for n, _ in srcs if n not in dosed]
    f["sources_overridden"], f["booster_sources"], f["entry_points"] = over, boost, list(src_nodes)
    if not srcs:
        rep.add("Sources in the file", "none: the dose you set enters at every entry point", FROM_FILE)
    else:
        rep.add("Sources in the file", "; ".join(([f"at {_ids_list(over)}, an entry point: replaced by the dose you "
                                                   f"set"] if over else [])
                                                 + ([f"at {_ids_list(boost)}, not an entry point (a booster): kept in "
                                                     f"the file, but at the tool's dose scale it adds almost nothing"]
                                                    if boost else [])), FROM_FILE)
    if over:
        rep.warnings.append(f"The file's own source at {_ids_list(over)} is overridden: the tool doses every entry point "
                            f"with the dose you set.")
    if boost:
        rep.warnings.append(f"{len(boost)} source{'s' if len(boost) != 1 else ''} at node{'s' if len(boost) != 1 else ''} "
                            f"the tool does not dose ({_ids_list(boost)}): boosters, kept as the file has them. The tool "
                            f"writes its own dose to EPANET in the model's units, a thousand times the file's mg/L, and "
                            f"leaves these as they are, so they stay in the run at about a thousandth of their strength "
                            f"and add almost nothing. How to treat boosters is an open decision; until then the map "
                            f"has almost none of their chlorine.")
    if len(src_nodes) > 1:
        rep.warnings.append(f"{len(src_nodes)} entry points ({_ids_list(list(src_nodes))}): every one is dosed at the "
                            f"same value, the dose you set. If they carry different residuals, the map is off near "
                            f"them; a residual per inlet is planned.")
    # 8. structure: pumps and controls, demand and patterns, connected parts, export connectors
    loose = []
    for pn in wn.pump_name_list:
        p = wn.get_link(pn)
        if getattr(p, "speed_pattern_name", None) or pn in ctl["links"]:
            continue
        loose.append((pn, "closed" if p.initial_status == LinkStatus.Closed else "open"))
    f["pumps_without_controls"] = [pn for pn, _ in loose]
    f["controls"], f["rules"] = ctl["n_controls"], ctl["n_rules"]
    rep.add("Pumps and controls", f"{wn.num_pumps} pump{'s' if wn.num_pumps != 1 else ''}; {ctl['n_controls']} "
                                  f"control{'s' if ctl['n_controls'] != 1 else ''} and {ctl['n_rules']} "
                                  f"rule{'s' if ctl['n_rules'] != 1 else ''}", FROM_FILE)
    if loose:
        opened = [pn for pn, st in loose if st == "open"]
        closed = [pn for pn, st in loose if st == "closed"]
        desc = "; ".join(([f"open all the time: {_ids_list(opened)}"] if opened else [])
                         + ([f"closed all the time: {_ids_list(closed)}"] if closed else []))
        rep.warnings.append(f"Pumps but no controls: {len(loose)} of {wn.num_pumps} pump{'s' if wn.num_pumps != 1 else ''} "
                            f"{'has' if len(loose) == 1 else 'have'} no control, rule or speed pattern ({desc}), so "
                            f"{'it stays' if len(loose) == 1 else 'each stays'} as the file starts it for all 7 days. If "
                            f"the real pumps cycle, the water age and the tanks' chlorine are wrong: export the model "
                            f"with its controls.")
    total = sum(ts.base_value for j in junctions for ts in wn.get_node(j).demand_timeseries_list if ts.base_value > 0)
    f["patterns"], f["zero_total_demand"] = len(wn.pattern_name_list), total == 0
    rep.add("Demand", f"{len(wn.pattern_name_list)} time pattern{'s' if len(wn.pattern_name_list) != 1 else ''}"
                      + ("; no junction takes water" if total == 0 else ""), FROM_FILE)
    if not wn.pattern_name_list:
        rep.warnings.append("No time patterns: every demand is the same all day, so the daily low point of the residual "
                            "is not represented.")
    if total == 0:
        rep.warnings.append("Zero total demand: no junction takes water, so water moves only as the tanks fill and "
                            "drain, and its age is not the network's.")
    g = nx.Graph()
    g.add_nodes_from(wn.node_name_list)
    g.add_edges_from((link.start_node_name, link.end_node_name) for _, link in wn.links())
    parts = sorted((sorted(c) for c in nx.connected_components(g)), key=lambda c: (-len(c), c[0]))
    f["parts"] = [len(c) for c in parts]
    if len(parts) > 1:
        dry = [c for c in parts if not set(c) & (dosed | set(wn.tank_name_list) | set(wn.reservoir_name_list))]
        rep.add("Connected parts", f"{len(parts)} parts that no pipe, pump or valve joins (sizes "
                                   f"{', '.join(str(len(c)) for c in parts[:6])})", DERIVED)
        rep.warnings.append(f"{len(parts)} parts that no pipe, pump or valve joins (sizes "
                            f"{', '.join(str(len(c)) for c in parts[:6])}{' and more' if len(parts) > 6 else ''}): "
                            f"each part needs its own entry point"
                            + (f", and {len(dry)} part{'s have' if len(dry) != 1 else ' has'} none (for example "
                               f"{_ids_list(dry[0], 3)}), so no chlorine reaches it" if dry else "")
                            + ". If the zone was cut from a larger model, check the cut.")
    else:
        rep.add("Connected parts", "one connected network", DERIVED)
    conn = [pn for pn in wn.pipe_name_list
            if abs(wn.get_link(pn).length - CONNECTOR_LENGTH_M) < 0.01 and abs(wn.get_link(pn).diameter
                                                                                 - CONNECTOR_DIAM_M) < 0.01]
    f["connector_pipes"], f["connector_tags"] = len(conn), len(ctl["connector_tags"])
    if ctl["connector_tags"] or len(conn) >= CONNECTOR_MIN_PIPES:   # one pipe of that size alone says nothing (Net6)
        rep.add("Export connectors", f"{len(conn)} pipe{'s' if len(conn) != 1 else ''} of 1 ft and 99 in (the size a "
                                     f"WaterGEMS EPANET export uses for its connector pipes) and "
                                     f"{len(ctl['connector_tags'])} element{'s' if len(ctl['connector_tags']) != 1 else ''} "
                                     f"tagged AddedForINPConversion (the tag it gives them); kept as they are (for "
                                     f"information)", FROM_FILE)
        return set(conn) | (set(ctl["connector_tags"]) & set(wn.pipe_name_list))
    return set()


def _dw_report(rep: Readiness, conv: dict) -> None:
    """Rows, fix or block for a Darcy-Weisbach file and its gate (task 15b, item 3)."""
    from .simulate import roughness_factor
    keep = ("pipes", "c_min", "c_max", "age_change_median_h", "age_change_p95_h", "passes")
    rep.flags["headloss_conversion"] = {k: (round(conv[k], 4) if isinstance(conv[k], float) else conv[k]) for k in keep}
    med, p95 = conv["age_change_median_h"], conv["age_change_p95_h"]
    rep.add("Headloss conversion", f"water age moved by a median of {med:.2f} h and a 95th percentile of {p95:.2f} h "
                                   f"over the scored day's junction-hours between your file and the copy; the copy is "
                                   f"kept up to {DW_GATE_MEDIAN_H:g} h and {DW_GATE_P95_H:g} h", DERIVED)
    if conv["passes"]:
        rep.fixes.append(f"Converted the Darcy-Weisbach headloss formula to Hazen-Williams: each of the {conv['pipes']} "
                         f"pipes gets the C ({conv['c_min']:.0f} to {conv['c_max']:.0f}) that gives its Darcy-Weisbach "
                         f"headloss at its reference velocity (the median over the scored day, at least "
                         f"{V_REF_FLOOR_MS:g} m/s). The wall decay reads each pipe's roughness as a C, and a "
                         f"Darcy-Weisbach roughness (a fraction of a millimetre) read as a C near 0 would multiply every "
                         f"wall rate by about {roughness_factor(0.0):.0f} at the old-pipe factor 1. The conversion is "
                         f"exact at that velocity only, which is why the water age is compared.")
    else:
        rep.blocking.append(f"The Darcy-Weisbach model could not be converted closely enough: water age moved by a "
                            f"median of {med:.2f} h (at most {DW_GATE_MEDIAN_H:g} h is kept) and a 95th percentile of "
                            f"{p95:.2f} h (at most {DW_GATE_P95_H:g} h) between the file and a Hazen-Williams copy. Export "
                            f"the model with Hazen-Williams, or ask your modeller.")


def _timed_member_run(path: str, gamma: float) -> float:
    """Seconds for one 7-day EPANET run of the file exactly as a grid member runs it (simulate_nominal_chlorine at
    kb 0.40, kw 0.70), in a temporary folder that is removed; it raises when EPANET cannot run it."""
    from .simulate import simulate_nominal_chlorine
    with tempfile.TemporaryDirectory(prefix="rm_test_run_") as d, warnings.catch_warnings():
        warnings.simplefilter("ignore")
        t0 = time.perf_counter()
        simulate_nominal_chlorine(path, 0.40, 0.70, gamma, file_prefix=os.path.join(d, "member"))
        return time.perf_counter() - t0


def _read_and_report(read_path: str, out_path: str, rep: Readiness, ids: IdScan, disinfectant: str,
                     text: str) -> tuple[bool, str | None]:
    """WNTR's load of the prepared copy, the report rows that need it, the physics checks (task 15b), EPANET's test run
    of the first hour, the Darcy-Weisbach conversion and its gate, and one timed 7-day run for the run cost.
    (accepted, replacement): accepted is False (with a blocking problem) when WNTR cannot read it, the model has no
    entry point, a physics check blocks or a run fails; replacement is the converted text of a Darcy-Weisbach model,
    which _prepare saves instead of the copy read here."""
    import wntr
    from .simgp import GRIDS
    from .simulate import source_nodes
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wn = wntr.network.WaterNetworkModel(read_path)
    except Exception as e:  # noqa: BLE001
        msg = str(e).replace(read_path, os.path.basename(out_path))     # the file's name, never its folder
        rep.blocking.append(f"EPANET (through WNTR) could not read this model: {type(e).__name__}: {msg[:300]}")
        return False, None
    for item, n in (("Junctions", wn.num_junctions), ("Reservoirs", wn.num_reservoirs), ("Tanks", wn.num_tanks),
                    ("Pipes", wn.num_pipes), ("Pumps", wn.num_pumps), ("Valves", wn.num_valves)):
        rep.add(item, n, FROM_FILE)
    rep.add("Flow units", wn.options.hydraulic.inpfile_units, FROM_FILE)
    n_nodes = wn.num_nodes
    have_xy = ids.coordinate_ids & set(wn.node_name_list)
    if have_xy:
        rep.add("Node coordinates", f"{len(have_xy)} of {n_nodes} nodes", FROM_FILE)
    else:
        rep.add("Node coordinates", "none: the maps need them (EPANET or the modelling package writes them in "
                                    "[COORDINATES])", MISSING)
    src_nodes = source_nodes(wn)
    if not src_nodes:
        rep.blocking.append("No reservoir, inflow junction or tank found: the model needs a place where water enters.")
        return False, None
    how_in = ("reservoirs" if wn.reservoir_name_list else
              "junctions with a negative demand (inflows)" if set(src_nodes) <= set(wn.junction_name_list) else "tanks")
    rep.add("Where water enters", f"{how_in}: {', '.join(src_nodes[:6])}{' and more' if len(src_nodes) > 6 else ''}",
            DERIVED)
    rep.add("Disinfectant", f"{DISINFECTANT_WORDS[disinfectant]} (your setting; every entry point gets the same dose)",
            ASSUMED)
    connectors = _physics_report(rep, wn, control_scan(text), have_xy, src_nodes)
    if rep.blocking:
        return False, None
    headloss = wn.options.hydraulic.headloss
    # the uniform-C test leaves out the export's connector pipes (the second review: a WaterGEMS-style export of a
    # one-C network has a C 130 connector of 1 ft at each pump, which is not a second C value of the network)
    scored = [pn for pn in wn.pipe_name_list if pn not in connectors] or list(wn.pipe_name_list)
    roughness = [float(wn.get_link(pn).roughness) for pn in scored]
    left_out = len(wn.pipe_name_list) - len(scored)
    n_junctions = wn.num_junctions
    problem = _test_run(wn)            # it changes wn's time and quality options: nothing below reads them
    if problem:
        rep.blocking.append(problem)
        return False, None
    replacement, final_path = None, read_path
    with tempfile.TemporaryDirectory(prefix="rm_test_run_") as work:
        if headloss == "D-W":
            try:
                replacement, conv = dw_to_hw(read_path, text, work)
            except Exception as e:  # noqa: BLE001
                rep.blocking.append(f"The Darcy-Weisbach model could not be converted ({type(e).__name__}: "
                                    f"{str(e)[:200]}). Export the model with Hazen-Williams, or ask your modeller.")
                return False, None
            _dw_report(rep, conv)
            if not conv["passes"]:
                return False, None
            final_path = conv["converted_path"]
            keep = set(scored)
            converted = sorted({round(float(v), 6) for k, v in conv["c"].items() if k in keep})
            rep.flags["converted_c_values"] = len(converted)
        # the uniform test reads the file's own roughness, before any conversion: a Darcy-Weisbach file with one
        # roughness on every pipe gets many converted C values, which differ by diameter and velocity, not by age.
        # It compares them to 6 significant digits, not 6 decimals (the second review: a Darcy-Weisbach roughness in
        # metres is about 1e-6 for plastic pipe, so 6 decimals made 0.0015 mm and 0.0024 mm one value)
        own = sorted({float(f"{v:.6g}") for v in roughness})
        uniform = len(own) == 1
        rep.flags["c_values"], rep.flags["uniform_c"] = len(own), uniform
        but = (f" (the export's {left_out} connector pipe{'s' if left_out != 1 else ''} left out)" if left_out else "")
        if uniform:
            what = (f"one C value: every pipe has C {own[0]:g}{but}, so the factor would scale every wall rate alike "
                    f"and cannot be told from the wall rate itself" if headloss != "D-W" else
                    f"one roughness value: every pipe has the Darcy-Weisbach roughness {own[0] * 1000:g} mm in your "
                    f"file{but}, so "
                    f"the C values of the converted copy ({converted[0]:g} to {converted[-1]:g}) differ only by each "
                    f"pipe's diameter and velocity, not by its age, and the factor would read them as age")
            if grid_for(disinfectant, True) == "full_fixed_gamma":
                fixed_runs = math.prod(len(a) for a in GRIDS["full_fixed_gamma"])
                held = f"it is held at 0 (the free-chlorine grid with the factor held at 0 has {fixed_runs} runs)"
            else:
                held = ("the chloramine grid still varies it (a chloramine grid with the factor held at 0 is not built "
                        "yet), so its calibrated value means nothing for this file and the app shows it as not "
                        "identifiable")
            rep.add("Old-pipe factor", f"not identifiable from {what}; {held}", DERIVED)
        elif headloss == "D-W" and converted:
            rep.add("Old-pipe factor", f"learned from the {len(converted)} C values the conversion gives "
                                       f"({converted[0]:g} to {converted[-1]:g}), from your file's {len(own)} "
                                       f"roughness values{but}", DERIVED)
        elif own:
            rep.add("Old-pipe factor", f"learned from your file's {len(own)} C values ({own[0]:g} to "
                                       f"{own[-1]:g}){but}", FROM_FILE)
        try:
            secs = _timed_member_run(final_path, 0.0 if uniform else 1.0)
        except Exception as e:  # noqa: BLE001
            msg = str(e).replace(work, "")
            rep.blocking.append(f"EPANET could not run the 7-day simulation that every grid member is ({type(e).__name__}: "
                                f"{msg[:200]}). Open the model in EPANET or the modelling package, fix what it reports, "
                                f"and export it again.")
            return False, None
    rep.run_cost = run_cost(n_junctions, disinfectant, uniform, secs)
    c = rep.run_cost
    est = c["estimate_s"]
    rep.add(RUN_COST_ROW, f"{c['runs_per_dose']} EPANET runs per dose (the {c['grid']} grid), once per file: about "
                          + (f"{max(1, round(est)):d} s" if est < 60 else f"{c['minutes']:g} min")
                          + f" here, estimated as {c['runs_per_dose']} times one 7-day run ({c['seconds_per_run']:.2f} s) "
                          f"over an assumed parallel speedup of {EFFECTIVE_SPEEDUP:g}. The real speedup grows with the "
                          f"model's size, because starting the workers takes a fixed time: a small model's grid takes "
                          f"longer than this estimate, a large one's less. A grid file of about {c['grid_file_mb']:g} MB "
                          f"and up to "
                          f"{c['temp_files_mb']:g} MB of temporary files while it runs", DERIVED)
    if c["large"]:
        rep.warnings.append(f"{LARGE_WARNING_START} {n_junctions} junctions, more than {LARGE_JUNCTIONS}. The grid is "
                            f"{c['runs_per_dose']} runs and a file of about {c['grid_file_mb']:g} MB (the run cost row "
                            f"has the time), so building it needs an explicit confirm (a checkbox in the app's readiness "
                            f"report, or --allow-large in the pilot). The model is not skeletonized: trimming branches would "
                            f"delete the dead ends, where residuals are lowest.")
    rep.add(TEST_RUN_ROW, "EPANET ran the first hour (hydraulics and water age, set up as the app's first run) "
                          "without an error, and one full 7-day run of the grid's kind; the grid's runs come next",
            DERIVED)
    return True, replacement


# ----------------------------------------------------------------------------- task 16: grab logs, tap maps, clock time
# A grab log arrives as the utility's lab or field system exports it: one datetime column or a date and a time, US or
# European dates, headers such as 'Total Chlorine (mg/L)' or 'Sample Point', semicolons and decimal commas, a Windows
# code page or a byte-order mark, '<0.02' and 'ND' for readings below detection, an Excel workbook with typed date and
# time cells whose numeric tap IDs lost their leading zeros, or 15-minute data from an online analyzer.  read_grab_log
# turns it into the pilot's rows (junction, model hour, reading, month, date, tap) and says what it did with every row.
LOG_WINDOW_MONTHS = 3                  # the app's date window: the last 3 calendar months that hold a reading, as the
                                       # pilot's --window counts months
ANALYZER_MIN_PER_DAY = 12              # an online analyzer: 12 or more readings a day ...
ANALYZER_MAX_STEP_MIN = 30             # ... a median step of at most 30 minutes between them (an analyzer day) ...
ANALYZER_MIN_DAYS = 2                  # ... on 2 or more days; or 12 or more a day on ANALYZER_DENSE_DAYS days or more
ANALYZER_DENSE_DAYS = 7                # (a historian's hourly export).  An hourly grab-sampling study of a day or two
                                       # is neither, so every one of its readings is kept (the review of task 16)
ANALYZER_HOUR = 10.0                   # an analyzer series is thinned to the reading nearest this clock hour, per day
ANALYZER_MAX_GAP_MIN = 60              # ... when one lies within this many minutes of it; otherwise that day has none
UTC_DAY_MIN = (6 * 60, 18 * 60)        # most grab samples outside 06:00 to 18:00 suggests times exported in UTC
DAY_START_H, DAY_END_H = 7, 17         # simgp.DAY_HOURS (07:00 to 17:00), the hours an operator can take a grab sample
LOG_MAX_BYTES = 100 * 1024 * 1024      # a grab log or tap map read whole up to this size
LOG_MAX_ROWS = 1_000_000
LOG_MAX_COLS = 200                     # an Excel sheet is read up to this many columns (a stray cell far right is cut)
LOG_LAYOUT_RECORDS = 50                # a text table's separator and header line are found in its first 50 records
READING_MAX_MGL = 10.0                 # a reading above this is no chlorine residual: listed, not used (check the unit)
READING_HIGH_MEDIAN_MGL = 5.0          # a median reading above this warns (the unit or the column may be wrong)
DROPPED_MAX_RUNS = 20                  # the not-used list gives at most this many runs of rows per reason, and counts
DROPPED_MAX_VALUES = 5                 # a reason that names a value ('not a number: '0.452 J'') becomes one line for
                                       # all its values when more than this many differ (the final verification)
DROPPED_MAX_LINES = 30                 # and the list gives at most this many lines, the last counting the rest
LOG_MAX_LINES = 2 * LOG_MAX_ROWS + LOG_LAYOUT_RECORDS   # line breaks counted before a text is split (quoted
                                       # comments over several lines leave room); more is refused unread
DATE_ORDERS = {"MDY": "month first (MM/DD/YYYY)", "DMY": "day first (DD/MM/YYYY)", "YMD": "year first (YYYY-MM-DD)"}
_BLANK_READINGS = {"", "-", "--", "na", "n/a", "nan", "none", "null", "#n/a"}
_ND_WORDS = {"nd", "n.d.", "n/d", "non-detect", "nondetect", "non detect", "not detected", "bdl", "below detection",
             "<dl", "<mdl", "<rl", "<lod", "<mrl"}
# a reading column's header names chlorine: a word containing 'chlorine', or one of these words; never a header with
# one of the _NOT_CHLORINE words (chloride, chlorate and chlorite contain 'chlor' but are other analytes)
_READING_WORDS = ("cl2", "chlor", "chloramine", "chloramines", "residual")
_NOT_CHLORINE = {"chloride", "chlorides", "chlorate", "chlorates", "chlorite", "chlorites", "perchlorate", "chlorophyll",
                 "chloroform", "dioxide", "clo2", "monochloramine", "dichloramine", "trichloramine", "combined"}
_NOT_READING_WORDS = {"method", "methods", "note", "notes", "unit", "units", "qualifier", "qual", "flag", "flags",
                      "comment", "comments", "remark", "remarks", "dose", "date", "time", "analyst", "instrument",
                      "id", "code", "type", "demand", "pressure", "psi", "temperature", "temp"}
# a detection or reporting limit is no reading, even when its header names chlorine ('Free Chlorine MDL', 'Free
# Chlorine (RL)'): looked for with the bracket contents kept (the final verification of task 16)
_LIMIT_WORDS = {"mdl", "mrl", "rl", "lod", "loq", "pql", "dl", "limit", "limits", "detection"}
# the species, read with the bracket contents kept ('Chlorine (Free)', 'Cl2 (total)'; the final verification):
# chloramine is measured as total chlorine; 'combined' is total minus free, never read as a reading (_NOT_CHLORINE)
# but never free either
_FREE_WORDS = {"free", "fac", "frc", "fcl2", "cl2f"}                # free available, free residual; LIMS codes
_TOTAL_WORDS = {"total", "trc", "tcl2", "cl2t", "chloramine", "chloramines", "combined"}
_CHLORINE_CODES = {"fcl2", "cl2f", "tcl2", "cl2t"}  # LIMS codes that name chlorine and its species by themselves
_UNIT_WORDS = {"mg", "l", "mgl", "ug", "µg", "ugl", "µgl", "mcg", "mcgl", "ppm", "ppb", "as", "per", "liter", "litre",
               "1"}                             # left out of the species words ('mg/L', 'ug per L', 'mg L-1')
# a long-format export (one row per analyte): the column that names each row's analyte, and the one with its unit
_ANALYTE_WORDS = {"analyte", "analytes", "parameter", "parameters", "param", "determinand", "characteristic",
                  "constituent", "analysis", "test", "component", "measurand", "compound"}
_ANALYTE_NOT = {"date", "time", "result", "results", "value", "values", "unit", "units", "id", "code", "method"}
_UNIT_EXACT = ("unit", "units", "uom", "unit of measure", "units of measure", "result unit", "result units",
               "reporting unit", "reporting units")
_RESULT_EXACT = ("result", "results", "value", "result value", "reported value", "final result", "measurement",
                 "concentration", "conc", "reading", "numeric result", "result numeric", "result num")
# units, on the text as _unit_text writes it (lower case, a camel-case join split, µ written one way): 'ug/L',
# 'ugL', '_ugl', 'µg L-1', 'ug per L', 'mcg/L', 'micrograms per litre', 'ppb' are micro; 'mg/L', 'mgL', 'mg Cl2/L',
# 'mgCl2/L', 'mg/litre', 'ppm' are milli (the final verification of task 16: 'ugL' and 'mg Cl2/L' were not read)
_UNIT_TAIL = r"[\s_./]*(?:cl2[\s_./]*)?(?:per[\s_]*)?(?:l|lit(?:er|re)s?)(?:[\s_]*-\s*1)?(?![a-z])"
_MICRO = re.compile(r"(?<![a-z])(?:u|µ|mc|micro)g(?:rams?)?" + _UNIT_TAIL + r"|(?<![a-z])ppb(?![a-z])")
_MILLI = re.compile(r"(?<![a-z])(?:mg|milligrams?)" + _UNIT_TAIL + r"|(?<![a-z])ppm(?![a-z])")
_CAMEL = re.compile(r"(?<=[a-z0-9µ])(?=[A-Z])")
_MOJIBAKE_MU = (("Âµ", "µ"), ("Î¼", "µ"), ("μ", "µ"))     # µ in UTF-8 read back as cp1252, and the Greek mu
_SEP_HINT = re.compile(r"^\s*sep\s*=\s*(.)\s*$", re.I)
_MONTHS = {m: i for i, ms in enumerate((("jan", "january"), ("feb", "february"), ("mar", "march"), ("apr", "april"),
                                        ("may",), ("jun", "june"), ("jul", "july"), ("aug", "august"),
                                        ("sep", "sept", "september"), ("oct", "october"), ("nov", "november"),
                                        ("dec", "december")), start=1) for m in ms}
_TAP_EXACT = ("tap id", "tap", "tapid", "tap name", "sample point", "sample point id", "sampling point",
              "sampling point id", "sample site", "sample site id", "site", "site id", "site name", "sample location",
              "location", "location id", "station", "station id", "point id", "sample tap", "sp id", "loc id")
_TAP_WORDS = ("tap", "point", "site", "location", "station", "loc")
_JUNCTION_EXACT = ("junction id", "junction", "junctionid", "node id", "node", "nodeid", "model node", "model node id",
                   "model id", "model junction", "epanet id", "epanet node")
_ISO_DATE = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:(?:[T ]|,\s*)+(.*))?$")
_NUM_DATE = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4}|\d{2})(?:(?:[T ]|,\s*)+(.*))?$")
# month names (an Excel or Oracle export's '10-Mar-26', 'Mar 10, 2026'): the name settles the order
_MON_DATE = re.compile(r"^(\d{1,2})[-/. ]([A-Za-z]{3,9})\.?[-/. ](\d{4}|\d{2})(?:(?:[T ]|,\s*)+(.*))?$")
_MON_FIRST = re.compile(r"^([A-Za-z]{3,9})\.?[- ](\d{1,2})(?:st|nd|rd|th)?,?[- ](\d{4})(?:(?:[T ]|,\s*)+(.*))?$")
# a time: '9:15', '09:15:30', '9:15 AM', or with dots, as European and Oracle exports write it ('09.15', '09.15.00 AM')
_CLOCK = re.compile(r"^(\d{1,2})[:.](\d{2})(?:[:.](\d{2})(?:[.,]\d+)?)?\s*(?:([AaPp])\.?\s*[Mm]\.?)?$")
_HHMM = re.compile(r"^(\d{1,2})(\d{2})$")
_ZONE = re.compile(r"(?:\s*[Zz]|\s*(?:UTC|GMT)|\s*[+\-]\d{2}:?\d{2})$")


class LogRefused(ValueError):
    """A grab log or tap map that cannot be read as it is.  The message says why in plain words, and asks a question
    when the file leaves one open (a date that reads two ways); nothing is guessed in its place."""


def have_openpyxl() -> bool:
    """openpyxl (task 16: Excel grab logs and the .xlsx template) is in this environment."""
    import importlib.util
    return importlib.util.find_spec("openpyxl") is not None


if not have_openpyxl():         # task 16: without openpyxl an Excel grab log is refused, and the sniff says so
    KINDS["xlsx"] = (KINDS["xlsx"][0], KINDS["xlsx"][1],
                     "Excel workbooks need openpyxl, which is not installed here: save the sheet as CSV (File, Save As, "
                     "CSV UTF-8) and upload that. It is not a model.")


@dataclass
class LogTable:
    header: list                       # the header row's names, made unique ('name (2)' for a repeat)
    rows: list                         # data rows: lists of cells (str, or Excel's typed date, time and number cells)
    row_numbers: list                  # each data row's line in the file (or row of the sheet), 1-based
    how: str                           # how it was read, in plain words


def _unique(names: list[str]) -> list[str]:
    seen, out = {}, []
    for n in names:
        n = n.strip() or "column"
        seen[n] = seen.get(n, 0) + 1
        out.append(n if seen[n] == 1 else f"{n} ({seen[n]})")
    return out


def _cell_text(v) -> str:
    """A cell as text: an integral number without '.0' (Excel keeps tap 7 as 7.0 in some exports)."""
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).strip()


def _iter_lines(text: str):
    """The text's lines, as _LINE_BREAK.split(text) gives them, one at a time (no list of every line is built: the
    final verification of task 16 measured about 0.3 GB per million short rows before a long file was refused)."""
    pos = 0
    for m in _LINE_BREAK.finditer(text):
        yield text[pos:m.start()]
        pos = m.end()
    yield text[pos:]


def _records(lines, d: str, limit: int | None = None) -> list[tuple[int, list]]:
    """(line, cells) of each record of the text that is not blank, read with delimiter d: line is the record's first
    line in the file, 1-based (csv.reader's line_num, so a quoted field that runs over several lines, a LIMS comment,
    does not shift the rows after it; the review of task 16).  lines is any iterable of lines (_iter_lines).  Stops
    after `limit` records."""
    out, prev = [], 0
    rd = csv.reader(lines, delimiter=d)
    for row in rd:
        start, prev = prev + 1, rd.line_num
        if not any(c.strip() for c in row):
            continue
        out.append((start, row))
        if limit and len(out) >= limit:
            break
    return out


def _table_layout(text: str) -> tuple[str, int, str | None] | None:
    """(delimiter, how many records come before the header, a note) of a delimited text (',', ';', tab or '|').
    A first line 'sep=;' (Excel's hint, common in European exports) names the delimiter and is skipped.  Otherwise,
    first as before task 16's review: the delimiter that splits the first record into two or more cells and gives at
    least 80% of the first 25 records its count, csv.Sniffer only ordering the candidates (on a sample with its quote
    characters taken out, as _sniff_table does); every file read before is read the same way.  When no delimiter
    passes, the file may open with a title or a 'Lab:/Exported' block: each delimiter's most common cell count over the
    first LOG_LAYOUT_RECORDS records (two or more cells) is tried, the header being the first record with that count
    and two filled cells, and at least 80% of the records from it on must have that count; the delimiter that skips
    the fewest lines wins (then the larger count, then csv.Sniffer's order).  Since the final verification, a delimiter
    that passes the first test gives way to one whose most common count is larger, when the records up to that one's
    header do not all have the first delimiter's count: a title with a comma ('Residuals report, March 2026') before a
    semicolon export with decimal commas splits the title and every row into 2 cells at a comma, which passed, while
    the header line, with no comma, did not.  None when nothing fits.  The text is read line by line, never split
    into a list of every line."""
    from collections import Counter
    from itertools import islice
    nonblank = lambda: (ln for ln in _iter_lines(text) if ln.strip())       # noqa: E731
    first_line = next(nonblank(), "")
    m = _SEP_HINT.match(first_line)
    if m and m.group(1) in _DELIMS:
        return m.group(1), 1, f"its first line, {first_line.strip()!r}, names the separator"
    sample = "\n".join(islice(nonblank(), 25))
    order = list(_DELIMS)
    try:
        sniffed = csv.Sniffer().sniff(_sniffer_sample(sample), delimiters="".join(_DELIMS)).delimiter
        order = [sniffed] + [x for x in _DELIMS if x != sniffed]
    except csv.Error:
        pass
    recs, fast = {}, None
    for d in order:
        try:
            recs[d] = _records(_iter_lines(text), d, LOG_LAYOUT_RECORDS)
        except csv.Error:
            continue
        rows = [r for _, r in recs[d][:25]]
        if rows and len(rows[0]) >= 2 and sum(len(r) == len(rows[0]) for r in rows) / len(rows) >= 0.8:
            fast = d
            break
    modal = {}
    for rank, d in enumerate(order):
        if d not in recs and fast is not None:
            try:
                recs[d] = _records(_iter_lines(text), d, LOG_LAYOUT_RECORDS)
            except csv.Error:
                continue
        rows = [r for _, r in recs.get(d, [])]
        counts = [len(r) for r in rows]
        if len(counts) < 2:
            continue
        n = max(Counter(counts).items(), key=lambda kv: (kv[1], kv[0]))[0]
        k = next((i for i, r in enumerate(rows) if len(r) == n and sum(c.strip() != "" for c in r) >= 2), None)
        if n < 2 or k is None or len(counts) - k < 2:
            continue
        if sum(x == n for x in counts[k:]) / (len(counts) - k) < 0.8:
            continue
        modal[d] = (k, n, rank)
    if fast is not None:
        n0 = len(recs[fast][0][1])
        head = [len(r) for _, r in recs[fast]]
        better = [(k, -n, rank, d) for d, (k, n, rank) in modal.items()
                  if d != fast and n > n0 and any(x != n0 for x in head[:k + 1])]
        if not better:
            return fast, 0, None
        k, _, _, d = min(better)
    elif modal:
        k, _, _, d = min((k, -n, rank, d) for d, (k, n, rank) in modal.items())
    else:
        return None
    if k == 0:
        return d, 0, None
    want = {n for n, _ in recs[d][:k][:3]}
    texts = {i: ln for i, ln in enumerate(islice(_iter_lines(text), max(want)), start=1) if i in want}
    skipped = "; ".join(texts[n].strip()[:60] for n, _ in recs[d][:k][:3])
    return d, k, (f"the {k} line{'s' if k != 1 else ''} before its header {'were' if k != 1 else 'was'} skipped "
                  f"({skipped!r}{', ...' if k > 3 else ''})")


def _xlsx_cells(data: bytes) -> tuple[list, list, str]:
    """The rows of a workbook's first sheet that holds a table (a cover sheet before it is passed over), through
    openpyxl in read-only mode: at most LOG_MAX_COLS columns and LOG_MAX_ROWS + LOG_LAYOUT_RECORDS rows are read, so a
    small file with one stray cell far down or far right (which pandas.read_excel padded to rows x columns cells) is
    cut, not unpacked.  Returns (rows, each row's number in the sheet, how)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        sheets = list(wb.worksheets)
        for k, ws in enumerate(sheets):
            cells, numbers, empty, wide = [], [], None, False
            for i, row in enumerate(ws.iter_rows(max_col=LOG_MAX_COLS, max_row=LOG_MAX_ROWS + LOG_LAYOUT_RECORDS,
                                                 values_only=True), start=1):
                if row is empty:                     # openpyxl yields one shared tuple for each missing row
                    continue
                r = list(row)
                while r and (r[-1] is None or (isinstance(r[-1], str) and not r[-1].strip())):
                    r.pop()
                if not r:
                    empty = row
                    continue
                wide = wide or len(r) >= LOG_MAX_COLS
                cells.append(r)
                numbers.append(i)
            # a table: two or more rows with two filled cells (a cover sheet's 'Report | March 2026' is not one)
            if sum(sum(_cell_text(c) != "" for c in r) >= 2 for r in cells[:LOG_LAYOUT_RECORDS]) >= 2:
                how = ("an Excel workbook (.xlsx), its first sheet, typed cells kept" if k == 0 else
                       f"an Excel workbook (.xlsx), its sheet {ws.title!r} (the first that holds a table), typed "
                       f"cells kept")
                if wide:
                    how += f", read up to column {LOG_MAX_COLS}"
                return cells, numbers, how
        return [], [], "an Excel workbook (.xlsx)"
    finally:
        wb.close()


def read_table(data: bytes, name: str = "") -> LogTable:
    """A grab log or tap map as a table: delimited text in any encoding decode_text reads (a byte-order mark, UTF-16, a
    Windows code page), separated by commas, semicolons, tabs or vertical bars (found by _table_layout, which also
    passes over a title or 'Lab:/Exported' lines before the header and Excel's 'sep=;' line), or an Excel workbook
    (.xlsx, its first sheet that holds a table, through openpyxl with the row and column caps of _xlsx_cells, typed
    cells kept).  The header is the first row with two or more filled cells (for text, the header _table_layout
    finds).  Each row's number is its line in the file (or its row in the sheet).  Raises LogRefused for anything else
    (a legacy .xls, a PDF, a model), with the sniff's advice."""
    if len(data) > LOG_MAX_BYTES:
        raise LogRefused(f"The file is {len(data) / 1e6:.0f} MB, more than the {LOG_MAX_BYTES / 1e6:.0f} MB a log is read "
                         f"up to. Send the last few months only.")
    s = sniff(data, name)
    skip = 0
    if s.kind == "xlsx":
        if not have_openpyxl():
            raise LogRefused("This is an Excel workbook, and openpyxl, which reads one, is not installed here. Save "
                             "the sheet as CSV (File, Save As, CSV UTF-8) and upload that.")
        try:
            zip_listing(data)                      # the size limits of a zip, before openpyxl unpacks anything
            cells, numbers, how = _xlsx_cells(data)
        except ZipRefused as e:
            raise LogRefused(f"The workbook is refused: {e}.") from None
        except Exception as e:  # noqa: BLE001  (a damaged workbook)
            raise LogRefused(f"The workbook could not be read ({type(e).__name__}). Save the sheet as CSV and upload "
                             f"that.") from None
    elif s.category == "table" or s.kind in ("text_other", "empty"):
        if s.kind == "xls_legacy":
            raise LogRefused("This is a legacy Excel workbook (.xls), which is not read here. Save it as .xlsx or as "
                             "CSV and upload that.")
        text, enc = decode_text(data)
        # the final verification of task 16: the lines are counted before anything is split or built, and records
        # are read only up to the row limit, so a small file of very short rows is refused without its memory
        n_lines = text.count("\n") + text.count("\r") - text.count("\r\n") + 1
        if n_lines > LOG_MAX_LINES:
            raise LogRefused(f"The file has {n_lines:,} lines, more than the {LOG_MAX_ROWS:,} rows a log is read up to. "
                             f"Send the last few months only.")
        if sum(1 for _ in zip(range(2), (ln for ln in _iter_lines(text) if ln.strip()))) < 2:
            raise LogRefused("The file has fewer than two lines: a header row and at least one reading are needed.")
        try:
            layout = _table_layout(text)
        except csv.Error as e:
            raise LogRefused(f"The file could not be read as a table ({e}). Save it again as CSV.") from None
        if layout is None:
            raise LogRefused("The file is text, but no separator (comma, semicolon, tab or vertical bar) splits it into "
                             "the same columns on every line. Save it again as CSV.")
        d, skip, layout_note = layout
        limit = skip + LOG_MAX_ROWS + LOG_LAYOUT_RECORDS + 1
        try:
            recs = _records(_iter_lines(text), d, limit)
        except csv.Error as e:
            raise LogRefused(f"The file could not be read as a table ({e}). Save it again as CSV.") from None
        if len(recs) >= limit:
            raise LogRefused(f"The file has more than {LOG_MAX_ROWS:,} rows. Send the last few months only.")
        numbers, cells = [n for n, _ in recs[skip:]], [r for _, r in recs[skip:]]
        how = f"text separated by {_DELIMS[d]}" + ("" if enc == "utf-8" else f", decoded as {enc}")
        if layout_note:
            how += f"; {layout_note}"
    else:
        raise LogRefused(f"{s.sentence()} A grab log or tap map is a table: CSV, or an Excel workbook (.xlsx). "
                         f"{s.advice}")
    k = next((i for i, r in enumerate(cells) if sum(_cell_text(c) != "" for c in r) >= 2), None)
    if s.kind == "xlsx" and k is not None:
        # the final verification of task 16: a title row with two filled cells ('Residuals report | March 2026') is
        # no header: the header is the first row with more filled cells than half the most common count of the
        # sheet's first rows (two or more filled cells each), and the rows before it are named in the notes
        from collections import Counter
        filled = [sum(_cell_text(c) != "" for c in r) for r in cells[:LOG_LAYOUT_RECORDS]]
        two = [x for x in filled if x >= 2]
        n = max(Counter(two).items(), key=lambda kv: (kv[1], kv[0]))[0] if two else 2
        k = next((i for i, x in enumerate(filled) if x >= 2 and x > n / 2), k)
        if k:
            shown = "; ".join(" | ".join(_cell_text(c) for c in r if _cell_text(c) != "")[:60] for r in cells[:k][:3])
            how += (f"; the {k} row{'s' if k != 1 else ''} before its header {'were' if k != 1 else 'was'} skipped "
                    f"({shown!r}{', ...' if k > 3 else ''})")
    if k is None:
        raise LogRefused("No row of the file has two or more filled cells, so no header row was found.")
    header = _unique([_cell_text(c).replace("﻿", "") for c in cells[k]])
    body, nums = [], []
    for r, n in zip(cells[k + 1:], numbers[k + 1:]):
        if all(_cell_text(c) == "" for c in r):
            continue
        r = list(r)[:len(header)] + [None] * max(0, len(header) - len(r))
        body.append(r)
        nums.append(n)
        if len(body) > LOG_MAX_ROWS:
            raise LogRefused(f"The file has more than {LOG_MAX_ROWS:,} rows. Send the last few months only.")
    return LogTable(header=header, rows=body, row_numbers=nums, how=how)


def _norm_header(h: str) -> str:
    """'Total Chlorine (mg/L)' -> 'total chlorine'; 'tap_id' -> 'tap id'; 'Sample Date/Time' -> 'sample date time'."""
    s = re.sub(r"\(.*?\)|\[.*?\]", " ", str(h).replace("﻿", ""))
    return re.sub(r"[^0-9a-z]+", " ", s.lower()).strip()


def _words(h, keep_brackets: bool) -> list[str]:
    """The words of a header or a cell, in lower case, a camel-case join split ('FreeCl' -> free, cl), with the
    contents of brackets kept or left out."""
    s = str(h).replace("\ufeff", "")
    if not keep_brackets:
        s = re.sub(r"\(.*?\)|\[.*?\]", " ", s)
    return re.sub(r"[^0-9a-zµ]+", " ", _CAMEL.sub(" ", s).lower()).split()


def _species_words(h) -> list[str]:
    """The words a species is read from: the bracket contents kept ('Chlorine (Free)', 'Cl2 (total)'), the unit
    words left out (the final verification of task 16: brackets were stripped first, so their species was lost)."""
    return [w for w in _words(h, True) if w not in _UNIT_WORDS]


def reading_species(h: str) -> str | None:
    """'free' or 'total' when a reading column's header (or a long-format row's analyte) names the species, None when
    it does not (a plain 'Chlorine Residual') or names both.  Read with the bracket contents kept, so 'Chlorine
    (Free)' is free and 'Cl2 (total)' total; 'Chloramine' is total chlorine (what a chloraminated system measures);
    the LIMS codes FCL2, CL2-F, TCL2 and CL2-T name their species."""
    w = _species_words(h)
    ws = set(w) | ({a + b for a, b in zip(w, w[1:])} & _CHLORINE_CODES)
    free, total = bool(ws & _FREE_WORDS), bool(ws & _TOTAL_WORDS)
    if free != total:
        return "free" if free else "total"
    return None


def _names_chlorine(words: list[str]) -> bool:
    """The words name chlorine: a word containing 'chlorine' ('free_chlorine_mgL', 'FreeChlorine'), or 'cl2', 'chlor',
    'chloramine', 'residual' or a LIMS code (FCL2, TCL2) as a whole word, and none of chloride, chlorate, chlorite,
    chlorophyll, dioxide and the like (task 16's review: 'chlor' matched inside 'Chloride', which was then fitted as
    chlorine)."""
    return (any("chlorine" in x or x in _READING_WORDS or x in _CHLORINE_CODES for x in words)
            and not (_NOT_CHLORINE & set(words))
            and not any(x.startswith(("chlorid", "chlorat", "chlorit")) for x in words))


def names_chlorine(h) -> bool:
    """A header or an analyte names chlorine: _names_chlorine on its words outside brackets, or 'Cl' with a species
    ('Free Cl', 'Cl (total)'; 'Cl' alone may be chloride); never when a word of it, brackets included, is chloride,
    combined or the like ('Chlorine (Combined)')."""
    ws = _species_words(h)
    if _NOT_CHLORINE & set(ws) or any(x.startswith(("chlorid", "chlorat", "chlorit")) for x in ws):
        return False
    return _names_chlorine(_norm_header(h).split()) or ("cl" in ws and reading_species(h) is not None)


def is_reading_header(h: str) -> bool:
    """A column of chlorine readings, by its header: it names chlorine (names_chlorine), and is no method, unit,
    note, pressure or temperature column (words outside brackets) and no detection or reporting limit ('Free Chlorine
    MDL', 'Free Chlorine (RL)', words inside brackets too)."""
    return (names_chlorine(h) and not (_NOT_READING_WORDS & set(_norm_header(h).split()))
            and not (_LIMIT_WORDS & set(_species_words(h))))


def _unit_text(s) -> str:
    s = str(s or "")
    for a, b in _MOJIBAKE_MU:
        s = s.replace(a, b)
    return _CAMEL.sub(" ", s).lower()


def header_unit(h: str) -> float | None:
    """The factor to mg/L of a unit a header (or a units cell) names: 1 for mg/L or ppm, 0.001 for µg/L, ug/L, mcg/L
    or ppb, None when it names neither (a header with no unit is read as mg/L).  Since the final verification of task
    16 the unit may be written without a slash or with Cl2 ('ugL', 'free_chlorine_ugl', 'µg L-1', 'ug per L', 'mg
    Cl2/L', 'mgCl2/L', 'mg/litre'), and a µ that an Excel round trip turned into 'Âµ' is read as µ."""
    s = _unit_text(h)
    if _MICRO.search(s):
        return 0.001
    if _MILLI.search(s):
        return 1.0
    return None


def analyte_kind(v) -> str:
    """What a long-format row's analyte names: 'free' or 'total' chlorine, 'chlorine' with no species ('Chlorine
    Residual'), or 'other' (pH, temperature, chloride, monochloramine).  Since the final verification of task 16 the
    species is read with the bracket contents kept ('Chlorine (Free)'), and 'Free Cl', 'Total Cl' and the codes FCL2
    and TCL2 name chlorine."""
    s = _cell_text(v)
    if not names_chlorine(s) or {"pressure", "psi", "temperature", "temp"} & set(_species_words(s)):
        return "other"
    return reading_species(s) or "chlorine"


def _guess(header: list[str], exact: tuple, words: tuple = ()) -> str | None:
    norm = {h: _norm_header(h) for h in header}
    for e in exact:
        for h in header:
            if norm[h] == e:
                return h
    for h in header:
        w = norm[h].split()
        if any(x in w for x in words) and not {"date", "time", "chlorine", "cl2"} & set(w):
            return h
    return None


def guess_log_columns(header: list[str], disinfectant: str = "free_chlorine") -> dict:
    """The columns of a grab log, guessed from its header: {'tap', 'reading', 'date', 'time'} (header names, or None)
    and the reading columns by species ('free_columns', 'total_columns', 'other_columns'; 'reading_columns' all of
    them).  'both_species' is True when the reading column must be chosen: the header has two or more reading columns
    and they are not all of one named species (a free and a total column, or a column whose header names no species
    beside another reading column; since the final verification of task 16, and the species is read inside brackets
    too, 'Chlorine (Free)').  The reading guess is the first reading column of the disinfectant's species, else the
    first whose header names no species, a header naming chlorine before one that says only 'residual'.  'time' is None when there is no separate time column (one datetime column, or a date column
    that holds the time).  For a long-format export (one row per analyte, the review of task 16), 'analyte' is the
    column that names each row's analyte ('Analyte', 'Parameter', 'Test'), 'unit' its units column, and the reading
    guess, when no header names chlorine, its result column ('Result', 'Value')."""
    want = "total" if disinfectant == "chloramine" else "free"
    norm = {h: _norm_header(h).split() for h in header}
    # a header that names chlorine comes before one that says only 'residual' (task 16's review)
    readings = sorted((h for h in header if is_reading_header(h)),
                      key=lambda h: not any("chlorine" in x or x in ("cl", "cl2", "chlor", "chloramine", "chloramines")
                                            or x in _CHLORINE_CODES for x in norm[h]))
    by = {"free": [h for h in readings if reading_species(h) == "free"],
          "total": [h for h in readings if reading_species(h) == "total"],
          None: [h for h in readings if reading_species(h) is None]}
    other = "free" if want == "total" else "total"
    reading = (by[want] or by[None] or by[other] or [None])[0]
    # a long-format export (one row per analyte): the analyte column, its units column, and its result column
    analyte = next((h for h in header if set(norm[h]) & _ANALYTE_WORDS
                    and not set(norm[h]) & (_ANALYTE_NOT | set(_TAP_WORDS)) and h != reading), None)
    unit = next((h for h in header if " ".join(norm[h]) in _UNIT_EXACT), None)
    if reading is None and analyte is not None:
        reading = next((h for e in _RESULT_EXACT for h in header if " ".join(norm[h]) == e), None)
    both = [h for h in header if "date" in norm[h] and "time" in norm[h] or "datetime" in norm[h]
            or "timestamp" in norm[h]]
    dates = [h for h in header if "date" in norm[h] and "time" not in norm[h]
             or norm[h] in (["collected"], ["sampled"], ["collection"])]
    times = [h for h in header if "time" in norm[h] and "date" not in norm[h] and "zone" not in norm[h]]
    if dates and times:
        date, time_ = dates[0], times[0]
    elif both:
        date, time_ = both[0], None
    else:
        date, time_ = (dates or [None])[0], (times or [None])[0] if dates else None
    return {"tap": _guess(header, _TAP_EXACT, _TAP_WORDS), "reading": reading, "date": date, "time": time_,
            "analyte": analyte, "unit": unit, "free_columns": by["free"], "total_columns": by["total"],
            "other_columns": by[None], "reading_columns": readings,
            "both_species": len(readings) >= 2 and (bool(by[None]) or bool(by["free"] and by["total"]))}


def guess_tap_columns(header: list[str]) -> dict:
    """The tap map's columns: {'tap', 'junction'}.  With two columns: when neither header is recognised, the first is
    the tap and the second the junction; when one is, the other column takes the missing role, except that a column
    that names a junction or a node is never the tap (the final verification of task 16: 'Node ID,Elevation', a
    junction export, was offered as a tap map whose junction column was 'Elevation')."""
    tap = _guess(header, _TAP_EXACT, _TAP_WORDS)
    junction = _guess([h for h in header if h != tap], _JUNCTION_EXACT, ("junction", "node"))
    if len(header) == 2:
        if tap is None and junction is None:
            tap, junction = header
        elif junction is None:
            junction = next(h for h in header if h != tap)
    return {"tap": tap, "junction": junction}


def parse_reading(v) -> tuple[float | None, str | None]:
    """(value, None) for a reading; (None, reason) when it is not used.  A reason starting 'censored' marks a reading
    below detection ('<0.02', 'ND', 'BDL'), and one starting 'above' a reading above the measuring range ('>2.2');
    both are kept and listed, never fitted.  Decimal commas ('0,45') and a trailing 'mg/L' or 'ppm' are read."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if isinstance(v, float) and math.isnan(v):
            return None, "blank reading"
        return (float(v), None) if v >= 0 else (None, f"negative reading: {v!r}")
    s = _cell_text(v)
    low = re.sub(r"\s+", " ", s.lower()).strip()
    if low in _BLANK_READINGS:
        return None, "blank reading"
    if low in _ND_WORDS:
        return None, f"censored, not detected ({s!r}): kept and listed, not used in the fit"
    m = re.fullmatch(r"([<>])\s*=?\s*(\d+(?:[.,]\d+)?)\s*(?:mg\s*/?\s*l|ppm)?", low)
    if m:
        if m.group(1) == "<":
            return None, f"censored, below detection ({s!r}): kept and listed, not used in the fit"
        return None, f"above the measuring range ({s!r}): not used in the fit"
    m = re.fullmatch(r"(-?\d+(?:[.,]\d+)?|-?[.,]\d+)\s*(?:mg\s*/?\s*l|ppm)?", low)
    if not m:
        return None, f"not a number: {s!r}"
    x = float(m.group(1).replace(",", "."))
    if x < 0:
        return None, f"negative reading: {s!r}"
    return x, None


def _clock_text(s: str) -> int | None:
    """Minutes of the day of a time written as text ('9:15', '09:15:30', '9:15 AM', '0915'), seconds truncated."""
    s = s.strip()
    m = _CLOCK.fullmatch(s)
    if m:
        h, mi = int(m.group(1)), int(m.group(2))
        ap = (m.group(4) or "").lower()
        if ap:
            if not 1 <= h <= 12:
                return None
            h = h % 12 + (12 if ap == "p" else 0)
        return h * 60 + mi if h < 24 and mi < 60 else None
    m = _HHMM.fullmatch(s)
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return int(m.group(1)) * 60 + int(m.group(2))
    return None


def _time_cell(v) -> int | None:
    """Minutes of the day of a time cell: Excel's typed time or datetime, a fraction of a day, or text.  Text that
    holds a date and a time ('1899-12-30 09:15:00', as an Access or SQL time field exports it) gives its time of day,
    as a typed datetime does (the final verification of task 16)."""
    import datetime as _dt
    if isinstance(v, _dt.datetime):
        return v.hour * 60 + v.minute
    if isinstance(v, _dt.time):
        return v.hour * 60 + v.minute
    if isinstance(v, float) and 0 <= v < 1:
        return int(round(v * 86400)) // 60 % 1440
    s = _cell_text(v)
    t = _clock_text(s)
    if t is None:
        m = _ISO_DATE.match(s) or _NUM_DATE.match(s)
        if m and m.group(4):
            z = _ZONE.search(m.group(4))
            t = _clock_text(m.group(4)[:z.start()] if z else m.group(4))
    return t


def _date_order(parts: list[tuple[int, int]], column: str, example: str, date_format: str | None) -> str:
    """'MDY' or 'DMY' for the numeric dates of a column ((first, second) number of each): the given format, or what
    the dates themselves settle (a first number above 12 is a day, a second above 12 is a day).  A column where no date
    settles it, and some date reads two ways, is refused with the question."""
    if date_format in ("MDY", "DMY"):
        return date_format
    a = any(p > 12 for p, _ in parts)
    b = any(q > 12 for _, q in parts)
    if a and b:
        raise LogRefused(f"The dates in column '{column}' are written both ways: some have a day above 12 first and some "
                         f"second. Export the log again with one date format.")
    if a:
        return "DMY"
    if b:
        return "MDY"
    if all(p == q for p, q in parts):
        return "MDY"
    p, q = next((p, q) for p, q in parts if p != q)
    raise LogRefused(f"The dates in column '{column}' can be read two ways: {example} is month {p}, day {q} with the "
                     f"month first (MM/DD/YYYY), or day {p}, month {q} with the day first (DD/MM/YYYY), and no date in "
                     f"the file has a day above 12 to settle it. Which is it? Choose the date format (in the app, "
                     f"'Date format' with the grab log's columns; in the pilot, --date-format MDY or DMY).")


def _parse_dates(cells: list, column: str, date_format: str | None,
                 midnight: set | None = None) -> tuple[list, int]:
    """Each date cell as (date or None, minutes of the day written in it or None), and how many carried a time zone
    mark (Z, UTC or an offset), which is dropped: the time is read as written.  date_format is 'MDY', 'DMY' or 'YMD'
    (how numeric dates are ordered), a strptime format (with '%'), or None to let the dates settle it.  A date with
    its month's name ('10-Mar-26', '10 March 2026', 'Mar 10, 2026') needs no format.  An Excel datetime at exactly
    midnight in a column where another datetime carries a time is given 00:00 and its index is added to midnight
    (when a set is passed): read_grab_log keeps it only at an online analyzer's tap."""
    import datetime as _dt
    out: list = [None] * len(cells)
    zoned = 0
    numeric = []
    # an Excel datetime at exactly midnight is 00:00 when another datetime cell of the column carries a time (an
    # analyzer's series exported to Excel, the review of task 16), and a date only when every one is at midnight;
    # since the final verification read_grab_log keeps such a 00:00 only at an analyzer's tap (a date typed with no
    # time, in a column of datetimes, is no midnight sample)
    timed = any(isinstance(v, _dt.datetime) and (v.hour or v.minute or v.second) for v in cells)
    for i, v in enumerate(cells):
        if isinstance(v, _dt.datetime):
            try:
                t = v.hour * 60 + v.minute
                exact = not (v.hour or v.minute or v.second)
                out[i] = (v.date(), t if (timed or not exact) else None)
                if timed and exact and midnight is not None:
                    midnight.add(i)
            except (ValueError, TypeError):      # pandas' NaT is a datetime with no date
                pass
            continue
        if isinstance(v, _dt.date):
            out[i] = (v, None)
            continue
        s = _cell_text(v)
        if not s:
            continue
        if date_format and "%" in date_format:
            try:
                d = _dt.datetime.strptime(s, date_format)
            except ValueError:
                continue
            has_t = any(x in date_format for x in ("%H", "%I", "%M"))
            out[i] = (d.date(), d.hour * 60 + d.minute if has_t else None)
            continue
        z = _ZONE.search(s) if ":" in s or "T" in s else None
        if z and (_ISO_DATE.match(s[:z.start()]) or _NUM_DATE.match(s[:z.start()])):
            s = s[:z.start()]
            zoned += 1
        m = _ISO_DATE.match(s)
        if m:
            y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
            t = _clock_text(m.group(4)) if m.group(4) else None
            if m.group(4) and t is None:
                continue
            try:
                out[i] = (_dt.date(y, mo, d), t)
            except ValueError:
                pass
            continue
        m = _NUM_DATE.match(s)
        if m:
            t = _clock_text(m.group(4)) if m.group(4) else None
            if m.group(4) and t is None:
                continue
            numeric.append((i, int(m.group(1)), int(m.group(2)), int(m.group(3)), t, s))
            continue
        m = _MON_DATE.match(s)               # '10-Mar-26', '10 March 2026': the month's name settles the order
        n = None if m else _MON_FIRST.match(s)
        if m or n:
            d, mon, y, rest = ((int(m.group(1)), m.group(2), int(m.group(3)), m.group(4)) if m else
                               (int(n.group(2)), n.group(1), int(n.group(3)), n.group(4)))
            mo = _MONTHS.get(mon.lower())
            t = _clock_text(rest) if rest else None
            if mo is None or (rest and t is None):
                continue
            try:
                out[i] = (_dt.date(y + 2000 if y < 100 else y, mo, d), t)
            except ValueError:
                pass
    if numeric:
        if date_format == "YMD":
            raise LogRefused(f"The dates in column '{column}' are written like {numeric[0][5]!r}, with the year last, "
                             f"not first as the chosen format says. Choose month first or day first.")
        order = _date_order([(a, b) for _, a, b, _, _, _ in numeric], column, numeric[0][5], date_format)
        for i, a, b, y, t, _ in numeric:
            y = y + 2000 if y < 100 else y
            mo, d = (a, b) if order == "MDY" else (b, a)
            try:
                out[i] = (_dt.date(y, mo, d), t)
            except ValueError:
                pass
    return out, zoned


def model_hour(clock_minute: int, start_clock_h: float = 0.0) -> int:
    """The model hour of a time of day, in minutes: (clock time - START CLOCKTIME) mod 24 h, truncated to the hour
    (minutes are truncated, as the pilot always did).  The start is taken to the minute once, as the readiness row
    shows it (importers._clock_minute), so a sample gets the hour the report's start time implies; readiness.flags
    keeps start_clock_h unrounded."""
    return ((int(clock_minute) - _clock_minute(start_clock_h)) % 1440) // 60


def clock_minute_of(h: int, start_clock_h: float = 0.0) -> int:
    """The time of day, in minutes, of model hour h of a file that starts at start_clock_h."""
    return (int(h) * 60 + _clock_minute(start_clock_h)) % 1440


def clock_label(h: int, start_clock_h: float = 0.0) -> str:
    """'HH:MM', the time of day of model hour h ('09:00' for model hour 3 of a file that starts at 06:00)."""
    m = clock_minute_of(h, start_clock_h)
    return f"{m // 60:02d}:{m % 60:02d}"


def shows_clock(start_clock_h: float | None) -> bool:
    """True when a file's hours need converting for display: its start, to the minute, is not midnight."""
    return bool(start_clock_h) and _clock_minute(float(start_clock_h)) != 0


def daytime_model_hours(start_clock_h: float = 0.0) -> list[int]:
    """The model hours whose time of day falls in the sampling day, 07:00 to 17:00 (the hour starting at 17:00
    included, as in simgp.DAY_HOURS), in order of the time of day; for a file that starts at midnight, DAY_HOURS."""
    hours = sorted(range(24), key=lambda h: clock_minute_of(h, start_clock_h))
    return [h for h in hours if DAY_START_H * 60 <= clock_minute_of(h, start_clock_h) < (DAY_END_H + 1) * 60]


@dataclass
class GrabLog:
    """A grab log read for the model.  frame has load_log's columns (junction, hour as a model hour, y in mg/L,
    month YYYY-MM, date YYYY-MM-DD, tap_id), one row per reading that is used; clock_min is each row's time of day in
    minutes (same index).  censored lists the censored readings, kept and never fitted, its column 'kind' 'below'
    for a reading below detection ('<0.02', 'ND') and 'above' for one above the measuring range ('>2.2'), which the
    app and the pilot never call below detection (the review of task 16); dropped lists every other row that is not
    used, with its reason; notes say how the file was read, warnings what to check."""
    frame: object
    clock_min: object
    censored: object
    dropped: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    mapping: dict = field(default_factory=dict)
    analyzer: dict = field(default_factory=dict)
    start_clock_h: float = 0.0
    n_rows: int = 0                     # data rows in the file
    row_no: object = None               # each frame row's line in the file (same index as frame)

    def only_junctions(self, junctions) -> "GrabLog":
        """The same log with the rows whose junction is not in the model moved to dropped (and censored readings at
        such a junction too), with the reason."""
        import pandas as pd
        js = set(map(str, junctions))
        keep = self.frame.junction.isin(js)
        dropped = list(self.dropped)
        for (_, r), row_no in zip(self.frame[~keep].iterrows(), self.row_no[~keep]):
            dropped.append({"row": int(row_no), "tap_id": r.tap_id,
                            "reason": f"junction {r.junction} (tap {r.tap_id}) is not in the model"})
        c = self.censored
        ck = c.junction.isin(js) if len(c) else pd.Series(dtype=bool)
        for _, r in (c[~ck].iterrows() if len(c) else []):
            dropped.append({"row": int(r.row), "tap_id": r.tap_id,
                            "reason": f"junction {r.junction} (tap {r.tap_id}) is not in the model"})
        out = GrabLog(frame=self.frame[keep].reset_index(drop=True), clock_min=self.clock_min[keep].reset_index(drop=True),
                      censored=c[ck].reset_index(drop=True) if len(c) else c, dropped=sorted(dropped, key=lambda d: d["row"]),
                      notes=list(self.notes), warnings=list(self.warnings), mapping=dict(self.mapping),
                      analyzer=dict(self.analyzer), start_clock_h=self.start_clock_h, n_rows=self.n_rows,
                      row_no=self.row_no[keep].reset_index(drop=True))
        return out

    def window(self, months: int = LOG_WINDOW_MONTHS) -> tuple:
        """(rows, censored rows, info) of the last `months` calendar months that hold a reading (used or censored),
        as the pilot's --window counts them; info gives the months, the window's first and last date, and the log's,
        and 'outside': every row before the window (a reading or a censored one), with its line and the reason, so
        the not-used list can hold it (the review of task 16); n_outside counts them."""
        all_months = sorted(set(self.frame.month) | (set(self.censored.month) if len(self.censored) else set()))
        use = all_months[-months:]
        inside = self.frame.month.isin(use)
        f = self.frame[inside].reset_index(drop=True)
        c = self.censored[self.censored.month.isin(use)].reset_index(drop=True) if len(self.censored) else self.censored
        dates = sorted(set(f.date) | (set(c.date) if len(c) else set()))
        all_dates = sorted(set(self.frame.date) | (set(self.censored.date) if len(self.censored) else set()))
        why = f"dated before the window (the last {months} months that hold a reading, from {use[0] if use else ''})"
        outside = [{"row": int(n), "tap_id": t, "reason": why}
                   for n, t in zip(self.row_no[~inside], self.frame.tap_id[~inside])]
        if len(self.censored):
            outside += [{"row": int(r.row), "tap_id": r.tap_id, "reason": why}
                        for r in self.censored[~self.censored.month.isin(use)].itertuples()]
        info = {"months": use, "n_months_in_log": len(all_months), "first": dates[0] if dates else None,
                "last": dates[-1] if dates else None, "log_first": all_dates[0] if all_dates else None,
                "log_last": all_dates[-1] if all_dates else None, "n_used": len(f), "n_censored": len(c),
                "n_outside": len(outside), "outside": sorted(outside, key=lambda d: d["row"])}
        return f, c, info


def read_tap_map(data: bytes, name: str = "tap_map.csv", mapping: dict | None = None) -> tuple[dict, LogTable, dict]:
    """(tap ID -> junction ID, the table, the columns used).  IDs are kept as text, leading zeros included.  A tap
    listed twice with two different junctions is refused; listed twice with one junction, it is kept once."""
    t = read_table(data, name)
    cols = dict(guess_tap_columns(t.header))
    cols.update({k: v for k, v in (mapping or {}).items() if v})
    if cols["tap"] not in t.header or cols["junction"] not in t.header:
        raise LogRefused(f"The tap map needs a tap column and a junction column; its header is "
                         f"{', '.join(t.header[:8])}. Name them tap_id and junction_id.")
    it, ij = t.header.index(cols["tap"]), t.header.index(cols["junction"])
    out, clash = {}, []
    for r in t.rows:
        tap, j = _cell_text(r[it]), _cell_text(r[ij])
        if not tap or not j:
            continue
        if tap in out and out[tap] != j:
            clash.append(tap)
        out[tap] = out.get(tap, j)
    if clash:
        raise LogRefused(f"The tap map lists {len(clash)} tap{'s' if len(clash) > 1 else ''} with two different "
                         f"junctions: {', '.join(sorted(set(clash))[:6])}. Keep one junction per tap.")
    return out, t, cols


def _analyte_by_values(t: LogTable, skip: set) -> str | None:
    """A long-format log's analyte column found by its values when its header says nothing ('Description', 'Name'):
    the first column (not in skip) filled on at least 90% of the first 2,000 rows, with at most 50 distinct values,
    one of which names free or total chlorine."""
    for i, h in enumerate(t.header):
        if h in skip:
            continue
        vals = [_cell_text(r[i]) for r in t.rows[:2000]]
        filled = [v for v in vals if v]
        distinct = set(filled)
        if (vals and len(filled) >= 0.9 * len(vals) and 0 < len(distinct) <= 50
                and any(analyte_kind(v) in ("free", "total") for v in distinct)):
            return h
    return None


def distinct_values(t: LogTable, column: str, cap: int = 200) -> list[str]:
    """The distinct filled values of one column, as read_grab_log compares them (the app's 'Analyte value' choice),
    sorted, at most cap of them."""
    if column not in t.header:
        return []
    i = t.header.index(column)
    return sorted({_cell_text(r[i]) for r in t.rows} - {""}, key=str.lower)[:cap]


def guess_columns(t: LogTable, disinfectant: str = "free_chlorine") -> dict:
    """guess_log_columns on a table's header, and, when no header names an analyte column, the analyte column found
    by its values (_analyte_by_values), so the app's guesses and read_grab_log's agree."""
    g = guess_log_columns(t.header, disinfectant)
    if g["analyte"] is None:
        g["analyte"] = _analyte_by_values(t, {g[k] for k in ("tap", "reading", "date", "time", "unit")} - {None})
        if g["reading"] is None and g["analyte"] is not None:
            norm = {h: " ".join(_norm_header(h).split()) for h in t.header}
            g["reading"] = next((h for e in _RESULT_EXACT for h in t.header if norm[h] == e), None)
    return g


_CHOOSE_READING = ("Which column is the reading? Choose it (in the app, 'Reading column' with the grab log's columns; in "
                   "the pilot, --columns 'reading=<its header>'), and only that column is read; or drop or rename the "
                   "other column (or, if the log mixes a free-chlorine and a chloraminated system, split it by "
                   "disinfectant).")
_CHOOSE_VALUE = ("Which value is the chlorine reading? Choose it (in the app, 'Analyte value' with the grab log's "
                 "columns; in the pilot, --columns 'analyte=<the column>;analyte_value=<the value>'), and only the rows "
                 "of that value are read; or, if every row is a chlorine reading and the column names something else, "
                 "choose '(every row)' (analyte_value=*).")
EVERY_ROW = "*"                  # analyte_value: the analyte column does not pick the rows, every row is a reading
_ANALYTE_LIKE_NOT = {"date", "time", "result", "results", "value", "values", "unit", "units", "id", "method"}


def _species_question(guess: dict, what: str) -> str:
    """The question for a log with two or more reading columns that are not all of one named species: a free and a
    total column, or a column whose header names no species beside another (the final verification of task 16)."""
    cols = guess["reading_columns"]
    if guess["free_columns"] and guess["total_columns"]:
        head = (f"The grab log has both a free chlorine column ({guess['free_columns'][0]}) and a total chlorine "
                f"column ({guess['total_columns'][0]}).")
    else:
        unnamed = guess["other_columns"]
        head = (f"The grab log has {len(cols)} chlorine columns ({', '.join(cols[:4])}), and "
                + (f"the header of '{unnamed[0]}' does not say whether it is free or total chlorine."
                   if len(unnamed) == 1 else "their headers do not say which is free and which is total chlorine."))
    return (f"{head} Free and total chlorine are different measurements and one model never mixes them: {what}. "
            + _CHOOSE_READING)


def _values_text(values) -> str:
    vs = sorted(values, key=lambda v: (v == "", v.lower()))
    return ", ".join(repr(v) for v in vs[:6]) + (f" and {len(vs) - 6} more" if len(vs) > 6 else "")


def read_grab_log(log_data: bytes, log_name: str = "grab_log.csv", taps_data: bytes | None = None,
                  taps_name: str = "tap_map.csv", disinfectant: str = "free_chlorine", date_format: str | None = None,
                  mapping: dict | None = None, start_clock_h: float = 0.0, analyzer_hour: float = ANALYZER_HOUR,
                  tap_mapping: dict | None = None, thin_analyzers: bool = True) -> GrabLog:
    """A grab log and its tap map read for the model (task 16).  The columns are guessed (guess_columns) and any
    given in mapping ({'tap', 'reading', 'date', 'time', 'analyte', 'unit'}, header names; 'time' None for one datetime
    column, 'analyte' or 'unit' None for none; and 'analyte_value', the one analyte whose rows are read) replace the
    guess.  Free and total chlorine are never mixed: a log with two reading columns that are not both of one named
    species (a free and a total column, or one whose header names no species) is refused unless mapping names the
    reading column, and then only that column is read; a reading column that names the other species is refused; in a
    long-format export (one row per analyte) only the rows whose analyte is the disinfectant's species are read, and
    the rest are listed with their reason; a long-format export whose analytes do not pick the species (two chlorine
    analytes that name no species, or none that names chlorine while the reading column is a plain 'Result') is
    refused with the question until analyte_value is given, and so is a plain 'Result' column beside an analyte-like
    column ('Parameter Code') that is not chosen; a reading column that names neither species is read as the
    disinfectant chosen, with a warning.  A
    reading's unit is the one its header or the units column names (µg/L, ug/L and ppb are divided by 1000, with a
    warning; mg/L, ppm or none are mg/L), and a reading above READING_MAX_MGL is listed, not used.  Each time of day
    becomes a model hour (model_hour: the time minus START CLOCKTIME, mod 24 h, minutes truncated); start_clock_h is
    0 for every bundled example, so their hours are unchanged.  A tap with ANALYZER_MIN_PER_DAY or more readings a
    day at a median step of ANALYZER_MAX_STEP_MIN minutes or less on ANALYZER_MIN_DAYS days or more (or with that
    many readings a day on ANALYZER_DENSE_DAYS days or more) is an online analyzer: it keeps one reading per day, the
    one nearest analyzer_hour, unless thin_analyzers is False (a sampling study the operator wants kept whole).  Tap
    IDs that are numbers in the log and not found as written are matched by number (Excel drops leading zeros), with
    a warning.  taps_data None reads each tap ID as its own junction (the Get started page's check of a log alone).
    Raises LogRefused when the log cannot be read as it is (with a question when a date reads two ways or a log has
    both species and no reading column was chosen)."""
    import pandas as pd
    if disinfectant not in DISINFECTANT_WORDS:
        raise ValueError(f"disinfectant must be one of {sorted(DISINFECTANT_WORDS)}")
    if date_format is not None and "%" not in date_format and date_format not in DATE_ORDERS:
        raise ValueError(f"date_format must be one of {sorted(DATE_ORDERS)} or a strptime format with '%'")
    t = read_table(log_data, log_name)
    guess = guess_columns(t, disinfectant)
    cols = {k: guess[k] for k in ("tap", "reading", "date", "time", "analyte", "unit")}
    for k, v in (mapping or {}).items():
        if k in cols:
            cols[k] = v or None
    want = "total" if disinfectant == "chloramine" else "free"
    other = "free" if want == "total" else "total"
    word = {"free": "free chlorine", "total": "total chlorine"}
    what = ("the free-chlorine model reads free chlorine only" if want == "free" else
            "the chloramine mode reads total chlorine only")
    chosen = bool((mapping or {}).get("reading"))
    if guess["both_species"] and not chosen:
        raise LogRefused(_species_question(guess, what))
    missing = [k for k in ("tap", "reading", "date") if not cols[k] or cols[k] not in t.header]
    missing += [k for k in ("time", "analyte", "unit") if cols[k] and cols[k] not in t.header]
    if missing:
        raise LogRefused(f"No {', '.join(missing)} column was found in the grab log's header ({', '.join(t.header[:10])}"
                         f"). Choose the columns (in the app, with the grab log's columns; in the pilot, --columns), "
                         f"or name them date, time, tap_id and {'total' if want == 'total' else 'free'}_chlorine_mgL.")
    sp = reading_species(cols["reading"])
    if sp is not None and sp != want:
        raise LogRefused(f"The reading column '{cols['reading']}' is {word[sp]}, and "
                         + ("the free-chlorine model reads free chlorine only: choose 'Chloramine (total chlorine)' "
                            "(--disinfectant chloramine) for a chloraminated system." if want == "free" else
                            "the chloramine mode reads total chlorine only: choose 'Free chlorine' for a free-chlorine "
                            "system."))
    warns, notes = [], [f"Read as {t.how}; {len(t.rows)} data row{'s' if len(t.rows) != 1 else ''}."]
    notes.append(f"Columns: tap '{cols['tap']}', reading '{cols['reading']}', date '{cols['date']}'"
                 + (f", time '{cols['time']}'" if cols["time"] else " (the time is in the date column)")
                 + (f", analyte '{cols['analyte']}'" if cols["analyte"] else "")
                 + (f", unit '{cols['unit']}'" if cols["unit"] else "") + ".")
    if guess["both_species"]:
        rest = [h for h in guess["reading_columns"] if h != cols["reading"]]
        notes.append(f"The log's other chlorine column{'s' if len(rest) != 1 else ''} ({', '.join(rest)}) "
                     f"{'are' if len(rest) != 1 else 'is'} not read: the reading column chosen is "
                     f"'{cols['reading']}', and one model never mixes free and total chlorine.")
    ix = {k: (t.header.index(v) if v else None) for k, v in cols.items()}
    # a long-format export: which rows are the disinfectant's species (the final verification of task 16: rows are
    # never all read when the analytes cannot say which are the reading)
    value = (mapping or {}).get("analyte_value") or None
    generic = not is_reading_header(cols["reading"])     # 'Result', 'Value': the header does not say what it holds
    kinds, pick = {}, None
    if value is not None and not cols["analyte"]:
        raise LogRefused(f"An analyte value ({value!r}) was chosen, but no analyte column: choose the column that names "
                         f"each row's analyte too.")
    if cols["analyte"] and value == EVERY_ROW:
        warns.append(f"The column '{cols['analyte']}' was chosen as not picking the rows (every row): every row of "
                     f"'{cols['reading']}' is read as {word[want]}, whatever that column says. Check that it is.")
        cols["analyte"], ix["analyte"] = None, None
    elif cols["analyte"]:
        values = {_cell_text(r[ix["analyte"]]) for r in t.rows}
        kinds = {v: analyte_kind(v) for v in values}
        A = cols["analyte"]
        if value is not None:
            if value not in values:
                raise LogRefused(f"No row's analyte in column '{A}' is {value!r}; its values are {_values_text(values)}.")
            vsp = reading_species(value)
            if vsp is not None and vsp != want:
                raise LogRefused(f"The analyte chosen, {value!r}, is {word[vsp]}, and "
                                 + ("the free-chlorine model reads free chlorine only: choose 'Chloramine (total "
                                    "chlorine)' (--disinfectant chloramine) for a chloraminated system." if want == "free"
                                    else "the chloramine mode reads total chlorine only: choose 'Free chlorine' for a "
                                         "free-chlorine system."))
            pick = value
        elif all(k == "other" for k in kinds.values()):
            if generic:
                raise LogRefused(f"The column '{A}' names each row's analyte ({_values_text(values)}), and none of its "
                                 f"values names chlorine in words ResidualMap reads (such as 'Free Chlorine', 'Total "
                                 f"Cl', 'Chlorine (Free)' or FCL2), so it cannot tell which rows of '{cols['reading']}' "
                                 f"are the reading; reading every row would mix in every other analyte. "
                                 + _CHOOSE_VALUE)
            warns.append(f"The column '{A}' looks like an analyte column, but none of its values names "
                         f"chlorine, so it does not pick the rows.")
            kinds, cols["analyte"], ix["analyte"] = {}, None, None
        elif want not in kinds.values():
            unnamed = [v for v, k in kinds.items() if k == "chlorine"]
            others = [v for v, k in kinds.items() if k == other]
            if len(unnamed) > 1 or (unnamed and others):
                raise LogRefused(f"The analyte column '{A}' names chlorine in {len(unnamed) + len(others)} ways "
                                 f"({_values_text(unnamed + others)}), and "
                                 + ("none of them says" if not others else
                                    f"{_values_text(unnamed)} {'does' if len(unnamed) == 1 else 'do'} not say")
                                 + f" whether it is free or total chlorine. Free and total chlorine are different "
                                   f"measurements and one model never mixes them: {what}. " + _CHOOSE_VALUE)
    elif generic:
        used = {v for v in cols.values() if v}
        for i, h in enumerate(t.header):
            w = set(_norm_header(h).split())
            if h in used or not w & _ANALYTE_WORDS or w & (_ANALYTE_LIKE_NOT | set(_TAP_WORDS)):
                continue
            vals = {_cell_text(r[i]) for r in t.rows[:2000]} - {""}
            if len(vals) >= 2:
                raise LogRefused(f"The column '{h}' looks like it names each row's analyte ({_values_text(vals)}), and "
                                 f"the reading column '{cols['reading']}' does not say what it measures, so every row "
                                 f"would be read, whatever its analyte. Choose '{h}' as the analyte column (in the app, "
                                 f"'Analyte column'; in the pilot, --columns 'analyte={h}'), then the analyte value of "
                                 f"the chlorine reading if it is not found by its name, or '(every row)' "
                                 f"(analyte_value=*) if every row is a chlorine reading.")
    if not guess["both_species"] and not chosen and len(guess["reading_columns"]) > 1:
        rest = [h for h in guess["reading_columns"] if h != cols["reading"]]
        notes.append(f"The log has {len(guess['reading_columns'])} {word.get(sp, 'chlorine')} columns: "
                     f"'{cols['reading']}', the first, is read, and {', '.join(repr(h) for h in rest)} "
                     f"{'is' if len(rest) == 1 else 'are'} not; choose the reading column if it is another.")
    named = [k for k in kinds.values() if k in ("free", "total")]
    if pick is not None:
        if reading_species(pick) is None:
            warns.append(f"The analyte chosen, {pick!r} (column '{cols['analyte']}'), does not say free or total "
                         f"chlorine{'' if names_chlorine(pick) else ', nor name chlorine in words ResidualMap reads'}; "
                         f"its rows are read as {word[want]}, the disinfectant chosen. Check that it is.")
    elif sp is None and not (kinds and want in named) and value != EVERY_ROW:
        warns.append(f"The reading column '{cols['reading']}' does not say free or total chlorine; it is read as "
                     f"{word[want]}, the disinfectant chosen. Check that it is.")
    head_factor = header_unit(cols["reading"]) or 1.0
    midnight: set = set()
    dates, zoned = _parse_dates([r[ix["date"]] for r in t.rows], cols["date"], date_format, midnight)
    if zoned:
        warns.append(f"{zoned} date{'s' if zoned != 1 else ''} carry a time zone mark (Z, UTC or an offset): each time "
                     f"is read as written. If the log's times are UTC, export them in local time.")
    taps, tmap_cols = (None, {}) if taps_data is None else read_tap_map(taps_data, taps_name, tap_mapping)[0::2]
    by_number = {}
    if taps is not None:
        for k in taps:
            if k.isascii() and k.isdigit():         # '1²' passes isdigit() and fails int() (the final verification)
                by_number.setdefault(int(k), []).append(k)
    rows, dropped, cens, matched_by_number, commas, micro, n_other = [], [], [], {}, 0, 0, 0
    for i_row, (r, row_no, dt_) in enumerate(zip(t.rows, t.row_numbers, dates)):
        tap = _cell_text(r[ix["tap"]])
        raw = r[ix["reading"]]
        if kinds:
            a = _cell_text(r[ix["analyte"]])
            k = kinds[a]
            keep = a == pick if pick is not None else (k == want or (k == "chlorine" and not named))
            if not keep:
                why = (f"{word[k]}, not {word[want]} ({a!r}): one model never mixes them" if k in word and k != want
                       else f"another analyte ({a!r}), not the one chosen ({pick!r})" if pick is not None and k != "other"
                       else f"the analyte does not say free or total chlorine ({a!r})" if k == "chlorine" else
                       f"another analyte ({a!r})")
                dropped.append({"row": row_no, "tap_id": tap, "reason": why})
                n_other += 1
                continue
        if not tap:
            dropped.append({"row": row_no, "tap_id": "", "reason": "no tap ID"})
            continue
        junction = tap if taps is None else taps.get(tap)
        if junction is None and tap.isascii() and tap.isdigit() and len(by_number.get(int(tap), [])) == 1:
            key = by_number[int(tap)][0]
            junction = taps[key]
            matched_by_number[tap] = key
        if junction is None:
            dropped.append({"row": row_no, "tap_id": tap, "reason": f"tap {tap} is not in the tap map"})
            continue
        if dt_ is None:
            dropped.append({"row": row_no, "tap_id": tap, "reason": "date not read",
                            "example": _cell_text(r[ix["date"]])})
            continue
        day, minute = dt_
        if ix["time"] is not None:
            minute = _time_cell(r[ix["time"]])
        if minute is None:
            dropped.append({"row": row_no, "tap_id": tap, "reason": "no time of day",
                            "example": _cell_text(r[ix["time"]]) if ix["time"] is not None else
                            _cell_text(r[ix["date"]])})
            continue
        factor = head_factor
        if ix["unit"] is not None and _cell_text(r[ix["unit"]]):
            u = _cell_text(r[ix["unit"]])
            factor = header_unit(u)
            if factor is None:
                dropped.append({"row": row_no, "tap_id": tap, "reason": f"the unit is {u!r}, not mg/L"})
                continue
        y, why = parse_reading(raw)
        rec = {"row": row_no, "tap_id": matched_by_number.get(tap, tap), "junction": str(junction),
               "date": day.isoformat(), "month": day.isoformat()[:7], "clock_min": int(minute),
               "hour": model_hour(minute, start_clock_h)}
        if ix["time"] is None and i_row in midnight:     # an Excel 00:00 that may be a date typed with no time
            rec.update(midnight=True, date_text=_cell_text(r[ix["date"]]))
        if y is None and why.startswith(("censored", "above")):
            cens.append({**rec, "raw": _cell_text(raw), "reason": why,
                         "kind": "above" if why.startswith("above") else "below"})
        elif y is None:
            dropped.append({"row": row_no, "tap_id": tap, "reason": why})
        elif y * factor > READING_MAX_MGL:
            dropped.append({"row": row_no, "tap_id": tap,
                            "reason": f"above {READING_MAX_MGL:g} mg/L, more than any chlorine residual: check the "
                                      f"reading column and its unit", "example": _cell_text(raw)})
        else:
            rows.append({**rec, "y": y * factor if factor != 1.0 else y})
            commas += isinstance(raw, str) and "," in raw
            micro += factor != 1.0
    if commas:
        notes.append(f"{commas} reading{'s' if commas != 1 else ''} written with a decimal comma (0,45) "
                     f"{'were' if commas != 1 else 'was'} read as decimals (0.45).")
    if micro:
        warns.append(f"{micro} reading{'s are' if micro != 1 else ' is'} in µg/L (micrograms per litre), as the "
                     f"{'units column' if ix['unit'] is not None and head_factor == 1.0 else 'reading column'} says: "
                     f"each was divided by 1000 to mg/L. Check that the unit is right.")
    if pick is not None:
        notes.append(f"Each row's analyte is in column '{cols['analyte']}': the rows of {pick!r}, the analyte chosen, are "
                     f"read, and the {n_other} other row{'s' if n_other != 1 else ''} are listed as not used.")
    elif kinds and want in named:
        notes.append(f"Each row's analyte is in column '{cols['analyte']}': the rows of {word[want]} are read, and the "
                     f"{n_other} row{'s' if n_other != 1 else ''} of other analytes or of {word[other]} are listed as "
                     f"not used.")
    n_high = sum(d["reason"].startswith(f"above {READING_MAX_MGL:g} mg/L") for d in dropped)
    if n_high and rows and not micro:      # the final verification of task 16: a µg/L log its header does not name
        warns.append(f"{n_high} reading{'s are' if n_high != 1 else ' is'} above {READING_MAX_MGL:g} mg/L and "
                     f"{len(rows)} {'are' if len(rows) != 1 else 'is'} not: a log in µg/L (micrograms per litre) whose "
                     f"header and units do not say so looks like this, and its readings of {READING_MAX_MGL:g} or less "
                     f"would be fitted 1000 times too high. Check the unit: if it is µg/L, name it in the reading "
                     f"column's header (such as 'Free Chlorine (ug/L)') or in a units column.")
    if matched_by_number:
        ex = ", ".join(f"{a} is {b}" for a, b in sorted(matched_by_number.items())[:6])
        warns.append(f"{len(matched_by_number)} tap ID{'s' if len(matched_by_number) != 1 else ''} in the log "
                     f"{'are numbers' if len(matched_by_number) != 1 else 'is a number'} with no leading zeros, as "
                     f"Excel stores them; matched by number to the tap map ({ex}). Check that these are the same taps.")
    # online analyzers (by their cadence: many readings a day, minutes apart, on several days): one reading per tap
    # per day, the one nearest analyzer_hour
    analyzer = {}
    if rows:
        per_day = {}
        for x in rows:
            per_day.setdefault(x["tap_id"], {}).setdefault(x["date"], []).append(x)
        target = int(round(analyzer_hour * 60))
        keep_ids = set()
        for tap_, days in per_day.items():
            dense = fast = 0
            for xs in days.values():
                if len(xs) < ANALYZER_MIN_PER_DAY:
                    continue
                dense += 1
                ms = sorted(x["clock_min"] for x in xs)
                steps = sorted(b - a for a, b in zip(ms, ms[1:]))
                fast += steps[len(steps) // 2] <= ANALYZER_MAX_STEP_MIN
            if fast < ANALYZER_MIN_DAYS and dense < ANALYZER_DENSE_DAYS:
                continue
            counts = sorted(len(v) for v in days.values())
            n = sum(counts)
            if not thin_analyzers:
                analyzer[tap_] = {"readings": n, "days": len(days), "kept": n, "days_without": 0,
                                  "per_day": counts[len(counts) // 2], "thinned": False}
                continue
            kept, gaps = [], 0
            for day, xs in sorted(days.items()):
                best = min(xs, key=lambda x: (abs(x["clock_min"] - target), x["clock_min"]))
                if abs(best["clock_min"] - target) <= ANALYZER_MAX_GAP_MIN:
                    kept.append(id(best))
                else:
                    gaps += 1
            analyzer[tap_] = {"readings": n, "days": len(days), "kept": len(kept), "days_without": gaps,
                              "per_day": counts[len(counts) // 2], "thinned": True}
            keep_ids |= set(kept)
            if gaps:
                dropped.append({"row": min(x["row"] for v in days.values() for x in v), "tap_id": tap_,
                                "reason": f"analyzer series: no reading within {ANALYZER_MAX_GAP_MIN} min of "
                                          f"{_clock(analyzer_hour)} on {gaps} day{'s' if gaps != 1 else ''}"})
        if analyzer:
            names = f"{len(analyzer)} tap{'s' if len(analyzer) != 1 else ''} ({', '.join(sorted(analyzer)[:6])})"
            look = "look" if len(analyzer) != 1 else "looks"
            n_in = sum(a["readings"] for a in analyzer.values())
            per = sorted(a["per_day"] for a in analyzer.values())[len(analyzer) // 2]
            if thin_analyzers:
                for x in rows:
                    if x["tap_id"] in analyzer and id(x) not in keep_ids:
                        dropped.append({"row": x["row"], "tap_id": x["tap_id"],
                                        "reason": f"analyzer series at tap {x['tap_id']}: thinned to one reading a "
                                                  f"day, the one nearest {_clock(analyzer_hour)}"})
                rows = [x for x in rows if x["tap_id"] not in analyzer or id(x) in keep_ids]
                n_out = sum(a["kept"] for a in analyzer.values())
                n_grab = sum(x["tap_id"] not in analyzer for x in rows)
                notes.append(f"{names} {look} like an online analyzer ({n_in:,} readings, {per} a day on the median "
                             f"day, minutes apart): thinned to one reading per tap per day, the one nearest "
                             f"{_clock(analyzer_hour)}, so each series counts once a day, not {per} times: {n_out:,} "
                             f"analyzer reading{'s' if n_out != 1 else ''} and {n_grab:,} grab "
                             f"sample{'s' if n_grab != 1 else ''} kept.")
            else:
                notes.append(f"{names} {look} like an online analyzer ({n_in:,} readings, {per} a day on the median "
                             f"day): every reading is kept, as asked (a sampling study, not thinned), so each counts in "
                             f"the fit.")
    # an Excel datetime at exactly 00:00 among timed ones: kept at an online analyzer's tap (its midnight reading),
    # listed as 'no time of day' elsewhere, a date typed without its time (the final verification of task 16)
    late = [x for x in rows + cens if x.get("midnight") and x["tap_id"] not in analyzer]
    if late:
        gone = {id(x) for x in late}
        rows = [x for x in rows if id(x) not in gone]
        cens = [x for x in cens if id(x) not in gone]
        for x in late:
            dropped.append({"row": x["row"], "tap_id": x["tap_id"], "reason": "no time of day",
                            "example": x["date_text"]})
        notes.append(f"{len(late)} Excel date{'s' if len(late) != 1 else ''} at exactly 00:00 in a column of dates "
                     f"with times, at taps that are not online analyzers, {'are' if len(late) != 1 else 'is'} read as "
                     f"a date whose time was not typed and listed as 'no time of day'; an online analyzer's 00:00 "
                     f"readings are kept.")
    grab = [x for x in rows if x["tap_id"] not in analyzer]
    if grab:
        out_ = sum(not (UTC_DAY_MIN[0] <= x["clock_min"] <= UTC_DAY_MIN[1]) for x in grab)
        if out_ > len(grab) / 2:
            warns.append(f"{out_} of {len(grab)} sample times fall outside 06:00 to 18:00, when grab samples are "
                         f"usually taken. The times may have been exported in UTC rather than local time: check one "
                         f"sample against its field sheet, and export the log in local time if so.")
    if rows:
        ys = sorted(x["y"] for x in rows)
        med = ys[len(ys) // 2]
        if med > READING_HIGH_MEDIAN_MGL:
            warns.append(f"The readings' median is {med:g} mg/L, above {READING_HIGH_MEDIAN_MGL:g} mg/L, higher than a "
                         f"distribution residual usually is: check that the reading column is chlorine and in mg/L.")
    if shows_clock(start_clock_h):
        notes.append(f"Times of day are converted to model hours: the model starts at {_clock(start_clock_h)} (START "
                     f"CLOCKTIME), so a sample taken at {clock_label(3, start_clock_h)} is model hour 3. Minutes are "
                     f"truncated.")
    cols_out = ["junction", "hour", "y", "month", "date", "tap_id"]
    frame = pd.DataFrame({"junction": [x["junction"] for x in rows], "hour": [int(x["hour"]) for x in rows],
                          "y": [float(x["y"]) for x in rows], "month": [x["month"] for x in rows],
                          "date": [x["date"] for x in rows], "tap_id": [x["tap_id"] for x in rows]},
                         columns=cols_out)
    frame["hour"] = frame["hour"].astype(int)
    frame["y"] = frame["y"].astype(float)
    censored = pd.DataFrame(cens, columns=["row", "tap_id", "junction", "date", "month", "clock_min", "hour", "raw",
                                           "reason", "kind"])
    gl = GrabLog(frame=frame, clock_min=pd.Series([x["clock_min"] for x in rows], dtype=int), censored=censored,
                 dropped=sorted(dropped, key=lambda d: d["row"]), notes=notes, warnings=warns,
                 mapping={**cols, "analyte_value": EVERY_ROW if value == EVERY_ROW else pick, "date_format": date_format,
                          "tap_map": tmap_cols},
                 analyzer=analyzer,
                 start_clock_h=float(start_clock_h or 0.0), n_rows=len(t.rows),
                 row_no=pd.Series([x["row"] for x in rows], dtype=int))
    return gl


def _row_runs(rows: list[int]) -> list[str]:
    """['3', '5 to 9', '12']: every row, a run of consecutive rows written as a range."""
    rs, out, i = sorted(set(rows)), [], 0
    while i < len(rs):
        j = i
        while j + 1 < len(rs) and rs[j + 1] == rs[j] + 1:
            j += 1
        out.append(str(rs[i]) if j == i else f"{rs[i]} to {rs[j]}")
        i = j + 1
    return out


def _row_ranges(rows: list[int]) -> str:
    """'3, 5 to 9, 12': every row, with a run of consecutive rows written as a range."""
    return ", ".join(_row_runs(rows))


# the reasons that name a value, and the words of their group: when more than DROPPED_MAX_VALUES values of one group
# are listed (a lab qualifier on every reading, '0.452 J'; a tap map from another system), dropped_lines gives the
# group one line, with its count of values and a few examples (the final verification of task 16)
_REASON_GROUPS = (
    (re.compile(r"^tap (.+) is not in the tap map$"), "a tap that is not in the tap map"),
    (re.compile(r"^junction (.+) \(tap .+\) is not in the model$"), "a junction that is not in the model"),
    (re.compile(r"^not a number: (.+)$"), "not a number"),
    (re.compile(r"^negative reading: (.+)$"), "a negative reading"),
    (re.compile(r"^the unit is (.+), not mg/L$"), "a unit that is not mg/L"),
    (re.compile(r"^another analyte \((.+?)\)(, not the one chosen .*)?$"), "another analyte"),
    (re.compile(r"^the analyte does not say free or total chlorine \((.+)\)$"),
     "an analyte that does not say free or total chlorine"),
    (re.compile(r"^(free chlorine|total chlorine), not (?:free|total) chlorine \((.+)\): one model never mixes them$"),
     None),
    (re.compile(r"^censored, (below detection|not detected) \((.+)\): kept and listed, not used in the fit$"), None),
    (re.compile(r"^above the measuring range \((.+)\): not used in the fit$"),
     "above the measuring range: not used in the fit"),
)


def _reason_group(reason: str) -> tuple[str, str] | None:
    """(the group's words, the value) of a reason that names a value, None for a reason that does not."""
    for rx, words in _REASON_GROUPS:
        m = rx.match(reason)
        if m:
            if words is None:            # the species and censored reasons keep their first words
                head, val = m.group(1), m.group(2)
                words = (f"{head}, not {'total' if head.startswith('free') else 'free'} chlorine: one model never "
                         f"mixes them" if head.endswith("chlorine") else
                         f"censored, {head}: kept and listed, not used in the fit")
                return words, val
            return words, m.group(1)
    return None


def dropped_lines(dropped: list[dict], max_runs: int = DROPPED_MAX_RUNS) -> list[str]:
    """Every row that is not used, grouped by its reason, with one example of what the file holds when the reason has
    one ('date not read (for example '24-Jan-26')'), and its rows (a run of consecutive rows as a range):
    'tap T99 is not in the tap map: rows 5, 9 to 12'.  A reason with more than max_runs runs of rows (an analyzer
    export interleaved by time) gives its count, its first and last row and its first max_runs runs, so the list
    stays short; the count still covers every row (the review of task 16).  Since the final verification, a reason
    that names a value ('not a number: '0.452 J'', 'tap X is not in the tap map', 'the unit is ...') is one line for
    all its values when more than DROPPED_MAX_VALUES of them differ, with their count and three examples, and the
    list has at most DROPPED_MAX_LINES lines, the last counting the reasons and rows it leaves out."""
    by: dict[str, list[int]] = {}
    example: dict[str, str] = {}
    for d in dropped:
        by.setdefault(d["reason"], []).append(int(d["row"]))
        if d.get("example") and d["reason"] not in example:
            example[d["reason"]] = d["example"]
    groups: dict[str, list[str]] = {}
    for why in by:
        g = _reason_group(why)
        if g:
            groups.setdefault(g[0], []).append(why)
    merged: dict[str, tuple] = {}            # a reason -> (the group's line head, the group's rows), once per group
    skip = set()
    for words, whys in groups.items():
        if len(whys) > DROPPED_MAX_VALUES:
            vals = [v if v[:1] in "'\"" else repr(v) for v in (_reason_group(w)[1] for w in whys)]
            head = (f"{words} ({len(whys):,} different values, for example {', '.join(vals[:3])})")
            merged[whys[0]] = (head, [r for w in whys for r in by[w]])
            skip |= set(whys[1:])
    out = []
    for why, rows in by.items():
        if why in skip:
            continue
        if why in merged:
            head, rows = merged[why]
        else:
            head = why + (f" (for example {example[why]!r})" if why in example else "")
        runs = _row_runs(rows)
        n = len(set(rows))
        if len(runs) <= max_runs:
            out.append(f"{head}: row{'s' if n != 1 else ''} {', '.join(runs)}")
        else:
            out.append(f"{head}: {n:,} rows, from row {min(rows):,} to row {max(rows):,} ({', '.join(runs[:max_runs])}, "
                       f"and {len(runs) - max_runs:,} more runs of rows)")
    if len(out) > DROPPED_MAX_LINES:
        left = [w for w in by if w not in skip][DROPPED_MAX_LINES - 1:]       # the reasons of the lines left out
        n_rows = len({r for w in left for r in (merged[w][1] if w in merged else by[w])})
        out = out[:DROPPED_MAX_LINES - 1] + [f"and {len(left):,} more reasons, for {n_rows:,} more "
                                             f"row{'s' if n_rows != 1 else ''}"]
    return out


def xlsx_bytes(rows: list[list], sheet: str = "grab log") -> bytes:
    """A one-sheet workbook of these rows, written by openpyxl (the Get started page's .xlsx template; nothing binary
    is committed).  Read back by read_table, this is same-library self-consistency."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- tap maps by coordinates (task 16b)
# A tap map may give each tap's position instead of its junction: latitude and longitude from a phone or a GPS unit,
# or x and y in the model's own coordinates.  Each position is put on the model's nearest pipe and that pipe's
# downstream junction is suggested, for the operator to check before anything uses it (pilot protocol section 2: not
# the nearest junction by straight line, and never a big main the tap is not on).  Nothing here runs inside
# read_grab_log: a suggestion, once confirmed, is written as a tap map by ID (tap_id, junction_id, ...) and read
# through read_tap_map like any other.  Task 16b stopped at its pre-registered bar before the app's table and confirm
# step and the pilot's --accept-snapped were built, so no operator path reaches this section yet.
MAIN_MIN_IN = 16.0          # pipes of 16 in or more (nominal) are left out of a tap's snap unless its row says on_main
MAIN_CUT_RATIO = 0.9375     # ... judged on the model's own diameter at this share of it (15.0 in for 16 in), so a
                            # metric 400 mm main (15.75 in) and a 16 in main given by an internal diameter of 15.0 in
                            # or more count as mains too (the review of task 16b)
MAIN_BAND_RATIO = 0.8       # an untagged tap put on a pipe from this share of it (12.8 in for 16 in) up to the cut is
                            # flagged, with the nearest smaller pipe: a 16 in plastic main given by its internal
                            # diameter (14.9 in for C905 DR14, 13.1 in for HDPE DR11, 12.9 in for PE100 SDR11) is under
                            # the cut (the final verification of task 16b)
SECOND_RATIO = 1.5          # a second candidate within this many times the first one's distance is shown, and flagged;
                            # a main left out within this many times the distance is flagged as well
SNAP_FAR_M = 30.0           # a snap farther than this is flagged
LOO_FLAG_M = 10.0           # a control tap whose leave-one-out residual is above this is flagged
LOO_LEVERAGE_MAX = 0.99     # a control tap whose leverage is above this cannot be checked: without it the others lie
                            # almost on one line, so its leave-one-out residual measures nothing (found in the final
                            # verification's fix pass: a right control tap got 644.7 m and a flag)
MIN_CONTROLS, CHECKED_CONTROLS = 3, 4   # 3 control taps fix the affine's 6 terms exactly, with nothing to check them
COLLINEAR_MAX_COND = 20.0   # nearly on one line: the control taps' centred east-north positions, largest over smallest
                            # singular value, above this,
COLLINEAR_MIN_AREA = 0.02   # or their largest triangle over the square of their largest distance below this,
ACROSS_MIN_M = 50.0         # and their RMS distance from their own best line under this: refused (a long, narrow zone
                            # whose control taps spread this far across it is accepted, with a warning)
SAME_PLACE_M = 50.0         # control taps all within this many metres of each other were taken at about one place:
                            # refused (the final verification of task 16b)
FIT_GPS_ASSUMED_M = 5.0     # the control taps' GPS error assumed (radial RMS, a phone outdoors) when the fit's error at
FIT_PREDICTED_FLAG_M = 10.0     # each tap is predicted from its leverage; a tap whose predicted error is above this is
                                # flagged (far from the control taps, or across the line they lie along)
FIT_SCALE_TOL = 0.2         # a fit whose scale is more than 20% off the model's coordinate units warns
FLOW_ZERO_M3S = 1e-7        # a pipe whose every flow in the day window is at most this (0.1 mL/s) has no flow
DRAWN_RATIO_BAND = (0.5, 2.0)   # a pipe drawn at under half or over twice the model's median drawn length per metre of
                                # declared length is left out of the snap (a node drawn at a placeholder such as 0,0)
DRAWN_SHARE_WARN = 0.9      # fewer of the model's pipes drawn (both ends with a position) than this share warns
GEOGRAPHIC_MAX_RATIO = 1e-4     # a drawing within +-180 on both axes (latitude as x or as y) whose median drawn length
                                # per metre is below this is drawn in degrees of longitude and latitude
FAR_NODE_SPANS = 100.0      # a node farther from the drawing's median position than this many times the larger
                            # interquartile range of its x and y is far outside the drawing (a no-data value such as
                            # -3.4e38): treated as not drawn, with a warning, as is a node whose position is not a
                            # number (the final verification of task 16b)
NEAR_MODEL_FACTOR, NEAR_MODEL_MIN_M = 1000.0, 10_000.0   # a position farther from the drawing's centre than this many
                                # times the drawing's extent, and this many metres, is not near the model
SNAP_CHUNK_ELEMENTS = 2_000_000     # taps times segments per chunk of the distance matrix (a few hundred MB at most)
WGS84_A, WGS84_F = 6378137.0, 1 / 298.257223563
SUGGESTED_NAME = "tap_map_suggested.csv"
SUGGESTED_COLUMNS = ("tap_id", "junction_id", "snap_m", "pipe", "second_choice", "second_m", "flags")
_LAT_EXACT = ("lat", "latitude", "gps lat", "gps latitude", "lat dd", "latitude dd", "lat deg", "latitude deg",
              "y lat", "lat y", "wgs84 lat", "lat wgs84", "wgs84 latitude")
_LON_EXACT = ("lon", "long", "lng", "longitude", "gps lon", "gps long", "gps longitude", "lon dd", "long dd",
              "longitude dd", "lon deg", "longitude deg", "x lon", "lon x", "x long", "wgs84 lon", "lon wgs84",
              "wgs84 longitude")
_X_EXACT = ("x", "easting", "east", "x coord", "x coordinate", "xcoord", "point x", "model x", "x model", "map x")
_Y_EXACT = ("y", "northing", "north", "y coord", "y coordinate", "ycoord", "point y", "model y", "y model", "map y")
_ON_MAIN_EXACT = ("on main", "onmain", "on large main", "on transmission main", "transmission main", "large main")
_YES = {"yes", "y", "true", "1", "x"}
_COORD_NUM = re.compile(r"^([NSEWnsew])?\s*([-+]?\d+(?:[.,]\d+)?(?:[eE][-+]?\d+)?)\s*°?\s*([NSEWnsew])?$")
_THOUSANDS = re.compile(r"^[-+]?\d{1,3}(?:,\d{3})+$")           # '1,197': a thousands comma, or a decimal comma
_THOUSANDS_DOT = re.compile(r"^[-+]?\d{1,3}(?:,\d{3})+\.\d*$")  # '1,197.5': a thousands comma, for certain
_DOT_GROUP = re.compile(r"^[-+]?[1-9]\d{0,2}\.\d{3}$")          # '1.197': a decimal point, or a dot thousands separator
_DOT_THOUSANDS = re.compile(r"^[-+]?\d{1,3}(?:\.\d{3}){2,}$|^[-+]?\d{1,3}(?:\.\d{3})+,\d*$")   # '1.234.567', '1.197,5':
                                                                # a dot thousands separator, for certain
_HEMISPHERE = {"lat": "NS", "lon": "EW", "xy": ""}             # the letters each kind of column may carry
_COLUMN_WORDS = {"tap": "tap", "junction": "junction", "lat": "latitude", "lon": "longitude", "x": "x", "y": "y",
                 "on_main": "on_main"}
_USE_IDS = "Use junction IDs instead: a tap map with two columns, the tap ID and the junction ID."


class GeorefRefused(LogRefused):
    """Taps in latitude and longitude that cannot be placed on the model: too few control taps, or control taps that
    lie almost on one line."""


@dataclass
class TapCoords:
    """A tap map by coordinates, read: kind 'latlon' (a, b are latitude and longitude in degrees) or 'xy' (a, b are
    x and y in the model's own coordinates).  taps: one dict per tap (tap_id, row, a, b, junction, on_main); a or b
    None for a row that gives only a junction; junction '' when none is given (a junction given with latitude and
    longitude makes a control tap).  dropped: every row that is not used, with its reason.  other_columns: the
    columns not used, with their cells (suggest_from says when one of them holds the model's junction IDs)."""
    kind: str
    taps: list = field(default_factory=list)
    dropped: list = field(default_factory=list)
    columns: dict = field(default_factory=dict)
    notes: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    n_rows: int = 0
    other_columns: dict = field(default_factory=dict)


def guess_coordinate_columns(header: list[str]) -> dict:
    """A tap map's coordinate columns, by exact header names only: {'lat', 'lon', 'x', 'y', 'on_main'}."""
    return {k: _guess(header, ex) for k, ex in (("lat", _LAT_EXACT), ("lon", _LON_EXACT), ("x", _X_EXACT),
                                                 ("y", _Y_EXACT), ("on_main", _ON_MAIN_EXACT))}


def _tap_columns(header: list[str], coords: dict) -> dict:
    """The tap and junction columns of a tap map by coordinates, guessed as guess_tap_columns guesses them, with the
    coordinate and on_main columns left out ('Point X' is no tap column)."""
    taken = {v for v in coords.values() if v}
    rest = [h for h in header if h not in taken]
    tap = _guess(rest, _TAP_EXACT, _TAP_WORDS)
    junction = _guess([h for h in rest if h != tap], _JUNCTION_EXACT, ("junction", "node"))
    return {"tap": tap, "junction": junction}


def _given(mapping: dict | None) -> dict:
    return {k: v for k, v in (mapping or {}).items() if v}


def _columns_of(t: LogTable, mapping: dict | None) -> tuple[dict, dict]:
    """(coordinate columns, tap and junction columns) of a tap map: guessed, then any column the mapping names."""
    m = _given(mapping)
    c = guess_coordinate_columns(t.header)
    c.update({k: m[k] for k in ("lat", "lon", "x", "y", "on_main") if k in m})
    tc = _tap_columns(t.header, c)
    tc.update({k: m[k] for k in ("tap", "junction") if k in m})
    return c, tc


def _pairs_in(t: LogTable, c: dict) -> list[tuple[str, str, str]]:
    """The coordinate pairs a tap map holds: ('latlon', lat column, lon column) and ('xy', x column, y column)."""
    return [(kind, c[a], c[b]) for kind, a, b in (("latlon", "lat", "lon"), ("xy", "x", "y"))
            if c.get(a) in t.header and c.get(b) in t.header]


def _filled(t: LogTable, column: str | None) -> list[bool]:
    if column not in t.header:
        return [False] * len(t.rows)
    i = t.header.index(column)
    return [bool(_cell_text(r[i])) for r in t.rows]


def _choose_pair(t: LogTable, c: dict, junction: str | None, mapping: dict | None):
    """The pair of columns that places the taps: (kind, column a, column b, a note or None), or None for a table with
    no pair.  A pair the mapping names wins.  With both pairs in the file, the one filled on more rows; on a tie,
    latitude and longitude, unless they have fewer than MIN_CONTROLS control taps (rows with a junction too), which
    their fit needs and x and y do not."""
    m = _given(mapping)
    pairs = {kind: (a, b) for kind, a, b in _pairs_in(t, c)}
    if not pairs:
        return None
    if len(pairs) == 1:
        kind, (a, b) = next(iter(pairs.items()))
        return kind, a, b, None
    ll, xy = pairs["latlon"], pairs["xy"]
    named_ll, named_xy = bool({"lat", "lon"} & set(m)), bool({"x", "y"} & set(m))
    if named_xy and not named_ll:
        return "xy", *xy, f"The columns '{ll[0]}' and '{ll[1]}' are not used: the taps are placed by the x and y chosen."
    if named_ll and not named_xy:
        return "latlon", *ll, (f"The columns '{xy[0]}' and '{xy[1]}' are not used: the taps are placed by the latitude "
                               f"and longitude chosen.")
    fl = [a and b for a, b in zip(_filled(t, ll[0]), _filled(t, ll[1]))]
    fx = [a and b for a, b in zip(_filled(t, xy[0]), _filled(t, xy[1]))]
    nl, nx = sum(fl), sum(fx)
    ctrl = sum(f and j for f, j in zip(fl, _filled(t, junction)))
    if nx > nl or (nx == nl and nx > 0 and ctrl < MIN_CONTROLS):
        why = ("x and y are filled on more rows" if nx > nl else
               f"latitude and longitude have {ctrl} control tap{'s' if ctrl != 1 else ''} (a row with a junction too), "
               f"fewer than the {MIN_CONTROLS} their fit needs, and x and y need none")
        return "xy", *xy, (f"The columns '{ll[0]}' and '{ll[1]}' are not used: {why} ({nx} rows filled against {nl}). "
                           f"To place the taps by latitude and longitude, choose those columns.")
    why = "latitude and longitude are filled on more rows" if nl > nx else "both pairs are filled"
    return "latlon", *ll, (f"The columns '{xy[0]}' and '{xy[1]}' are not used: {why} ({nl} rows filled against {nx}). "
                           f"To place the taps by x and y in the model's own coordinates, choose those columns.")


def tap_map_kind(t: LogTable, mapping: dict | None = None) -> str:
    """'coordinates' for a tap map that places taps by position (a latitude and longitude pair or an x and y pair of
    columns, and no junction column, or a row with a position and no junction); 'id' for every other table, which
    read_tap_map reads as before.  The columns are guessed, or named in mapping as read_tap_map and
    read_tap_coordinates take it (a junction column named there is used, so a map by ID whose junction header is
    known only through the mapping stays 'id')."""
    c, tc = _columns_of(t, mapping)
    pairs = _pairs_in(t, c)
    if not pairs:
        return "id"
    if tc["junction"] not in t.header:
        return "coordinates"
    ij = t.header.index(tc["junction"])
    pos = [(t.header.index(a), t.header.index(b)) for _, a, b in pairs]
    return "coordinates" if any(not _cell_text(r[ij]) and any(_cell_text(r[ia]) or _cell_text(r[ib]) for ia, ib in pos)
                                for r in t.rows) else "id"


def tap_map_kind_of(data: bytes, name: str = "tap_map.csv", mapping: dict | None = None) -> str:
    """tap_map_kind of a file; 'id' when it cannot be read as a table (read_tap_map then says why)."""
    try:
        return tap_map_kind(read_table(data, name), mapping)
    except LogRefused:
        return "id"


def coordinate_value(v, axis: str | None = None) -> float | None:
    """A coordinate cell as a number: '45.123', '45,123' (a decimal comma), '-122.4', '122.4 W', 'N 45.1', '45.1°'
    (S and W are negative); None for a blank or anything else (degrees, minutes and seconds included).  axis 'lat'
    takes only the letters N and S, 'lon' only E and W, 'xy' none; None takes any of the four."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return x if math.isfinite(x) else None
    m = _COORD_NUM.match(str(v).strip().replace(" ", " "))
    if not m or (m.group(1) and m.group(3)):
        return None
    letter = (m.group(1) or m.group(3) or "").upper()
    if letter and axis is not None and letter not in _HEMISPHERE[axis]:
        return None
    x = float(m.group(2).replace(",", "."))
    if not math.isfinite(x):
        return None
    return -abs(x) if letter in ("S", "W") else x


def _wrong_letter(cell: str, axis: str) -> str | None:
    """The letter of a cell such as '40.0 W' when it belongs to the other axis ('W' in a latitude column)."""
    m = _COORD_NUM.match(str(cell).strip().replace(" ", " "))
    letter = (m.group(1) or m.group(3) or "").upper() if m else ""
    return letter if letter and letter not in _HEMISPHERE[axis] and axis in ("lat", "lon") else None


def _comma_mode(cells: list) -> str:
    """How a column of x or y reads a value such as '1,197' (a comma then three digits: a thousands comma, or a
    decimal comma): 'thousands' when another value of the column settles it so ('1,197.5', '12,345,678', or a value
    of 1,000 or more written without a comma), 'decimal' when another value has a decimal comma ('0,5', '1197,25'),
    and 'ambiguous' when nothing does, or both do."""
    thou = dec = False
    for v in cells:
        if isinstance(v, bool) or v is None:
            continue
        if isinstance(v, (int, float)):
            thou |= math.isfinite(float(v)) and abs(float(v)) >= 1000
            continue
        s = str(v).strip()
        if _THOUSANDS_DOT.match(s) or (_THOUSANDS.match(s) and s.count(",") > 1):
            thou = True
        elif _DOT_THOUSANDS.match(s):          # '1.197,5', '1.234.567': dots for thousands, so a comma is decimal
            dec = True
        elif _THOUSANDS.match(s):
            continue
        elif "," in s:
            dec |= coordinate_value(s, "xy") is not None
        else:
            x = coordinate_value(s, "xy")
            thou |= x is not None and abs(x) >= 1000
    return "thousands" if thou and not dec else "decimal" if dec and not thou else "ambiguous"


def _dot_mode(cells: list, active: bool) -> str:
    """The mirror of _comma_mode, for a value such as '1.197' (a dot then three digits) in a column whose file writes
    decimal commas (active: the column has a decimal comma, or the file is separated by semicolons): 'thousands' when
    another value settles it so ('1.197,5', '1.234.567', or a value of 1,000 or more written without a dot),
    'decimal' when another value has a dot that can only be a decimal point ('503.25', '0.125', '1,197.5'), and
    'ambiguous' when nothing does, or both do.  Not active, a dot is a decimal point ('decimal'), as before.  The
    final verification of task 16b: '1.197' for 1197 was read as 1.197, and the suggestion looked right."""
    if not active:
        return "decimal"
    thou = dec = False
    for v in cells:
        if isinstance(v, bool) or v is None:
            continue
        if isinstance(v, (int, float)):
            thou |= math.isfinite(float(v)) and abs(float(v)) >= 1000
            continue
        s = str(v).strip()
        if _DOT_THOUSANDS.match(s):
            thou = True
        elif _THOUSANDS_DOT.match(s) or (_THOUSANDS.match(s) and s.count(",") > 1):
            dec = True
        elif _DOT_GROUP.match(s):
            continue
        elif "." in s:
            dec |= coordinate_value(s, "xy") is not None
        else:
            x = coordinate_value(s, "xy")
            thou |= x is not None and abs(x) >= 1000
    return "thousands" if thou and not dec else "decimal" if dec and not thou else "ambiguous"


def _xy_value(v, mode: str, dot: str = "decimal") -> tuple[float | None, str | None]:
    """An x or y cell as a number, its commas read by the column's mode (_comma_mode) and its dots by its dot mode
    (_dot_mode): (value, None), or (None, 'comma' or 'dot') for a value that reads two ways."""
    if isinstance(v, str):
        s = v.strip()
        if _DOT_THOUSANDS.match(s):
            return coordinate_value(s.replace(".", ""), "xy"), None
        if _DOT_GROUP.match(s) and dot != "decimal":
            if dot == "ambiguous":
                return None, "dot"
            return coordinate_value(s.replace(".", ""), "xy"), None
        if _THOUSANDS_DOT.match(s) or (_THOUSANDS.match(s) and s.count(",") > 1):
            return coordinate_value(s.replace(",", ""), "xy"), None
        if _THOUSANDS.match(s):
            if mode == "ambiguous":
                return None, "comma"
            return coordinate_value(s.replace(",", "") if mode == "thousands" else s, "xy"), None
    return coordinate_value(v, "xy"), None


def _lon_0_to_360(t: LogTable, ia: int, ib: int) -> bool:
    """True for longitudes written from 0 to 360, as some receivers and data sets write them (259.4 for 100.6 W): every
    longitude that reads as a number is between 0 and 360 and one is over 180 (the final verification of task 16b:
    such a file was refused as 'not degrees')."""
    lons = [x for x in (coordinate_value(r[ib], "lon") for r in t.rows) if x is not None]
    return bool(lons) and any(x > 180 for x in lons) and all(0 <= x <= 360 for x in lons)


def read_tap_coordinates(data: bytes, name: str = "tap_map.csv", mapping: dict | None = None) -> TapCoords:
    """A tap map by coordinates read into a TapCoords (columns guessed, or given in mapping: 'tap', 'junction', 'lat',
    'lon', 'x', 'y', 'on_main'; a column named there that the file does not have is refused).  With both pairs of
    columns, the pair is chosen by _choose_pair, with a note.  Every row that is not used is listed with its reason:
    no tap ID, a tap listed again (its first row is used; a warning when the rows differ), no position and no
    junction, a position that is not a number or reads two ways ('1,197' in x or y with nothing to settle it, or
    '1.197' in a file that writes decimal commas), a hemisphere letter of the other axis, a latitude or longitude out
    of range.  Longitudes written from 0 to 360 are read west of Greenwich, with a note.  Raises LogRefused when the
    file has no coordinate pair, no tap column, a tap given two different junctions, or latitudes and longitudes that
    are no degrees (and says so when the two columns look swapped)."""
    t = read_table(data, name)
    m = _given(mapping)
    missing = [(k, v) for k, v in m.items() if k in _COLUMN_WORDS and v not in t.header]
    if missing:
        raise LogRefused(f"The tap map has no column {' or '.join(repr(v) for _, v in missing)}, chosen for the "
                         f"{' and '.join(_COLUMN_WORDS[k] for k, _ in missing)}; its header is "
                         f"{', '.join(t.header[:8])}. Choose one of its columns.")
    c, tc = _columns_of(t, m)
    pair = _choose_pair(t, c, tc["junction"], m)
    if pair is None:
        raise LogRefused(f"The tap map has no columns of positions (lat and lon, or the model's x and y); its header is "
                         f"{', '.join(t.header[:8])}. Name them lat and lon (or x and y), or give each tap's junction "
                         f"ID.")
    kind, ca, cb, pair_note = pair
    if tc["tap"] not in t.header:
        raise LogRefused(f"The tap map needs a tap column; its header is {', '.join(t.header[:8])}. Name it tap_id.")
    out = TapCoords(kind=kind, columns={"tap": tc["tap"], "junction": tc["junction"], "a": ca, "b": cb,
                                        "on_main": c["on_main"]}, n_rows=len(t.rows))
    ix = {k: (t.header.index(v) if v in t.header else None) for k, v in out.columns.items()}
    words = ("latitude", "longitude") if kind == "latlon" else ("x", "y")
    axes = ("lat", "lon") if kind == "latlon" else ("xy", "xy")
    modes = [_comma_mode([r[ix[k]] for r in t.rows]) if kind == "xy" else "decimal" for k in ("a", "b")]
    semicolons = t.how.startswith(f"text separated by {_DELIMS[';']}")
    dots = [_dot_mode([r[ix[k]] for r in t.rows], semicolons or m == "decimal") if kind == "xy" else "decimal"
            for k, m in zip(("a", "b"), modes)]
    lon360 = kind == "latlon" and _lon_0_to_360(t, ix["a"], ix["b"])
    lon360_example = None
    seen: dict[str, dict] = {}
    clash, differ, out_of_range = [], [], []
    in_range = 0
    for r, row_no in zip(t.rows, t.row_numbers):
        tap = _cell_text(r[ix["tap"]])
        cells = [_cell_text(r[ix[k]]) for k in ("a", "b")]
        j = _cell_text(r[ix["junction"]]) if ix["junction"] is not None else ""
        if not tap:
            if any(cells) or j:
                out.dropped.append({"row": row_no, "tap_id": "", "reason": "no tap ID"})
            continue
        if tap in seen:
            first = seen[tap]
            if j and first["junction"] and j != first["junction"]:
                clash.append(tap)
            elif (cells, j) != (first["cells"], first["junction"]):
                differ.append((tap, first["row"], row_no))
            out.dropped.append({"row": row_no, "tap_id": tap,
                                "reason": f"tap {tap} is listed again: its first row (row {first['row']}) is used"})
            continue
        if not any(cells):
            if not j:
                out.dropped.append({"row": row_no, "tap_id": tap,
                                    "reason": f"no position ('{ca}' and '{cb}' blank) and no junction"})
                continue
            a = b = None
        else:
            if kind == "latlon":
                a, b = coordinate_value(r[ix["a"]], "lat"), coordinate_value(r[ix["b"]], "lon")
                amb_a = amb_b = None
                if lon360 and b is not None and b > 180:
                    lon360_example = lon360_example or (cells[1], b - 360.0)
                    b -= 360.0
            else:
                (a, amb_a), (b, amb_b) = (_xy_value(r[ix["a"]], modes[0], dots[0]),
                                          _xy_value(r[ix["b"]], modes[1], dots[1]))
            if amb_a or amb_b:
                ex, how = (cells[0], amb_a) if amb_a else (cells[1], amb_b)
                sep, other = ("comma", "decimal point") if how == "comma" else ("dot", "decimal comma")
                out.dropped.append({"row": row_no, "tap_id": tap, "example": ex,
                                    "reason": f"the {words[0]} or {words[1]} reads two ways ('{ex}' is a number of "
                                              f"thousands, or one with a decimal {'comma' if how == 'comma' else 'point'}"
                                              f"): write it without the {sep}, or with a {other}"})
                continue
            if a is None or b is None:
                k = 0 if a is None else 1
                letter = _wrong_letter(cells[k], axes[k])
                other = {"lat": "longitude", "lon": "latitude"}.get(axes[k])
                out.dropped.append({"row": row_no, "tap_id": tap, "example": cells[k],
                                    "reason": (f"the {words[k]} '{cells[k]}' has the letter {letter}, which belongs to "
                                               f"a {other}: are the two columns swapped?") if letter else
                                              f"the {words[0]} or {words[1]} is not a number"})
                continue
            if kind == "latlon" and not (abs(a) <= 90 and abs(b) <= 180):
                out_of_range.append((a, b, f"{cells[0]}, {cells[1]}"))
                out.dropped.append({"row": row_no, "tap_id": tap, "reason": "the latitude or longitude is out of range",
                                    "example": f"{cells[0]}, {cells[1]}"})
                continue
        in_range += a is not None
        on_main = ix["on_main"] is not None and _cell_text(r[ix["on_main"]]).lower() in _YES
        seen[tap] = {"row": row_no, "cells": cells, "junction": j}
        out.taps.append({"tap_id": tap, "row": row_no, "a": a, "b": b, "junction": j, "on_main": on_main})
    if clash:
        raise LogRefused(f"The tap map lists {len(set(clash))} tap{'s' if len(set(clash)) > 1 else ''} with two "
                         f"different junctions: {', '.join(sorted(set(clash))[:6])}. Keep one junction per tap.")
    if kind == "latlon" and out_of_range and not in_range:
        if all(90 < abs(a) <= 180 and abs(b) <= 90 for a, b, _ in out_of_range):
            raise LogRefused(f"The columns '{ca}' and '{cb}' look swapped: '{ca}' holds numbers between 90 and 180 (for "
                             f"example {out_of_range[0][2]}), which only a longitude can be, and '{cb}' numbers of at "
                             f"most 90. Swap the two column names.")
        raise LogRefused(f"The columns '{ca}' and '{cb}' hold numbers that are not degrees of latitude and longitude "
                         f"(for example {out_of_range[-1][2]}): if they are the model's own x and y, name the columns "
                         f"x and y.")
    if differ:
        out.warnings.append(f"{len(differ)} tap{'s are' if len(differ) != 1 else ' is'} listed again with another "
                            f"position or junction ("
                            + "; ".join(f"{tp}, rows {a} and {b}" for tp, a, b in differ[:6])
                            + ("; and more" if len(differ) > 6 else "")
                            + "): the first row is used. Check which row is right.")
    placed = sum(1 for x in out.taps if x["a"] is not None)
    given = sum(1 for x in out.taps if x["junction"])
    how = "latitude and longitude" if kind == "latlon" else "x and y in the model's own coordinates"
    out.notes.append(f"Tap map by coordinates: {len(out.taps)} tap{'s' if len(out.taps) != 1 else ''}, placed by {how} "
                     f"(columns '{ca}' and '{cb}'); {given} with a junction given"
                     + (" (control taps, when they have a position too)" if kind == "latlon" and given else "") + ".")
    if pair_note:
        out.notes.append(pair_note)
    if lon360_example:
        out.notes.append(f"The longitudes are written from 0 to 360, as some receivers write them: each one over 180 is "
                         f"read as that minus 360, west of Greenwich ('{lon360_example[0]}' as "
                         f"{lon360_example[1]:.6f}).")
    if kind == "xy" and placed and all(abs(x["a"]) <= 180 and abs(x["b"]) <= 90 for x in out.taps if x["a"] is not None):
        out.warnings.append(f"The columns '{ca}' and '{cb}' hold numbers that look like degrees: if they are longitude "
                            f"and latitude, name the columns lon and lat.")
    used = {tc["tap"], tc["junction"], *(v for v in c.values() if v)}
    out.other_columns = {h: [_cell_text(r[i]) for r in t.rows] for i, h in enumerate(t.header) if h not in used}
    return out


# ---- the model's geometry and its day-window flows
@dataclass
class TapGeometry:
    """The pipes a tap can be placed on: every pipe with both ends drawn (except those drawn out of proportion,
    left_out), its segments (vertices included) in metres around the drawing's centre (origin: the median position
    of these pipes' end nodes), its diameter in inches and its end nodes.  extent_m: the farthest segment end from the
    centre.  node_xy holds the drawn nodes, without those whose position is not a number or is far outside the
    drawing."""
    path: str
    scale: dict
    metres_per_unit: float
    origin: tuple
    pipes: list
    diam_in: object
    start: list
    end: list
    seg_a: object
    seg_b: object
    seg_start: object                 # each pipe's first segment (the segments are in pipe order)
    junctions: set
    node_xy: dict                     # every drawn node: (x, y) in the model's own units
    node_kind: dict                   # 'junction', 'tank' or 'reservoir'
    warnings: list = field(default_factory=list)
    left_out: list = field(default_factory=list)
    extent_m: float = 0.0


def _drawn_length(pts) -> float:
    return sum(math.dist(pts[k], pts[k + 1]) for k in range(len(pts) - 1))


def _far_nodes(xy: dict) -> list:
    """The nodes drawn far outside the rest of the drawing: farther from the nodes' median position, in x or in y,
    than FAR_NODE_SPANS times the larger of the interquartile ranges of their x and y (the width of the drawing's
    middle half).  None with fewer than 4 nodes, or when that width is 0."""
    import numpy as np
    if len(xy) < 4:
        return []
    names = list(xy)
    A = np.array([xy[n] for n in names], float)
    med = np.median(A, axis=0)
    q25, q75 = np.percentile(A, [25, 75], axis=0)
    span = float(np.max(q75 - q25))
    if not span > 0:
        return []
    d = np.abs(A - med).max(axis=1)
    return sorted(n for n, v in zip(names, d) if v > FAR_NODE_SPANS * span)


def tap_geometry(path: str) -> TapGeometry:
    """The model's pipes for placing taps by coordinates (path: a file, or the name of a network bundled with WNTR).
    A schematic drawing, a model with no coordinates or too few pipes drawn (coordinate_scale, as the readiness report
    types them), and a drawing in degrees of longitude and latitude, are refused, with 'use junction IDs'.  To scale
    in feet or metres, distances are in metres; to scale in other units, the units' size is the median drawn length
    per metre of declared length (unrounded), with a warning.  A pipe drawn out of proportion to the rest (its drawn
    length over its declared length outside DRAWN_RATIO_BAND times the median) is left out, with a warning, and so is
    a model of which fewer than DRAWN_SHARE_WARN of the pipes are drawn warned about.  Before anything is measured, a
    node whose position is not a number or is far outside the rest (_far_nodes) is treated as not drawn, with a
    warning that names it and its pipes."""
    import numpy as np
    import wntr
    from .simulate import net_path
    p = net_path(path)
    with open(p, "rb") as fh:
        text = decode_text(fh.read())[0]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wn = wntr.network.WaterNetworkModel(p)
    have_xy = scan_inp_ids(text).coordinate_ids & set(wn.node_name_list)
    # the final verification of task 16b: a node whose position is not a number, or one drawn far outside the rest
    # (a no-data value such as -3.4e38), is treated as not drawn before anything is measured, so it cannot move the
    # drawing's centre, its units or the refusals below
    raw = {n: tuple(map(float, wn.get_node(n).coordinates)) for n in have_xy}
    not_number = sorted(n for n, v in raw.items() if not all(math.isfinite(c) for c in v))
    finite = {n: v for n, v in raw.items() if n not in set(not_number)}
    far = _far_nodes(finite)
    xy = {n: v for n, v in finite.items() if n not in set(far)}
    have_xy = set(xy)
    cs = coordinate_scale(wn, have_xy)
    if cs["status"] == "no coordinates":
        raise LogRefused("This model has no coordinates, so a tap cannot be placed on it by its position. " + _USE_IDS)
    if cs["status"] == "too few pipes drawn":
        raise LogRefused("Too few of this model's pipes are drawn to tell whether its drawing is to scale, so a tap "
                         "cannot be placed on it by its position. " + _USE_IDS)
    if cs["status"] == "schematic":
        raise LogRefused("This model's drawing is a schematic, not to scale: its drawn pipe lengths are not in "
                         "proportion to the lengths in the file, so a tap cannot be placed on it by its position. "
                         + _USE_IDS)
    drawn = [pn for pn in wn.pipe_name_list
             if wn.get_link(pn).start_node_name in have_xy and wn.get_link(pn).end_node_name in have_xy]
    ratio, odd_vertex = {}, []
    for pn in drawn:
        pl = wn.get_link(pn)
        pts = [xy[pl.start_node_name], *pl.vertices, xy[pl.end_node_name]]
        if not all(math.isfinite(float(c)) for pt in pts for c in pt):
            odd_vertex.append(pn)
        elif pl.length > 0:
            ratio[pn] = _drawn_length(pts) / pl.length
    med = float(np.median([v for v in ratio.values() if v > 0]))       # coordinate_scale's median, not rounded
    # degrees on both axes, latitude as x or as y (the final verification of task 16b: a drawing with latitude as x
    # passed a test that wanted x within 180 and y within 90)
    if (med < GEOGRAPHIC_MAX_RATIO and max(abs(v[0]) for v in xy.values()) <= 180
            and max(abs(v[1]) for v in xy.values()) <= 180):
        raise LogRefused(f"This model is drawn in degrees of longitude and latitude, not in feet or metres (about "
                         f"{med:.2g} of its units per metre of pipe), so a tap cannot be placed on it by its position. "
                         + _USE_IDS)
    warns = []
    at = lambda names: sorted({pn for pn in wn.pipe_name_list for n in (wn.get_link(pn).start_node_name,
                                                                         wn.get_link(pn).end_node_name) if n in names})
    for names, what in ((not_number, "a position that is not a number"),
                        (far, f"a position far outside the rest of the drawing (more than {FAR_NODE_SPANS:g} times "
                              f"the width of its middle half from its middle), as an export writes a position it does "
                              f"not have (a no-data value)")):
        if names:
            k, pp = len(names), at(set(names))
            them = "them" if k != 1 else "it"
            gone = (f", so the pipe at {them} ({pp[0]}) is left out of the placement" if len(pp) == 1 else
                    f", so the {len(pp)} pipes at {them} ({', '.join(pp[:6])}{', and more' if len(pp) > 6 else ''}) "
                    f"are left out of the placement" if pp else "")
            warns.append(f"{k} node{'s have' if k != 1 else ' has'} {what}: "
                         + ", ".join(f"{n} at {raw[n][0]:.6g}, {raw[n][1]:.6g}" for n in names[:4])
                         + (", and more" if k > 4 else "")
                         + f". {'They are' if k != 1 else 'It is'} treated as not drawn{gone}.")
    if cs["units"] == "feet":
        mpu = 0.3048
    elif cs["units"] == "metres":
        mpu = 1.0
    else:
        mpu = 1.0 / med
        warns.append(f"The model's drawing is to scale, in units that are neither feet nor metres (about "
                     f"{med:.3g} per metre of pipe): distances use that scale.")
    lo, hi = DRAWN_RATIO_BAND
    left_out = [pn for pn in drawn if pn in odd_vertex or (pn in ratio and not lo * med <= ratio[pn] <= hi * med)]
    if left_out:
        n = len(left_out)
        warns.append(f"{n} pipe{'s are' if n != 1 else ' is'} left out of the placement ("
                     f"{', '.join(left_out[:6])}{', and more' if n > 6 else ''}): drawn at under half or over twice "
                     f"the model's usual drawn length per metre of pipe, as when a node is drawn at a placeholder "
                     f"position such as 0,0. A tap beside {'one of them' if n != 1 else 'it'} is put on the nearest "
                     f"other pipe.")
    if len(drawn) < DRAWN_SHARE_WARN * len(wn.pipe_name_list):
        warns.append(f"Only {len(drawn)} of the model's {len(wn.pipe_name_list)} pipes are drawn (both ends with a "
                     f"position): a tap beside a pipe that is not drawn is put on the nearest drawn pipe, which may be "
                     f"the wrong one.")
    out_set = set(left_out)
    pipes, diam, starts, ends, a, b, seg_start = [], [], [], [], [], [], []
    for pn in drawn:
        if pn in out_set:
            continue
        pl = wn.get_link(pn)
        pts = [xy[pl.start_node_name], *[tuple(map(float, v)) for v in pl.vertices], xy[pl.end_node_name]]
        seg_start.append(len(a))
        for k in range(len(pts) - 1):
            a.append(pts[k])
            b.append(pts[k + 1])
        pipes.append(pn)
        diam.append(pl.diameter / 0.0254)
        starts.append(pl.start_node_name)
        ends.append(pl.end_node_name)
    # the drawing's centre: the median position of the kept pipes' end nodes, not the mean of every node (the final
    # verification of task 16b: one node at -3.4e38 moved the mean so far that every position lost its precision)
    kept = sorted(set(starts) | set(ends))
    origin = ((float(np.median([xy[n][0] for n in kept])), float(np.median([xy[n][1] for n in kept]))) if kept else
              (0.0, 0.0))
    o = np.array(origin)
    kind = {n: "junction" for n in wn.junction_name_list}
    kind.update({n: "tank" for n in wn.tank_name_list})
    kind.update({n: "reservoir" for n in wn.reservoir_name_list})
    sa = (np.array(a, float).reshape(-1, 2) - o) * mpu
    sb = (np.array(b, float).reshape(-1, 2) - o) * mpu
    extent = float(max(np.hypot(*sa.T).max(), np.hypot(*sb.T).max())) if len(sa) else 0.0
    return TapGeometry(path=p, scale=cs, metres_per_unit=mpu, origin=origin, pipes=pipes,
                       diam_in=np.round(np.array(diam, float), 3), start=starts, end=ends,
                       seg_a=sa, seg_b=sb, seg_start=np.array(seg_start, int), junctions=set(wn.junction_name_list),
                       node_xy=xy, node_kind=kind, warnings=warns, left_out=left_out, extent_m=extent)


def tap_hydraulics(path: str, start_clock_h: float = 0.0) -> dict:
    """One 7-day run of the model, set up as simulate.load sets up every run (with water age), in a temporary folder:
    each pipe's flow (m3/s, positive from its start node to its end node) at each hour of the clock day window, 07:00
    to 17:00 (daytime_model_hours), on the scored day (the last), and each node's mean water age (h) over that
    window."""
    import wntr
    from .simulate import DAY, DURATION_DAYS, load, net_path
    wn = load(net_path(path))
    wn.options.quality.parameter = "AGE"
    with tempfile.TemporaryDirectory(prefix="rm_epanet_") as d, warnings.catch_warnings():
        warnings.simplefilter("ignore")
        res = wntr.sim.EpanetSimulator(wn).run_sim(file_prefix=os.path.join(d, "run"))
    hours = daytime_model_hours(start_clock_h)
    times = [(DURATION_DAYS - 1) * DAY + h * 3600 for h in hours]
    fl = res.link["flowrate"].loc[times, wn.pipe_name_list]
    age = res.node["quality"].loc[times, wn.node_name_list].mean() / 3600.0
    return {"flow": {p: [float(x) for x in fl[p].values] for p in wn.pipe_name_list},
            "age_h": {n: float(v) for n, v in age.items()},
            "window": (clock_label(hours[0], start_clock_h), clock_label(hours[-1], start_clock_h))}


def pipe_junctions(geom: TapGeometry, hyd: dict) -> list[tuple]:
    """For each pipe of geom: (junction, other end, codes).  The junction is the pipe's downstream end by the sign of
    its mean flow over the day window, never by drawing direction.  codes: 'reverses' (some hour of the window flows
    against the mean, or the mean is no flow while the hours are not: the other end is given too), 'no_flow' (every
    hour at most FLOW_ZERO_M3S: the end with the older water, by mean water age over the window; equal ages take the
    first end by name), 'end_not_junction' (the downstream end is a tank or reservoir: the other end).  A pipe with
    no junction at either end gives (None, None, codes): no tap is put on it."""
    import numpy as np
    out = []
    for p, s, e in zip(geom.pipes, geom.start, geom.end):
        q = np.array(hyd["flow"][p], float)
        codes = []
        m = float(q.mean())
        older = lambda: (s if hyd["age_h"].get(s, 0.0) > hyd["age_h"].get(e, 0.0) else
                         e if hyd["age_h"].get(e, 0.0) > hyd["age_h"].get(s, 0.0) else min(s, e))
        if float(np.abs(q).max()) <= FLOW_ZERO_M3S:
            codes.append("no_flow")
            down = older()
        elif abs(m) <= FLOW_ZERO_M3S:
            codes.append("reverses")
            down = older()
        else:
            down = e if m > 0 else s
            if bool(np.any((np.sign(q) == -np.sign(m)) & (np.abs(q) > FLOW_ZERO_M3S))):
                codes.append("reverses")
        up = s if down == e else e
        if down not in geom.junctions:
            if up not in geom.junctions:
                out.append((None, None, tuple(codes)))
                continue
            codes.append("end_not_junction")
            down, up = up, down
        out.append((down, up if ("reverses" in codes and up in geom.junctions) else None, tuple(codes)))
    return out


def _segment_distances(P, A, B):
    """Distances (k, n) from k points to n segments AB, every coordinate in metres."""
    import numpy as np
    AB = B - A
    L2 = (AB ** 2).sum(1)
    L2 = np.where(L2 > 0, L2, 1.0)
    PA = P[:, None, :] - A[None, :, :]
    u = np.clip((PA * AB[None]).sum(2) / L2[None], 0.0, 1.0)
    C = A[None] + u[..., None] * AB[None]
    return np.sqrt(((P[:, None, :] - C) ** 2).sum(2))


def snap_points(P, allow_main, geom: TapGeometry, rules: list[tuple], main_min_in: float = MAIN_MIN_IN,
                chunk: int | None = None, cut_in: float | None = None):
    """The nearest pipe to each point (k, 2), in metres around geom.origin, among the pipes with a junction and, unless
    allow_main is True for that point, under main_min_in inches (judged at MAIN_CUT_RATIO of it on the model's
    diameter, or under cut_in inches on it when given); and the nearest pipe whose junction differs from the first
    one's (the second candidate).  Returns
    (first, d1, second, d2, main_d): pipe indices into geom.pipes (-1 for none), distances in metres (inf for none),
    and the distance to the nearest pipe left out as a main (inf where mains are allowed, or there is none).  The
    points go in chunks of at most SNAP_CHUNK_ELEMENTS points times segments, so memory does not grow with the
    model's size beyond that."""
    import numpy as np
    P = np.asarray(P, float).reshape(-1, 2)
    allow = np.broadcast_to(np.asarray(allow_main, bool), (len(P),))
    names = sorted({j for j, _, _ in rules if j is not None})
    code = {j: i for i, j in enumerate(names)}
    pj = np.array([code[j] if j is not None else -1 for j, _, _ in rules], int)
    has_j = pj >= 0
    small = geom.diam_in < (main_min_in * MAIN_CUT_RATIO if cut_in is None else cut_in)
    main = has_j & ~small
    chunk = chunk or max(1, min(256, SNAP_CHUNK_ELEMENTS // max(1, len(geom.seg_a))))
    first, d1, second, d2, dmain = (np.full(len(P), -1, int), np.full(len(P), np.inf), np.full(len(P), -1, int),
                                    np.full(len(P), np.inf), np.full(len(P), np.inf))
    for k0 in range(0, len(P), chunk):
        Pc = P[k0:k0 + chunk]
        Dall = np.minimum.reduceat(_segment_distances(Pc, geom.seg_a, geom.seg_b), geom.seg_start, axis=1)
        ok = has_j[None, :] & (small[None, :] | allow[k0:k0 + chunk, None])
        D = np.where(ok, Dall, np.inf)
        f = np.argmin(D, axis=1)
        fd = D[np.arange(len(Pc)), f]
        D2 = np.where(pj[None, :] != pj[f][:, None], D, np.inf)
        s = np.argmin(D2, axis=1)
        sd = D2[np.arange(len(Pc)), s]
        first[k0:k0 + chunk] = np.where(np.isfinite(fd), f, -1)
        d1[k0:k0 + chunk] = fd
        second[k0:k0 + chunk] = np.where(np.isfinite(sd), s, -1)
        d2[k0:k0 + chunk] = sd
        if main.any():
            dm = np.where(main[None, :], Dall, np.inf).min(axis=1)
            dmain[k0:k0 + chunk] = np.where(allow[k0:k0 + chunk], np.inf, dm)
    return first, d1, second, d2, dmain


# ---- georeferencing: WGS84 latitude and longitude to the model's x and y
def _ecef(lat, lon):
    import numpy as np
    e2 = WGS84_F * (2 - WGS84_F)
    p, lam = np.radians(lat), np.radians(lon)
    n = WGS84_A / np.sqrt(1 - e2 * np.sin(p) ** 2)
    return np.column_stack([n * np.cos(p) * np.cos(lam), n * np.cos(p) * np.sin(lam), n * (1 - e2) * np.sin(p)])


def enu_from_latlon(lat, lon, lat0: float, lon0: float):
    """East and north in metres of each point on the local tangent plane at (lat0, lon0): WGS84 to earth-centred
    coordinates (ECEF), then the plane's east and north axes (the height above the ellipsoid is taken as 0)."""
    import numpy as np
    x = _ecef(np.atleast_1d(lat), np.atleast_1d(lon)) - _ecef(np.array([lat0]), np.array([lon0]))
    p, lam = np.radians(lat0), np.radians(lon0)
    east = np.array([-np.sin(lam), np.cos(lam), 0.0])
    north = np.array([-np.sin(p) * np.cos(lam), -np.sin(p) * np.sin(lam), np.cos(p)])
    return np.column_stack([x @ east, x @ north])


def centre_longitude(lon, how: str = "median") -> float:
    """The taps' middle longitude, their median (how 'median'; 'mean' as first built), taken across the antimeridian
    when they straddle it (longitudes written from -180 to 180 that span more than 180 degrees: -179.9 and 179.9 give
    180, not 0)."""
    import numpy as np
    lon = np.asarray(lon, float)
    f = np.median if how == "median" else np.mean
    if lon.size and float(lon.max() - lon.min()) > 180:
        m = float(f(np.where(lon < 0, lon + 360.0, lon)))
        return m - 360.0 if m > 180 else m
    return float(f(lon))


def plane_centre(lat, lon, how: str = "median") -> tuple[float, float]:
    """The east-north plane's tangent point: the taps' median latitude and longitude (centre_longitude), so one tap
    with a dropped digit cannot move it (the final verification of task 16b: with the mean of every tap, one typo in a
    small tap map moved it by degrees and put correct taps 13 m off); how 'mean' gives the plane as first built."""
    import numpy as np
    lat = np.asarray(lat, float)
    return float(np.median(lat) if how == "median" else lat.mean()), centre_longitude(lon, how)


def fit_affine(src, dst):
    """The least-squares 2D affine from src (n, 2) to dst (n, 2): a (3, 2) matrix M with dst ~ [src, 1] M."""
    import numpy as np
    A = np.column_stack([np.asarray(src, float), np.ones(len(src))])
    M, *_ = np.linalg.lstsq(A, np.asarray(dst, float), rcond=None)
    return M


def apply_affine(M, src):
    import numpy as np
    return np.column_stack([np.asarray(src, float), np.ones(len(src))]) @ M


def _hull_indices(E) -> list[int]:
    """The corners of the convex hull of points E (n, 2), as indices in ascending order (Andrew's monotone chain)."""
    order = sorted(range(len(E)), key=lambda i: (float(E[i][0]), float(E[i][1])))
    if len(order) < 3:
        return sorted(order)
    cross = lambda o, a, b: ((E[a][0] - E[o][0]) * (E[b][1] - E[o][1]) - (E[a][1] - E[o][1]) * (E[b][0] - E[o][0]))
    lower, upper = [], []
    for i in order:
        while len(lower) >= 2 and cross(lower[-2], lower[-1], i) <= 0:
            lower.pop()
        lower.append(i)
    for i in reversed(order):
        while len(upper) >= 2 and cross(upper[-2], upper[-1], i) <= 0:
            upper.pop()
        upper.append(i)
    return sorted(set(lower[:-1] + upper[:-1]))


def control_spread(enu_ctrl) -> dict:
    """How far the control taps are from lying on one line: the largest over the smallest singular value of their
    centred east-north positions (cond), their RMS distance in metres from their own best line (across_m), and their
    largest triangle's area over the square of their largest distance (area).  The largest distance and the largest
    triangle are both measured on the corners of the taps' convex hull, where they lie, so both use the same set
    (at most 150 corners for the triangle, evenly spaced, and 2,000 for the distance)."""
    import itertools
    import numpy as np
    E = np.asarray(enu_ctrl, float).reshape(-1, 2)
    s = np.linalg.svd(E - E.mean(0), compute_uv=False) if len(E) else np.zeros(0)
    cond = float(s[0] / s[1]) if len(s) > 1 and s[1] > 0 else math.inf
    across = float(s[1] / math.sqrt(len(E))) if len(s) > 1 else 0.0
    H = E[_hull_indices(E)] if len(E) else E
    Hd = H[np.linspace(0, len(H) - 1, 2000).astype(int)] if len(H) > 2000 else H
    dmax = float(np.hypot(*(Hd[:, None, :] - Hd[None, :, :]).transpose(2, 0, 1)).max()) if len(Hd) > 1 else 0.0
    area = 0.0
    Ht = H[np.linspace(0, len(H) - 1, 150).astype(int)] if len(H) > 150 else H
    if dmax > 0 and len(Ht) >= 3:
        idx = np.array(list(itertools.combinations(range(len(Ht)), 3)))
        a, b, c = Ht[idx[:, 0]], Ht[idx[:, 1]], Ht[idx[:, 2]]
        area = float(np.max(np.abs((b[:, 0] - a[:, 0]) * (c[:, 1] - a[:, 1]) - (b[:, 1] - a[:, 1]) * (c[:, 0] - a[:, 0])))
                     / 2 / dmax ** 2)
    return {"cond": cond, "area": area, "largest_distance_m": dmax, "across_m": across}


def _places(E, radius: float, most: int = 3) -> int:
    """How many places the points E (n, 2) were taken at, counting a point within radius of an earlier place as that
    place (at most `most`)."""
    centres = []
    for p in E:
        if not any(math.dist(p, c) <= radius for c in centres):
            centres.append(p)
            if len(centres) >= most:
                break
    return len(centres)


def georeference(lat, lon, controls: list[int], control_xy, metres_per_unit: float = 1.0,
                 names: list[str] | None = None, centre: str = "median") -> tuple:
    """Taps in WGS84 latitude and longitude put in the model's x and y: each goes to east and north in metres on the
    local tangent plane at the taps' median position (plane_centre; centre 'mean' gives the plane as first built),
    then through the 2D affine fitted from the control taps (indices controls, their model positions control_xy, in
    model units).  Returns (xy, info): xy (n, 2) in model units, info with the control count, whether the fit is
    checked (4 or more control taps), each control tap's leave-one-out residual in metres (the fit made without it,
    which needs 4 or more: in closed form, its residual over one minus its leverage; None, and listed as uncheckable,
    when its leverage is over LOO_LEVERAGE_MAX), the flagged ones (above LOO_FLAG_M), each tap's predicted error in metres (predicted_m: FIT_GPS_ASSUMED_M of GPS error on the control taps
    carried through the fit to that tap by its leverage, which grows with its distance from them and across the line
    they lie along), the fit's scale against the model's units, whether it mirrors, and whether the control taps lie
    along a narrow strip.  Raises GeorefRefused with fewer than MIN_CONTROLS control taps, with control taps all within
    SAME_PLACE_M of each other, or with control taps that lie almost on one line (control_spread: the relative tests,
    and fewer than ACROSS_MIN_M metres across)."""
    import numpy as np
    lat, lon = np.asarray(lat, float), np.asarray(lon, float)
    names = list(names) if names is not None else [str(i) for i in range(len(lat))]
    k = len(controls)
    if k < MIN_CONTROLS:
        raise GeorefRefused(f"Taps given by latitude and longitude need at least {MIN_CONTROLS} control taps, taps whose "
                            f"junction you know, with the position taken at that junction ({CHECKED_CONTROLS} or more, "
                            f"so the fit can be checked); this tap map has {k}. Give a junction ID on {MIN_CONTROLS - k} "
                            f"more tap{'s' if MIN_CONTROLS - k != 1 else ''} spread over the area, or give the taps' "
                            f"junction IDs.")
    enu = enu_from_latlon(lat, lon, *plane_centre(lat, lon, centre))
    ce, cxy = enu[controls], np.asarray(control_xy, float)
    sp = control_spread(ce)
    if sp["largest_distance_m"] < SAME_PLACE_M:
        where = ("at the same place" if sp["largest_distance_m"] < 0.5 else
                 f"at about the same place (within {sp['largest_distance_m']:.1f} m of each other)")
        raise GeorefRefused(f"The {k} control taps were all taken {where}, so the fit has nothing to measure the map's "
                            f"turn and stretch with. Take the control taps at junctions spread over the area: at least "
                            f"{MIN_CONTROLS} of them not on one line, and {CHECKED_CONTROLS} or more so the fit can be "
                            f"checked.")
    long_thin, small_area = sp["cond"] > COLLINEAR_MAX_COND, sp["area"] < COLLINEAR_MIN_AREA
    if (long_thin or small_area) and sp["across_m"] < ACROSS_MIN_M:
        if _places(ce, SAME_PLACE_M / 2) == 2:
            raise GeorefRefused(f"The {k} control taps were taken at only two places, about "
                                f"{sp['largest_distance_m']:.0f} m apart, so they lie on one line and the fit cannot tell "
                                f"the map's turn and stretch across it. Add a control tap at a junction well away from "
                                f"that line.")
        why = ([f"their positions are {sp['cond']:.0f} times longer than wide, above {COLLINEAR_MAX_COND:g}"
                if math.isfinite(sp["cond"]) else "their positions lie exactly on one line"] if long_thin else []) + \
              ([f"their largest triangle is {sp['area']:.3f} of the square of their largest distance, below "
                f"{COLLINEAR_MIN_AREA:g}"] if small_area else []) + \
              [f"they lie within about {sp['across_m']:.1f} m of that line, under {ACROSS_MIN_M:g} m"]
        raise GeorefRefused(f"The {k} control taps lie almost on one line, so the fit cannot tell the map's turn and "
                            f"stretch across that line ({'; '.join(why)}). Add a control tap well away from that "
                            f"line.")
    c0 = cxy.mean(0)
    M = fit_affine(ce, cxy - c0)
    xy = apply_affine(M, enu) + c0
    # leverage from the control taps' design matrix [east, north, 1], centred: A = QR, h_ii the rows' squared norms of
    # Q; any position's leverage solves R^T z = [e, n, 1]
    cm = ce.mean(0)
    Q, R = np.linalg.qr(np.column_stack([ce - cm, np.ones(k)]))
    h = (Q ** 2).sum(1)
    lev = (np.linalg.solve(R.T, np.column_stack([enu - cm, np.ones(len(enu))]).T) ** 2).sum(0)
    loo = []
    if k >= CHECKED_CONTROLS:
        # the final verification of task 16b: in closed form (O(n)), not one refit per control tap (2,000 control
        # taps took 0.4 s and 16,000 took 23 s); a control tap of leverage over LOO_LEVERAGE_MAX gets None
        resid = (cxy - c0) - apply_affine(M, ce)
        for i in range(k):
            r = (float(np.hypot(*(resid[i] / (1.0 - h[i])))) * metres_per_unit if h[i] <= LOO_LEVERAGE_MAX else None)
            loo.append({"tap_id": names[controls[i]], "residual_m": r})
    sv = np.linalg.svd(M[:2], compute_uv=False) * metres_per_unit      # model units per metre, in metres per metre
    measured = [x for x in loo if x["residual_m"] is not None]
    info = {"controls": k, "checked": k >= CHECKED_CONTROLS, "loo": loo,
            "loo_max_m": max((x["residual_m"] for x in measured), default=None),
            "flagged": [x["tap_id"] for x in measured if x["residual_m"] > LOO_FLAG_M],
            "uncheckable": [x["tap_id"] for x in loo if x["residual_m"] is None],
            "scale": [float(sv[0]), float(sv[1])], "mirrored": bool(np.linalg.det(M[:2]) < 0),
            "cond": sp["cond"], "area": sp["area"], "across_m": sp["across_m"],
            "narrow_strip": bool(long_thin or small_area),
            "predicted_m": [float(x) for x in FIT_GPS_ASSUMED_M * np.sqrt(lev)]}
    return xy, info


# ---- the suggestion
@dataclass
class TapSuggestion:
    """The suggested tap map: one row per tap (SUGGESTED_COLUMNS, flags a list of plain words), the rows of the tap
    map that are not used with their reasons, notes, warnings, and the georeference's info (latitude and longitude
    only).  A suggestion for the operator to check, never a tap map by itself: the app's table and confirm step and
    the pilot's --accept-snapped, which would take a checked one, are not built yet (task 16b stopped)."""
    kind: str
    rows: list = field(default_factory=list)
    dropped: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    georef: dict = field(default_factory=dict)
    n_rows: int = 0

    def mapping(self) -> dict:
        return {r["tap_id"]: r["junction_id"] for r in self.rows if r["junction_id"]}

    def table(self) -> list[dict]:
        """The rows as the CSV holds them: snap_m and second_m with one decimal, flags joined by '; '."""
        f = lambda v: "" if v is None else f"{v:.1f}"
        return [{"tap_id": r["tap_id"], "junction_id": r["junction_id"] or "", "snap_m": f(r["snap_m"]),
                 "pipe": r["pipe"] or "", "second_choice": r["second_choice"] or "", "second_m": f(r["second_m"]),
                 "flags": "; ".join(r["flags"])} for r in self.rows]

    def to_csv(self) -> bytes:
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=SUGGESTED_COLUMNS, lineterminator="\n")
        w.writeheader()
        w.writerows(self.table())
        return buf.getvalue().encode("utf-8")


def _flag_words(codes: tuple, other: str | None, window: tuple, kind_of: dict, up_name: str | None) -> list:
    """The plain words of pipe_junctions' codes for one tap."""
    out = []
    w = f"between {window[0]} and {window[1]}"
    if "no_flow" in codes:
        out.append(f"no flow {w}: the end with the older water")
    if "reverses" in codes:
        out.append(f"the flow reverses {w}: it may be the other end, junction {other}" if other else
                   f"the flow reverses {w}")
    if "end_not_junction" in codes and up_name:
        out.append(f"the downstream end is a {kind_of.get(up_name, 'node')} ({up_name}): the other end")
    return out


def suggest_from(coords: TapCoords, geom: TapGeometry, hyd: dict, main_min_in: float = MAIN_MIN_IN) -> TapSuggestion:
    """The suggested junction of each tap in coords on the model of geom (its flows in hyd).  A tap with a junction
    given keeps it (in latitude and longitude it is a control tap, with its leave-one-out residual; in x and y its
    distance from the junction's drawn position is checked); every other tap is put on the nearest pipe under
    main_min_in inches (any pipe when its row says on_main), its segments and vertices in metres, and gets that
    pipe's downstream junction (pipe_junctions); the nearest pipe with another junction is the second choice, shown
    and flagged when within SECOND_RATIO times the distance; a main left out within SECOND_RATIO times the distance is
    flagged; a snap over SNAP_FAR_M is flagged; an untagged tap put on a pipe from MAIN_BAND_RATIO of main_min_in up
    to the cut (a main given by its internal diameter may be there) is flagged with the nearest smaller pipe; a tap
    placed by latitude and longitude where the fit's predicted error is above FIT_PREDICTED_FLAG_M is flagged.  A
    given junction that is not in the model, a control tap's junction with no drawn position, and a position that is
    not near the model's drawing are listed as not used; a column left unused that holds the model's junction IDs is
    warned about."""
    import numpy as np
    rules = pipe_junctions(geom, hyd)
    cut = main_min_in * MAIN_CUT_RATIO
    band_lo = main_min_in * MAIN_BAND_RATIO
    has_j = np.array([j is not None for j, _, _ in rules], bool)
    sug = TapSuggestion(kind=coords.kind, dropped=list(coords.dropped), notes=list(coords.notes),
                        warnings=list(coords.warnings) + list(geom.warnings), n_rows=coords.n_rows)
    for h, vals in coords.other_columns.items():
        filled = [v for v in vals if v]
        hits = sum(v in geom.junctions for v in filled)
        if filled and hits >= max(1.0, len(filled) / 2):
            sug.warnings.append(f"The column '{h}' holds junction IDs of this model ({hits} of its {len(filled)} filled "
                                f"rows) and is not used: if it gives each tap's junction, choose it as the junction "
                                f"column, and those taps keep their junctions.")
    taps = []
    for t in coords.taps:
        if t["junction"] and t["junction"] not in geom.junctions:
            sug.dropped.append({"row": t["row"], "tap_id": t["tap_id"],
                                "reason": f"junction {t['junction']} (tap {t['tap_id']}) is not in the model"})
        elif coords.kind == "latlon" and t["junction"] and t["a"] is not None and t["junction"] not in geom.node_xy:
            sug.dropped.append({"row": t["row"], "tap_id": t["tap_id"],
                                "reason": f"junction {t['junction']} (tap {t['tap_id']}) has no position in the model's "
                                          f"drawing, so it cannot be a control tap"})
        else:
            taps.append(t)
    pos = [t for t in taps if t["a"] is not None]
    xy, predicted = {}, {}
    if coords.kind == "latlon" and pos:
        ctrl = [i for i, t in enumerate(pos) if t["junction"]]
        need = any(not t["junction"] for t in pos)
        if need or len(ctrl) >= MIN_CONTROLS:
            pts, info = georeference([t["a"] for t in pos], [t["b"] for t in pos], ctrl,
                                     [geom.node_xy[pos[i]["junction"]] for i in ctrl], geom.metres_per_unit,
                                     [t["tap_id"] for t in pos])
            xy = {t["tap_id"]: (float(p[0]), float(p[1])) for t, p in zip(pos, pts)}
            predicted = {t["tap_id"]: v for t, v in zip(pos, info["predicted_m"]) if not t["junction"]}
            sug.georef = info
            loo = info["loo"]
            measured = [x for x in loo if x["residual_m"] is not None]
            if info["checked"]:
                if measured:
                    worst = max(measured, key=lambda x: x["residual_m"])
                    sug.notes.append(f"Placed by latitude and longitude through a local east-north plane and an affine "
                                     f"fit to {info['controls']} control taps; each control tap's leave-one-out residual "
                                     f"(the fit made without it) is at most {worst['residual_m']:.1f} m (tap "
                                     f"{worst['tap_id']}).")
            else:
                sug.warnings.append(f"Unchecked: the fit uses {info['controls']} control taps, which fix it exactly and "
                                    f"leave nothing to check it with. Add a fourth control tap (a junction ID on a tap "
                                    f"whose position was taken at that junction).")
            if info["uncheckable"]:
                n = len(info["uncheckable"])
                sug.warnings.append(f"Control tap{'s' if n != 1 else ''} {', '.join(info['uncheckable'][:6])} cannot be "
                                    f"checked: without {'each' if n != 1 else 'it'}, the other control taps lie almost "
                                    f"on one line, so the fit made without {'each' if n != 1 else 'it'} measures "
                                    f"nothing. A control tap away from that line would let "
                                    f"{'them' if n != 1 else 'it'} be checked.")
            if info["narrow_strip"]:
                sug.warnings.append(f"The control taps lie along a narrow strip ({info['cond']:.0f} times longer than "
                                    f"wide, about {info['across_m']:.0f} m across): a tap far from that strip may be "
                                    f"placed less well. A control tap away from it would help.")
            if info["flagged"]:
                bad = sorted((x for x in measured if x["residual_m"] > LOO_FLAG_M), key=lambda x: -x["residual_m"])
                sug.warnings.append(f"Control tap{'s' if len(bad) != 1 else ''} "
                                    f"{', '.join(x['tap_id'] for x in bad)}: the fit made without "
                                    f"{'each' if len(bad) != 1 else 'it'} misses its junction by more than "
                                    f"{LOO_FLAG_M:g} m ({', '.join(f'{x['residual_m']:.1f} m' for x in bad)}). "
                                    + ("One wrong control tap pulls the others' residuals up too: check the "
                                       f"largest first, {bad[0]['tap_id']}, its junction ID and where its position was "
                                       f"taken." if len(bad) > 1 else
                                       "Check that tap's junction ID and where its position was taken."))
            if max(abs(s - 1) for s in info["scale"]) > FIT_SCALE_TOL:
                sug.warnings.append(f"The fit stretches the positions by {info['scale'][0]:.2f} and "
                                    f"{info['scale'][1]:.2f} against the model's coordinate units (1 when they agree): "
                                    f"check the control taps' junctions and the model's coordinate units.")
            if info["mirrored"]:
                sug.warnings.append("The fit mirrors the map: a control tap's junction ID is probably wrong, or the "
                                    "latitude and longitude columns are swapped.")
    elif coords.kind == "xy":
        xy = {t["tap_id"]: (t["a"], t["b"]) for t in pos}
    o = np.array(geom.origin)
    snap = [t for t in taps if not t["junction"] and t["tap_id"] in xy]
    far_off = set()
    res, alt = {}, {}
    if snap:
        P = (np.array([xy[t["tap_id"]] for t in snap], float) - o) * geom.metres_per_unit
        limit = max(NEAR_MODEL_FACTOR * geom.extent_m, NEAR_MODEL_MIN_M)
        with np.errstate(over="ignore", invalid="ignore"):
            near = np.hypot(P[:, 0], P[:, 1]) <= limit
        for t, ok in zip(snap, near):
            if not ok:
                far_off.add(t["tap_id"])
                sug.dropped.append({"row": t["row"], "tap_id": t["tap_id"],
                                    "reason": f"the position is not near the model's drawing (more than "
                                              f"{limit / 1000:,.0f} km from its centre): check its "
                                              + ("latitude and longitude" if coords.kind == "latlon" else "x and y")})
        snap = [t for t, ok in zip(snap, near) if ok]
        if snap:
            first, d1, second, d2, dmain = snap_points(P[near], [t["on_main"] for t in snap], geom, rules, main_min_in)
            res = {t["tap_id"]: (int(f), float(a), int(s), float(b), float(m))
                   for t, f, a, s, b, m in zip(snap, first, d1, second, d2, dmain)}
            # the final verification of task 16b: a 16 in plastic main given by its internal diameter is under the
            # cut, so an untagged tap put on a pipe in the band below it is flagged with the nearest smaller pipe
            band = [i for i, (t, f) in enumerate(zip(snap, first))
                    if f >= 0 and not t["on_main"] and band_lo <= geom.diam_in[f] < cut]
            if band:
                af, ad, *_ = snap_points(P[near][band], False, geom, rules, main_min_in, cut_in=band_lo)
                alt = {snap[i]["tap_id"]: (int(f), float(d)) for i, f, d in zip(band, af, ad)}
    taps = [t for t in taps if t["tap_id"] not in far_off]
    loo = {x["tap_id"]: x["residual_m"] for x in sug.georef.get("loo", [])}
    unchecked = coords.kind == "latlon" and sug.georef and not sug.georef.get("checked")
    for t in taps:
        row = {"tap_id": t["tap_id"], "junction_id": "", "snap_m": None, "pipe": None, "second_choice": None,
               "second_m": None, "flags": []}
        if t["junction"]:
            row["junction_id"] = t["junction"]
            if coords.kind == "latlon" and t["tap_id"] in xy:
                r = loo.get(t["tap_id"])
                row["flags"].append("control tap: junction given" + (
                    f"; leave-one-out residual {r:.1f} m" if r is not None else
                    "; its leave-one-out residual cannot be measured: without it the others lie almost on one line"
                    if t["tap_id"] in sug.georef.get("uncheckable", []) else ""))
                if r is not None and r > LOO_FLAG_M:
                    row["flags"].append(f"over {LOO_FLAG_M:g} m: check this control tap's junction and position")
            else:
                row["flags"].append("junction given in the tap map")
                if coords.kind == "xy" and t["a"] is not None and t["junction"] in geom.node_xy:
                    d = math.dist(geom.node_xy[t["junction"]], (t["a"], t["b"])) * geom.metres_per_unit
                    if d > SNAP_FAR_M:
                        row["flags"].append(f"the junction given is {d:.0f} m from the tap's position")
        elif t["tap_id"] in res:
            f, a, s, b, dm = res[t["tap_id"]]
            if f < 0:                    # worded by what happened (the final verification of task 16b)
                if not has_j.any():
                    row["flags"].append("no pipe to place it on: no drawn pipe has a junction at either end")
                elif not t["on_main"] and not (has_j & (geom.diam_in < cut)).any():
                    row["flags"].append(f"no pipe to place it on: every drawn pipe with a junction at one end is "
                                        f"{main_min_in:g} in or more (say on_main yes if the tap is on one)")
                else:
                    row["flags"].append("no pipe to place it on: its distance to the drawn pipes could not be "
                                        "measured; check its position")
            else:
                j, other, codes = rules[f]
                row.update(junction_id=j, snap_m=a, pipe=geom.pipes[f])
                if s >= 0 and b <= SECOND_RATIO * a:
                    row.update(second_choice=rules[s][0], second_m=b)
                    row["flags"].append(f"two pipes close: the second choice, junction {rules[s][0]} on pipe "
                                        f"{geom.pipes[s]}, is {b:.1f} m away")
                if dm <= SECOND_RATIO * a:
                    row["flags"].append(f"a main of {main_min_in:g} in or more is {dm:.1f} m away: if the tap is on it, "
                                        f"say on_main yes")
                if t["tap_id"] in alt:
                    af, ad = alt[t["tap_id"]]
                    row["flags"].append(
                        f"a pipe of {geom.diam_in[f]:.1f} in, under the {cut:.1f} in cut for mains: if it is a main of "
                        f"{main_min_in:g} in or more given by its internal diameter, "
                        + (f"and the tap is not on it, the nearest smaller pipe is {geom.pipes[af]} (junction "
                           f"{rules[af][0]}), {ad:.1f} m away; if the tap is on it, say on_main yes" if af >= 0 else
                           "check that the tap is on it, and say on_main yes if it is"))
                if a > SNAP_FAR_M:
                    row["flags"].append(f"far: {a:.0f} m from the nearest pipe, over {SNAP_FAR_M:g} m")
                up_name = (geom.end[f] if j == geom.start[f] else geom.start[f]) if "end_not_junction" in codes else None
                row["flags"] += _flag_words(codes, other, hyd["window"], geom.node_kind, up_name)
                if geom.diam_in[f] >= cut:
                    row["flags"].append(f"on a main of {main_min_in:g} in or more (on_main yes)")
                if unchecked:
                    row["flags"].append("the fit is unchecked: add a fourth control tap")
                if predicted.get(t["tap_id"], 0.0) > FIT_PREDICTED_FLAG_M:
                    row["flags"].append(f"the fit is uncertain here, about {predicted[t['tap_id']]:.0f} m (the control "
                                        f"taps' GPS error, taken as {FIT_GPS_ASSUMED_M:g} m, carried this far from "
                                        f"them): a control tap near it would help")
        else:
            row["flags"].append("no position and no junction")
        sug.rows.append(row)
    placed = [r for r in sug.rows if r["snap_m"] is not None]
    if placed:
        n_flag = sum(1 for r in placed if r["flags"])
        sug.notes.append(f"{len(placed)} tap{'s' if len(placed) != 1 else ''} put on the nearest pipe under "
                         f"{main_min_in:g} in (any pipe for a tap whose row says on_main yes), each given that pipe's "
                         f"downstream junction by the simulated flow between {hyd['window'][0]} and "
                         f"{hyd['window'][1]}; {n_flag} flagged for a look. A flag marks some of the wrong suggestions, "
                         f"not all of them: check every row.")
        unsure = [r["tap_id"] for r in placed if predicted.get(r["tap_id"], 0.0) > FIT_PREDICTED_FLAG_M]
        if unsure:
            n = len(unsure)
            sug.warnings.append(f"{n} tap{'s lie' if n != 1 else ' lies'} where the fit is uncertain ("
                                f"{', '.join(unsure[:6])}{', and more' if n > 6 else ''}): far from the control taps, "
                                f"or across the line they lie along, so {FIT_GPS_ASSUMED_M:g} m of GPS error on each "
                                f"control tap could move {'them' if n != 1 else 'it'} by more than "
                                f"{FIT_PREDICTED_FLAG_M:g} m. {'Each is' if n != 1 else 'It is'} flagged; a control tap "
                                f"near {'them' if n != 1 else 'it'} would help.")
        far = [r for r in placed if r["snap_m"] > 1000]
        if len(far) == len(placed) and coords.kind == "xy":
            sug.warnings.append("Every tap is more than 1 km from the model's pipes: check that x and y are in the "
                                "model's own coordinates (the same system and units as the model's drawing).")
    return sug


def suggest_tap_map(taps_data: bytes, taps_name: str, model_path: str, start_clock_h: float = 0.0,
                    mapping: dict | None = None, main_min_in: float = MAIN_MIN_IN) -> TapSuggestion:
    """A tap map by coordinates read (read_tap_coordinates), checked against the model's drawing (tap_geometry: a
    schematic drawing is refused), the model run once (tap_hydraulics) and every tap given its suggested junction
    (suggest_from).  Raises LogRefused (GeorefRefused) with the reason in plain words; any other error on the way is
    raised as a LogRefused that names it, so a page or the pilot shows a message, never a traceback."""
    try:
        coords = read_tap_coordinates(taps_data, taps_name, mapping)
        geom = tap_geometry(model_path)
        hyd = tap_hydraulics(model_path, start_clock_h)
        return suggest_from(coords, geom, hyd, main_min_in)
    except LogRefused:
        raise
    except Exception as e:  # noqa: BLE001
        raise LogRefused(f"The taps could not be placed on this model ({type(e).__name__}: {e}). " + _USE_IDS) from e


# ----------------------------------------------------------------------------- command line
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+")
    ap.add_argument("--prepare", action="store_true", help="prepare each EPANET model and print its readiness report")
    ap.add_argument("--out", default=os.path.join("outputs", "private", "inputs"),
                    help="where a fixed copy goes (default outputs/private/inputs, git-ignored)")
    ap.add_argument("--disinfectant", choices=sorted(DISINFECTANT_WORDS), default="free_chlorine")
    ap.add_argument("--json", action="store_true", help="print JSON instead of plain text")
    a = ap.parse_args(argv)
    bad = 0
    for f in a.files:
        if a.prepare:
            path, rep = prepare_inp(f, a.out, a.disinfectant)
            print(rep.to_json() if a.json else rep.to_markdown())
            if path:
                print(f"prepared: {path}\n")
            bad += path is None
        else:
            with open(f, "rb") as fh:
                s = sniff(fh.read(), os.path.basename(f))
            print(json.dumps({"file": f, **s.as_dict()}) if a.json else f"{os.path.basename(f)}: {s.sentence()} {s.advice}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
